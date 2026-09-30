import copy, sys, json, os
from docx import Document
from docx.shared import Pt, Inches, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH as AL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from PIL import Image, ImageOps

PAGES = json.load(open('pages.json')) if os.path.exists('pages.json') else [4,5,6,7,15]
FONT = 'Times New Roman'
d = Document('template.docx')
body = d.element.body
P = lambda i: docx_par(i)
from docx.text.paragraph import Paragraph
paras = [Paragraph(e, d) for e in body if e.tag == qn('w:p')]

def set_text(p, text):
    p.runs[0].text = text
    for r in p.runs[1:]: r._r.getparent().remove(r._r)

# ---------- cover page ----------
cover = {p.text: p for p in paras}
set_text(cover['<project title in caps>'], 'STUDENT REGISTRATION FORM FOR AN ONLINE COURSE PORTAL')
for r in cover['<project title in caps>'].runs: r.bold = True
namep = cover['<Name (Rollno)>']
students = [('AISHWARYA U', '71812401006'), ('AJAY IYANRAJ', '71812401007'),
            ('AKASHVARMAN R', '71812401008'), ('AKHILESH RAJ P', '71812401009')]
set_text(namep, f'{students[0][0]} ({students[0][1]})')
prev = namep._p
for n, r_ in students[1:]:
    new = copy.deepcopy(namep._p); prev.addnext(new); prev = new
    Paragraph(new, d).runs[0].text = f'{n} ({r_})'
# remove a few spare blank paragraphs after names so the cover keeps one page
nxt = prev.getnext(); removed = 0
while nxt is not None and removed < 3 and nxt.tag == qn('w:p') and not ''.join(nxt.itertext()).strip():
    n2 = nxt.getnext(); nxt.getparent().remove(nxt); nxt = n2; removed += 1

# explicit page breaks instead of relying on blank filler paragraphs
def blank(e): return e.tag == qn('w:p') and not ''.join(e.itertext()).strip()
def drop_blanks_after(par_el, keep=0):
    n = par_el.getnext(); k = 0
    while n is not None and blank(n):
        nn = n.getnext()
        if k >= keep: n.getparent().remove(n)
        k += 1; n = nn
drop_blanks_after(cover['20CS212 &WEB TECHNOLOGIES']._p, keep=2)
last_cover = [p for p in paras if 'Mr.N.Manoj' in p.text][0]
drop_blanks_after(last_cover._p)
dept2 = [p for p in paras if p.text.startswith('DEPARTMENT')][1]
dept2.paragraph_format.page_break_before = True
sig2 = [p for p in paras if 'Mr.N.Manoj' in p.text][1]
drop_blanks_after(sig2._p)

# ---------- index table ----------
t0, t1 = d.tables[0], d.tables[1]
for i, (n, r_) in enumerate(students, 1):
    cell = t0.rows[i].cells[1]
    p = cell.paragraphs[0]
    for r in list(p.runs): r._r.getparent().remove(r._r)
    p.paragraph_format.space_before = Pt(0); p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.first_line_indent = None
    run = p.add_run(n); run.bold = True; run.font.size = Pt(11); run.font.name = FONT
    p2 = cell.add_paragraph(); p2.paragraph_format.space_after = Pt(0)
    r2 = p2.add_run(r_); r2.font.size = Pt(11); r2.font.name = FONT
    p2.paragraph_format.left_indent = p.paragraph_format.left_indent

# ---------- contents table page numbers ----------
for i in range(1, 6):
    cell = t1.rows[i].cells[2]
    p = cell.paragraphs[0]
    p.alignment = AL.CENTER
    r = p.add_run(str(PAGES[i-1])); r.font.size = Pt(12); r.font.name = FONT

# ---------- helpers for new content ----------
sect = body.find(qn('w:sectPr'))
def add_el(el): sect.addprevious(el)

def para(text='', bold=False, size=12, align=AL.JUSTIFY, italic=False, space_after=6, keep=False, font=FONT, indent=None):
    p = d.add_paragraph()   # appended before sectPr by python-docx
    p.alignment = align
    pf = p.paragraph_format
    pf.space_after = Pt(space_after); pf.space_before = Pt(0); pf.line_spacing = 1.15
    if keep: pf.keep_with_next = True
    if indent: pf.left_indent = Inches(indent)
    if text:
        r = p.add_run(text); r.bold = bold; r.italic = italic; r.font.size = Pt(size); r.font.name = font
        r._r.rPr.rFonts.set(qn('w:eastAsia'), font)
    return p

def rich(parts, size=12, align=AL.JUSTIFY, space_after=4, indent=None, hanging=None):
    p = para('', align=align, space_after=space_after, indent=indent)
    for txt, b in parts:
        r = p.add_run(txt); r.bold = b; r.font.size = Pt(size); r.font.name = FONT
    if hanging: p.paragraph_format.first_line_indent = Inches(-hanging)
    return p

def heading(num, text, page_break=False):
    p = para(f'{num}. {text.upper()}', bold=True, size=14, align=AL.CENTER, space_after=8, keep=True)
    p.paragraph_format.space_before = Pt(10)
    if page_break: p.paragraph_format.page_break_before = True
    # bottom border
    pPr = p._p.get_or_add_pPr(); b = OxmlElement('w:pBdr'); bt = OxmlElement('w:bottom')
    for k, v in (('val', 'single'), ('sz', '8'), ('space', '1'), ('color', '000000')): bt.set(qn('w:' + k), v)
    b.append(bt)
    anc = next((c for c in pPr if c.tag in [qn('w:' + n) for n in ('shd', 'tabs', 'suppressAutoHyphens', 'spacing', 'ind', 'jc', 'rPr')]), None)
    if anc is not None: anc.addprevious(b)
    else: pPr.append(b)
    return p

def sub(text):
    p = para(text, bold=True, size=12, align=AL.LEFT, space_after=4, keep=True); p.paragraph_format.space_before = Pt(6); return p

def bullet(text, boldlead=None):
    parts = ([(boldlead, True)] if boldlead else []) + [(text, False)]
    return rich([('•  ', False)] + parts, indent=0.3, hanging=0.2)

def step(n, text, boldlead=None):
    parts = [(f'Step {n}: ', True)] + ([(boldlead, True)] if boldlead else []) + [(text, False)]
    return rich(parts, indent=0.0)

def set_borders(tbl, sz=6):
    tblPr = tbl._tbl.tblPr
    for old in tblPr.findall(qn('w:tblBorders')): tblPr.remove(old)
    b = OxmlElement('w:tblBorders')
    for e in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        x = OxmlElement('w:' + e)
        for k, v in (('val', 'single'), ('sz', str(sz)), ('space', '0'), ('color', '000000')): x.set(qn('w:' + k), v)
        b.append(x)
    anchor = next((c for c in tblPr if c.tag in [qn('w:' + n) for n in ('shd', 'tblLayout', 'tblCellMar', 'tblLook', 'tblCaption')]), None)
    if anchor is not None: anchor.addprevious(b)
    else: tblPr.append(b)

def shade(cell, color):
    tcPr = cell._tc.get_or_add_tcPr(); s = OxmlElement('w:shd')
    s.set(qn('w:val'), 'clear'); s.set(qn('w:color'), 'auto'); s.set(qn('w:fill'), color); tcPr.append(s)

def table(headers, rows, widths, size=10.5):
    t = d.add_table(rows=1, cols=len(headers)); t.autofit = False
    set_borders(t)
    def fill(cell, text, bold=False, w=None, center=False):
        cell.width = Inches(w); p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(2); p.paragraph_format.space_before = Pt(2)
        p.alignment = AL.CENTER if center else AL.LEFT
        r = p.add_run(text); r.bold = bold; r.font.size = Pt(size); r.font.name = FONT
    for i, h in enumerate(headers):
        fill(t.rows[0].cells[i], h, True, widths[i], True); shade(t.rows[0].cells[i], 'D9E2D5')
    trPr = t.rows[0]._tr.get_or_add_trPr(); th = OxmlElement('w:tblHeader'); trPr.append(th)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row): fill(cells[i], v, i == 0, widths[i])
    for i, w in enumerate(widths): t.columns[i].width = Inches(w)
    for r in t.rows:
        trPr = r._tr.get_or_add_trPr(); cs = OxmlElement('w:cantSplit'); trPr.append(cs)
    para('', space_after=4)
    return t

def code_block(lines, size=8):
    t = d.add_table(rows=1, cols=1); t.autofit = False
    set_borders(t, 8)
    cell = t.rows[0].cells[0]; cell.width = Inches(6.4); shade(cell, 'F5F5F5')
    first = True
    for ln in lines:
        p = cell.paragraphs[0] if first else cell.add_paragraph(); first = False
        pf = p.paragraph_format; pf.space_after = Pt(0); pf.space_before = Pt(0); pf.line_spacing = 1.0
        p.alignment = AL.LEFT
        r = p.add_run(ln if ln else ' '); r.font.name = 'Courier New'; r.font.size = Pt(size)
        r._r.rPr.rFonts.set(qn('w:hAnsi'), 'Courier New'); r._r.rPr.rFonts.set(qn('w:cs'), 'Courier New')
    return t

def figure(path, title, caption, width=6.4, maxh=8.8):
    im = Image.open(path).convert('RGB'); width = min(width, maxh * im.width / im.height); im = ImageOps.expand(im, border=2, fill=(0, 0, 0))
    out = path.replace('.png', '_b.png'); im.save(out)
    p = para(title, bold=True, size=12, align=AL.LEFT, space_after=4, keep=True)
    p.paragraph_format.space_before = Pt(6)
    p = para('', align=AL.CENTER, space_after=2, keep=True)
    p.add_run().add_picture(out, width=Inches(width))
    para(caption, italic=True, size=10.5, align=AL.CENTER, space_after=6)

# ================= CONTENT =================
heading(1, 'Scenario / Problem Statement', page_break=True)
sub('1.1 Scenario')
para('LearnHub is an online course portal that offers technology courses such as Full Stack Web Development, '
     'Data Science with Python, Artificial Intelligence, Cyber Security and Cloud Computing to students across the '
     'country. Every learner must first create an account before enrolling in a course. At present the portal '
     'collects data through plain text boxes, so students often enter wrong email addresses, invalid phone numbers, '
     'impossible birth dates and weak passwords. This produces bad records, failed confirmation mails and many '
     'support requests.')
sub('1.2 Problem Statement')
para('An online course portal needs a Student Registration Form. Select suitable HTML5 input types for email, '
     'phone number, date of birth, password and the other fields, so that the browser itself helps the student to '
     'enter correct data (proper keyboards on mobile devices, date picker, masked password, built-in validation) '
     'and the portal receives clean and complete data.')
sub('1.3 Requirements')
for lead, txt in [
    ('Correct input types: ', 'each field must use the HTML5 input type that best matches its data (email, tel, date, password, url, number, range, time, month, file, radio, checkbox, select and textarea).'),
    ('Validation: ', 'use HTML5 attributes (required, pattern, min, max, minlength, maxlength) together with a small JavaScript layer for friendly messages.'),
    ('Security: ', 'the password must be hidden, at least 8 characters long, and contain upper-case, lower-case and numeric characters; a confirm-password field must match it.'),
    ('Usability: ', 'labels, placeholders, hints, a password-strength meter, grouped sections and a responsive layout for phones and desktops.'),
    ('Feedback: ', 'invalid fields are highlighted in red with a message, valid fields in green, and a success message is shown after a correct submission.')]:
    bullet(txt, lead)

heading(2, 'ICT Tool Used')
para('The following Information and Communication Technology tools were used to design, code, test and document the '
     'Student Registration Form.')
table(['S. No.', 'Tool / Technology', 'Purpose in this activity'], [
    ['1', 'HTML5', 'Structure of the form; semantic elements (form, fieldset, legend, label) and modern input types such as email, tel, date, password, url, number, range, time and month.'],
    ['2', 'CSS3', 'Styling, grid layout, colours for valid / invalid states, and media queries for a responsive mobile view.'],
    ['3', 'JavaScript (ES6)', 'Constraint Validation API, custom error messages, password-match check, strength meter and success message.'],
    ['4', 'Visual Studio Code', 'Code editor used to write and format the source code (registration.html).'],
    ['5', 'Google Chrome / Chromium', 'Web browser used to run the page and to inspect the behaviour of each input type.'],
    ['6', 'Playwright (headless Chromium)', 'Used to open the page automatically, fill the fields and capture the output screenshots shown in this report.'],
    ['7', 'Microsoft Word', 'Preparation of this Active Learning Methodology report.']],
    [0.6, 1.9, 3.9])

heading(3, 'Activity Procedure')
sub('3.1 Steps followed')
step(1, 'Read the problem statement and list all the data that a student must give while registering.', 'Analyse: ')
step(2, 'Choose the most suitable HTML5 input type for every data item (see Table 3.2).', 'Select input types: ')
step(3, 'Create registration.html with the HTML5 boilerplate and the viewport meta tag for mobile devices.', 'Create the page: ')
step(4, 'Build a form with three fieldsets - Personal Details, Account Security and Course Preferences - and add a label for every control.', 'Design the form: ')
step(5, 'Add required, minlength, maxlength, min, max and pattern attributes so that the browser validates the data.', 'Add validation attributes: ')
step(6, 'Style the page with CSS (two-column grid, red / green states, responsive media query).', 'Apply CSS: ')
step(7, 'Write JavaScript to show messages, compare the two passwords, show the password strength and display a success message.', 'Add JavaScript: ')
step(8, 'Open the page in the browser, test empty submission, wrong values and correct values, and capture the screenshots.', 'Test and record output: ')
sub('3.2 Selection of HTML5 input types')
table(['Data', 'Input type', 'Key attributes', 'Reason for selection'], [
    ['Full name', 'type="text"', 'required, minlength, pattern', 'Free text; pattern allows only letters, space and dot.'],
    ['Email', 'type="email"', 'required', 'Browser checks the name@domain format; mobile shows the @ keyboard.'],
    ['Phone number', 'type="tel"', 'pattern="[6-9][0-9]{9}", maxlength', 'Numeric keypad on phones; pattern enforces a 10-digit mobile number.'],
    ['Date of birth', 'type="date"', 'required, min, max (today)', 'Built-in date picker; max blocks future dates.'],
    ['Password', 'type="password"', 'minlength=8, pattern', 'Characters are masked; pattern demands upper, lower and digit.'],
    ['Confirm password', 'type="password"', 'required', 'Must match the password (checked in JavaScript).'],
    ['Gender', 'type="radio"', 'required, same name', 'Only one option can be chosen.'],
    ['Course', 'input + datalist', 'list, required', 'Type-ahead suggestions of available courses.'],
    ['Qualification', '&lt;select&gt;'.replace('&lt;', '<').replace('&gt;', '>'), 'required', 'Fixed list of choices.'],
    ['Start month', 'type="month"', '-', 'Month and year picker.'],
    ['Class time', 'type="time"', 'value="18:00"', 'Time picker.'],
    ['Experience', 'type="number"', 'min=0, max=40, step=1', 'Numeric spinner with limits.'],
    ['Study hours', 'type="range"', 'min=1, max=40', 'Slider for an approximate value.'],
    ['Profile link', 'type="url"', '-', 'Browser checks for a valid URL.'],
    ['Photo', 'type="file"', 'accept="image/png, image/jpeg"', 'Restricts the upload to image files.'],
    ['Address', '<textarea>', 'rows, maxlength', 'Multi-line text.'],
    ['Terms / Updates', 'type="checkbox"', 'required (terms)', 'Yes / no choice; terms must be accepted.']],
    [1.15, 1.3, 1.85, 2.1], size=10)

heading(4, 'Student Activity / Implementation')
para('The complete source code of the Student Registration Form is given below. The HTML part defines the form and '
     'the input types, the CSS part gives the design, and the JavaScript part performs validation and displays '
     'the feedback. The file is saved as registration.html and can be opened in any modern web browser.')
sub('Program: registration.html')
lines = open('registration.html').read().split('\n')
code_block(lines, 8)
para('', space_after=4)
sub('Explanation of the code')
for lead, txt in [
    ('Form and fieldsets: ', 'the form uses novalidate so that our JavaScript can show custom messages, while the HTML5 constraint attributes still define the rules.'),
    ('Email and phone: ', 'type="email" validates the address format; type="tel" with pattern="[6-9][0-9]{9}" accepts only 10-digit mobile numbers.'),
    ('Date of birth: ', 'type="date" gives a calendar picker; JavaScript sets max to today and rejects students younger than 16 years.'),
    ('Password: ', 'type="password" hides the text; the lookahead pattern (?=.*[a-z])(?=.*[A-Z])(?=.*[0-9]).{8,} enforces the strength rules and the meter shows the score.'),
    ('Validation flow: ', 'validateField() uses checkValidity() and setCustomValidity() to decide whether a field is valid and then shows the message and the red / green border.')]:
    bullet(txt, lead)

sub('Explanation of the CSS and JavaScript')
for lead, txt in [
    ('CSS grid: ', 'the .grid class creates two equal columns; the media query (max-width: 640px) changes it to one column so the form fits phone screens.'),
    ('Visual feedback: ', 'the .invalid class gives a red border and pink background, and the .valid class gives a green border, so the student sees the result of every field immediately.'),
    ('Password meter: ', 'the input event on the password field scores length, mixed case, digits and symbols and changes the width and colour of the bar.'),
    ('Submit handler: ', 'preventDefault() stops the page reload, every required field is validated, and the success banner is shown only when all fields are valid.')]:
    bullet(txt, lead)
sub('Advantages of using suitable HTML5 input types')
for txt in ['Mobile devices open the correct keyboard automatically (email keyboard for email, numeric keypad for phone).',
            'Date, month and time pickers avoid typing mistakes and wrong formats.',
            'Built-in validation reduces the JavaScript code and the load on the server.',
            'Passwords are masked, and pattern rules enforce a strong password.',
            'Clean and correct data is stored, so confirmation mails and phone contacts do not fail.']:
    bullet(txt)
sub('HTML5 validation attributes used in the form')
table(['Attribute', 'Purpose', 'Used on'], [
    ['required', 'Field cannot be left empty', 'Name, date of birth, email, phone, password, course, terms'],
    ['minlength / maxlength', 'Limits the number of characters', 'Name, password, phone, address'],
    ['pattern', 'Regular expression the value must match', 'Name, phone, password'],
    ['min / max', 'Lowest and highest allowed value or date', 'Date of birth, experience, study hours'],
    ['step', 'Step size of a number input', 'Years of experience'],
    ['accept', 'Allowed file types', 'Profile photo (PNG, JPEG)'],
    ['list (datalist)', 'Type-ahead suggestions', 'Course'],
    ['placeholder / autocomplete', 'Sample text and browser auto-fill hints', 'Name, email, phone, passwords']],
    [1.7, 2.3, 2.4], size=10.5)
sub('Sample test data used for the outputs')
table(['Field', 'Valid data entered (Output 4 and 5)', 'Invalid data entered (Output 3)'], [
    ['Full name', 'Aishwarya U', 'Ajay Iyanraj (valid)'],
    ['Date of birth', '14-08-2005', '10-05-2012 (under 16)'],
    ['Email', 'aishwarya.u@example.com', 'ajay.iyanraj@'],
    ['Phone', '9876543210', '12345'],
    ['Password', 'Learn@2026', 'abc123'],
    ['Confirm password', 'Learn@2026', 'abc124 (mismatch)'],
    ['Course', 'Full Stack Web Development', '(not selected)'],
    ['Qualification', 'Undergraduate', '(not selected)'],
    ['Profile link', 'https://www.linkedin.com/in/aishwarya-u', 'linkedin.com/in/ajay'],
    ['Terms and Conditions', 'Accepted', 'Not accepted']],
    [1.6, 2.6, 2.2], size=10.5)
sub('Behaviour of the main input types in the browser')
table(['Input type', 'On a desktop browser', 'On a mobile browser'], [
    ['type="email"', 'Rejects values without @ and a domain', 'Shows a keyboard with @ and .com keys'],
    ['type="tel"', 'Accepts text; pattern checks 10 digits', 'Shows the numeric dial pad'],
    ['type="date"', 'Calendar drop-down picker', 'Native date wheel / calendar'],
    ['type="password"', 'Characters shown as dots', 'Characters masked; no auto-capitalisation'],
    ['type="number"', 'Up / down spinner within min and max', 'Numeric keyboard'],
    ['type="url"', 'Rejects values without a scheme', 'Keyboard with / and .com keys'],
    ['type="range"', 'Draggable slider', 'Touch slider']],
    [1.5, 2.5, 2.4], size=10.5)
sub('Learning outcomes')
for txt in ['Identified the HTML5 input type that best fits each kind of data in a registration form.',
            'Applied constraint-validation attributes (required, pattern, min, max, minlength, maxlength).',
            'Used JavaScript and CSS to give instant visual feedback and a responsive design.']:
    bullet(txt)
sub('Conclusion')
para('Choosing the right HTML5 input type is the first and simplest level of validation. With type="email", '
     'type="tel", type="date" and type="password", the browser itself gives the student the correct keyboard, '
     'picker or masking and checks the format before the data reaches the server. Combined with attributes such as '
     'required, pattern, min and max, and a little JavaScript for friendly messages, the LearnHub registration '
     'form collects accurate data with less effort for both the student and the portal.')
para('The next section shows the form running in the browser in its different states, starting with the blank form and '
     'ending with the mobile view and the test results.')
heading(5, 'Expected Output', page_break=True)
para('The screenshots below were captured from the running registration.html page in a Chromium browser. They show '
     'the form in its different states.')
figs = [
 ('shots/1_empty.png', 'Output 1: Blank Student Registration Form', 'Fig. 1 - Initial form with all HTML5 input controls (text, date, email, tel, radio, password, list, select, month, time, number, range, url, file, textarea and checkbox).'),
 ('shots/2_errors.png', 'Output 2: Validation messages on empty submission', 'Fig. 2 - Clicking Register Now with no data highlights every required field in red and shows a message below it.'),
 ('shots/3_invalid.png', 'Output 3: Invalid email, phone, date of birth, password and URL', 'Fig. 3 - Wrong email format, 5-digit phone, age below 16, weak password, mismatch of passwords and an invalid URL are rejected.'),
 ('shots/4_filled.png', 'Output 4: Correctly filled form', 'Fig. 4 - All fields contain valid data; borders turn green and the password-strength meter is full.'),
 ('shots/5_success.png', 'Output 5: Successful registration', 'Fig. 5 - After a valid submission a green confirmation message with the student name and email is displayed.'),
]
for j,(f, t, c) in enumerate(figs): figure(f, t, c, maxh=(7.9 if j==0 else 8.8))
p = para('Output 6: Responsive view on a mobile screen (390 px width)', bold=True, align=AL.LEFT, space_after=4, keep=True)
# mobile shot is very tall: split into three side-by-side strips
m = Image.open('shots/6_mobile.png').convert('RGB'); W, H = m.size; k = 3; h = H // k + 1
strips = []
for i in range(k):
    s = m.crop((0, i * h, W, min(H, (i + 1) * h))); s = ImageOps.expand(s, border=3, fill=(0, 0, 0)); strips.append(s)
canvas = Image.new('RGB', (sum(s.width for s in strips) + 40 * (k - 1), max(s.height for s in strips)), 'white')
x = 0
for s in strips: canvas.paste(s, (x, 0)); x += s.width + 40
canvas.save('shots/6_mobile_strips.png')
p = para('', align=AL.CENTER, space_after=2, keep=True); p.add_run().add_picture('shots/6_mobile_strips.png', width=Inches(6.2))
para('Fig. 6 - On a phone the two-column layout collapses into a single column (shown in three parts, top to bottom, left to right).', italic=True, size=10.5, align=AL.CENTER, space_after=8)
sub('Test cases')
table(['S. No.', 'Test input', 'Expected behaviour', 'Result'], [
    ['1', 'Submit the form with all fields empty', 'Every required field turns red with a message', 'Pass'],
    ['2', 'Email = ajay.iyanraj@', 'Message: valid email address required', 'Pass'],
    ['3', 'Phone = 12345', 'Message: 10-digit mobile number starting with 6-9', 'Pass'],
    ['4', 'Date of birth = 10-05-2012', 'Message: must be at least 16 years old', 'Pass'],
    ['5', 'Password = abc123', 'Weak strength bar; pattern message shown', 'Pass'],
    ['6', 'Confirm password different from password', 'Message: passwords do not match', 'Pass'],
    ['7', 'Profile link = linkedin.com/in/ajay', 'Message: URL must start with http:// or https://', 'Pass'],
    ['8', 'All fields valid and terms accepted', 'Green borders and success message', 'Pass']],
    [0.6, 2.4, 2.8, 0.7], size=10)
sub('Result')
para('The Student Registration Form for the online course portal was designed with suitable HTML5 input types - '
     'email for the email address, tel for the phone number, date for the date of birth and password for the '
     'password - together with number, range, time, month, url, file, radio, checkbox, select and textarea for '
     'the remaining data. The built-in validation attributes and the JavaScript feedback prevented invalid data and '
     'produced a clean, user-friendly and mobile-responsive registration page.')

# ---------- footer page numbers, margins, page border ----------
sec = d.sections[0]
sec.bottom_margin = Emu(int(0.6 * 914400)); sec.footer_distance = Emu(int(0.25 * 914400))
fp = sec.footer.paragraphs[0]; fp.alignment = AL.CENTER
def fld(p, instr):
    r = p.add_run(); r.font.size = Pt(10); r.font.name = FONT
    a = OxmlElement('w:fldChar'); a.set(qn('w:fldCharType'), 'begin'); r._r.append(a)
    it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve'); it.text = instr; r._r.append(it)
    c = OxmlElement('w:fldChar'); c.set(qn('w:fldCharType'), 'separate'); r._r.append(c)
    t = OxmlElement('w:t'); t.text = '1'; r._r.append(t)
    e = OxmlElement('w:fldChar'); e.set(qn('w:fldCharType'), 'end'); r._r.append(e)
fld(fp, 'PAGE')

sp = sect
pb = OxmlElement('w:pgBorders'); pb.set(qn('w:offsetFrom'), 'page')
for side in ('top', 'left', 'bottom', 'right'):
    e = OxmlElement('w:' + side)
    for k_, v in (('val', 'single'), ('sz', '12'), ('space', '20'), ('color', '000000')): e.set(qn('w:' + k_), v)
    pb.append(e)
sp.find(qn('w:pgMar')).addnext(pb)
# section break after page 2: header only on cover + index pages
sig2._p.get_or_add_pPr().append(copy.deepcopy(sp))
last = d.sections[-1]
for hr in last._sectPr.findall(qn('w:headerReference')): last._sectPr.remove(hr)
last.header.is_linked_to_previous = False
for p_ in last.header.paragraphs:
    for r_ in list(p_.runs): r_._r.getparent().remove(r_._r)
last.top_margin = Inches(0.8)
d.core_properties.title = 'ALM Report - Student Registration Form'
d.save('ALM_Student_Registration_Form.docx')
