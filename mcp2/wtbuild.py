import sys,os,json,glob; sys.path.insert(0,'.')
from wtlib import *
from docx.shared import Pt,Inches,Emu,Twips
from docx.enum.text import WD_ALIGN_PARAGRAPH as AL, WD_TAB_ALIGNMENT as TA
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image
OUT=sys.argv[1] if len(sys.argv)>1 else 'merged.docx'
REPL={}   # 'ref:imageNN.png' -> path
if os.path.exists('repl.json'): REPL=json.load(open('repl.json'))
SPC=json.load(open('spacers.json')) if os.path.exists('spacers.json') else {}
SHR=json.load(open('shrink.json')) if os.path.exists('shrink.json') else {}
FONT='Times New Roman'
SUBS=[('AMIRTHA VARSHINI S','AJAY.I'),('Amirtha Varshini S','Ajay I'),('Amirthavarshini S','Ajay I'),('S. Amirtha Varshini','Ajay I'),('Amirtha Varshini','Ajay I'),('Amirthavarshini','Ajay'),("Amirtha's","Ajay's"),('amirthavarshini693@gmail.com','ajay.i2007@gmail.com'),('amirthavarshini@example.com','ajay.i2007@example.com'),('amirthavarshini.2401015@srec.ac.in','ajay.2401007@srec.ac.in'),('71812401015','71812401007'),('7305687512','9025467813'),('+91 9025467813','+91 9025467813'),('9.03','8.74'),('varshini693@gmail.com','.i2007@gmail.com'),('amirtha','ajay'),('AMIRTHA','AJAY'),('Amirtha','Ajay')]
MONO=('consolas','courier','cascadia','lucida console','monaco','menlo')
FRONT=os.environ.get('FRONT')
if FRONT:
    from docx.enum.section import WD_SECTION
    dest=docx.Document(FRONT); dbody=dest.element.body
    nsec=dest.add_section(WD_SECTION.NEW_PAGE); sect=nsec._sectPr
    pn=OxmlElement('w:pgNumType'); pn.set(qn('w:start'),'4')
else:
    dest=docx.Document('ref.docx'); dbody=dest.element.body
    for e in list(dbody):
        if e.tag!=w('sectPr'): dbody.remove(e)
    sect=dbody.find(w('sectPr'))
# ---------- page setup ----------
for ch in list(sect):
    if ch.tag in (w('pgMar'),): sect.remove(ch)
pm=OxmlElement('w:pgMar')
for k,v in dict(top=1050,right=850,bottom=int(os.environ.get('BOT_M','1250')),left=850,header=int(os.environ.get('HDR_D','170')),footer=int(os.environ.get('FTR_D','620')),gutter=0).items(): pm.set(qn('w:'+k),str(v))
sect.find(w('pgSz')).addnext(pm)
for pb in sect.findall(w('pgBorders')): sect.remove(pb)
pb=OxmlElement('w:pgBorders'); pb.set(qn('w:offsetFrom'),'page')
for side,val in (('top','single'),('left','single'),('bottom','single'),('right','single')):
    e=OxmlElement('w:'+side)
    for k,v in dict(val=val,sz=12,space=28,color='000000').items(): e.set(qn('w:'+k),str(v))
    pb.append(e)
pm.addnext(pb)
if FRONT:
    for x in sect.findall(w('pgNumType')): sect.remove(x)
    pb.addnext(pn)
for dg in sect.findall(w('docGrid')): sect.remove(dg)
# header / footer
def runx(p,text,size=12,bold=False):
    size=size
    r=p.add_run(text); r.font.name=FONT; r.font.size=Pt(size); r.bold=bold; r._r.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),FONT); return r
def fld(p,instr,size=12):
    for typ in ('begin',None,'separate','t','end'):
        r=p.add_run(); r.font.name=FONT; r.font.size=Pt(size)
        if typ in('begin','separate','end'):
            fc=OxmlElement('w:fldChar'); fc.set(qn('w:fldCharType'),typ); r._r.append(fc)
        elif typ is None:
            it=OxmlElement('w:instrText'); it.set(qn('xml:space'),'preserve'); it.text=' %s '%instr; r._r.append(it)
        else: r.text='1'
sec=dest.sections[-1] if FRONT else dest.sections[0]
def setup(hf,items):
    for p in hf.paragraphs[1:]: p._p.getparent().remove(p._p)
    p=hf.paragraphs[0]
    for r in list(p._p): 
        if r.tag!=w('pPr'): p._p.remove(r)
    ppr=p._p.get_or_add_pPr()
    for t in ppr.findall(w('tabs'))+ppr.findall(w('ind'))+ppr.findall(w('jc')): ppr.remove(t)
    p.alignment=AL.LEFT; pf=p.paragraph_format; pf.left_indent=Emu(0); pf.space_before=Pt(0); pf.space_after=Pt(0)
    for pos,al in items[0]: pf.tab_stops.add_tab_stop(pos,al)
    tb=ppr.find(w('tabs'))
    for pos in (4513,9026):
        e=OxmlElement('w:tab'); e.set(qn('w:val'),'clear'); e.set(qn('w:pos'),str(pos)); tb.insert(0,e)
    items[1](p)
    if items[2]>5000: pf.space_before=Pt(float(os.environ.get('FTR_SP','0')))
    fp=OxmlElement('w:framePr')
    for k,v in dict(w=10210,hRule='auto',wrap='around',vAnchor='page',hAnchor='page',x=850,y=items[2]).items(): fp.set(qn('w:'+k),str(v))
sec.header.is_linked_to_previous=False; sec.footer.is_linked_to_previous=False
tw=Twips(10210)
setup(sec.header,([(tw,TA.RIGHT)],lambda p:(runx(p,os.environ.get('HDR_L','20CS279 – WEB TECHNOLOGIES LABORATORY'),float(os.environ.get('HDR_SZ','10.5'))),runx(p,'\t'+os.environ.get('HDR_R','COURSE INSTRUCTOR: Mr. N. Manoj, AP/CSE'),float(os.environ.get('HDR_SZ','10.5')))),230))
setup(sec.footer,([(Twips(5105),TA.CENTER),(tw,TA.RIGHT)],lambda p:(runx(p,'AJAY.I',10.5),runx(p,'\t',10.5),fld(p,'PAGE',10.5),runx(p,'\t71812401007',10.5)),16370))
# ---------- helpers ----------
def para(text='',bold=False,size=12,align=AL.JUSTIFY,before=0,after=4,left=0,first=0,keep=False,ls=1.12):
    p=dest.add_paragraph(); pf=p.paragraph_format; p.alignment=align
    pf.space_before=Pt(before); pf.space_after=Pt(after); pf.line_spacing=ls
    if left: pf.left_indent=Inches(left)
    if first: pf.first_line_indent=Inches(first)
    if keep: pf.keep_with_next=True
    if text: runx(p,text,size,bold)
    return p
def set_cell_border(cell):
    tcPr=cell._tc.get_or_add_tcPr(); b=OxmlElement('w:tcBorders')
    for e in('top','left','bottom','right'):
        x=OxmlElement('w:'+e); x.set(qn('w:val'),'single'); x.set(qn('w:sz'),'8'); x.set(qn('w:space'),'0'); x.set(qn('w:color'),'000000'); b.append(x)
    tcPr.append(b)
def title_table(no,date,title):
    t=dest.add_table(rows=2,cols=2); t.autofit=False; t.alignment=docx.enum.table.WD_TABLE_ALIGNMENT.CENTER
    widths=(1.55,5.45)
    for r in t.rows:
        r.height=Inches(0.42)
        for i,c in enumerate(r.cells): c.width=Inches(widths[i]); set_cell_border(c)
    for i,wd in enumerate(widths): t.columns[i].width=Inches(wd)
    def put(c,text,align):
        p=c.paragraphs[0]; p.alignment=align; p.paragraph_format.space_before=Pt(4); p.paragraph_format.space_after=Pt(4)
        runx(p,text,12,True)
        tcPr=c._tc.get_or_add_tcPr(); v=OxmlElement('w:vAlign'); v.set(qn('w:val'),'center'); tcPr.append(v)
    put(t.cell(0,0),os.environ.get('LABEL','EX.NO.')+f'{no}',AL.LEFT); put(t.cell(1,0),f'DATE: {date}' if date else 'DATE:',AL.LEFT)
    m=t.cell(0,1).merge(t.cell(1,1)); put(m,title,AL.CENTER)
    for r in t.rows: r._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
    return t
# ---------- conversion ----------
KEEP_R=['rFonts','b','bCs','i','iCs','u','strike','color','sz','szCs','highlight','vertAlign','shd','caps','smallCaps']
def mkel(tag,**kw):
    e=OxmlElement('w:'+tag)
    for k,v in kw.items(): e.set(qn('w:'+k),str(v))
    return e
def build_rpr(src,pstyle,run):
    rp=run.find(w('rPr')); rs=rp.find(w('rStyle')) if rp is not None else None
    eff=src.eff_r(pstyle,rs.get(w('val')) if rs is not None else None,rp)
    new=OxmlElement('w:rPr'); fam=FONT
    rf=eff.get(w('rFonts'))
    if rf is not None:
        a=rf.get(w('ascii')) or ''
        if any(m in a.lower() for m in MONO): fam=a
    new.append(mkel('rFonts',ascii=fam,hAnsi=fam,cs=fam,eastAsia=fam))
    for t in ('b','bCs','i','iCs','strike','caps','smallCaps'):
        e=eff.get(w(t))
        if e is not None and e.get(w('val')) not in ('0','false'): new.append(copy.deepcopy(e))
    u=eff.get(w('u'))
    if u is not None and u.get(w('val'))!='none': new.append(copy.deepcopy(u))
    c=eff.get(w('color'))
    if c is not None and c.get(w('val')) not in ('auto',None): new.append(copy.deepcopy(c))
    sz=eff.get(w('sz'))
    new.append(mkel('sz',val=sz.get(w('val')) if sz is not None else 24)); new.append(mkel('szCs',val=sz.get(w('val')) if sz is not None else 24))
    for t in ('highlight','vertAlign'):
        e=eff.get(w(t))
        if e is not None: new.append(copy.deepcopy(e))
    sh=eff.get(w('shd'))
    if sh is not None and sh.get(w('fill')) not in (None,'auto'): new.append(copy.deepcopy(sh))
    # schema order fix
    order=['rFonts','b','bCs','i','iCs','caps','smallCaps','strike','color','sz','szCs','highlight','u','shd','vertAlign']
    kids=sorted(list(new),key=lambda e:order.index(e.tag.split('}')[1]) if e.tag.split('}')[1] in order else 99)
    for k in list(new): new.remove(k)
    for k in kids: new.append(k)
    return new
did=[100]
def convert_drawing(src,dr,unit):
    dr=copy.deepcopy(dr)
    node=dr.find('{%s}anchor'%WP)
    if node is None: node=dr.find('{%s}inline'%WP)
    ext=node.find('{%s}extent'%WP); cx,cy=int(ext.get('cx')),int(ext.get('cy'))
    blip=node.find('.//{%s}blip'%A)
    if blip is None or blip.get('{%s}embed'%R) is None: return None
    rid=blip.get('{%s}embed'%R)
    part=src.rels[rid].target_part; blob=part.blob; fname=os.path.basename(str(part.partname))
    k=f'{src.tag}:{fname}'
    srect=node.find('.//{%s}srcRect'%A)
    if k in REPL:
        blob=open(REPL[k],'rb').read(); im=Image.open(io.BytesIO(blob)); cy=int(cx*im.height/im.width)
        if srect is not None: srect.getparent().remove(srect)
    sc=SHR.get(unit,1.0)
    cx=int(cx*sc); cy=int(cy*sc)
    mh=int(float(os.environ.get('IMG_MAXH','9'))*914400)
    if cy>mh: cx=int(cx*mh/cy); cy=mh
    mx=int(6.9*914400)
    if cx>mx: cy=int(cy*mx/cx); cx=mx
    nrid=dest.part.get_or_add_image(io.BytesIO(blob))[0]; blip.set('{%s}embed'%R,nrid)
    graphic=node.find('{%s}graphic'%A)
    for e in graphic.iter('{http://schemas.openxmlformats.org/drawingml/2006/picture}ext') if False else []: pass
    # fix inner pic extent
    for xf in graphic.iter('{%s}xfrm'%A):
        xe=xf.find('{%s}ext'%A)
        if xe is not None: xe.set('cx',str(cx)); xe.set('cy',str(cy))
    did[0]+=1
    inl=etree.SubElement(etree.Element('dummy'),'{%s}inline'%WP,distT='0',distB='0',distL='0',distR='0')
    etree.SubElement(inl,'{%s}extent'%WP,cx=str(cx),cy=str(cy))
    etree.SubElement(inl,'{%s}effectExtent'%WP,l='0',t='0',r='0',b='0')
    etree.SubElement(inl,'{%s}docPr'%WP,id=str(did[0]),name='Picture %d'%did[0])
    etree.SubElement(inl,'{%s}cNvGraphicFramePr'%WP)
    inl.append(copy.deepcopy(graphic))
    nd=OxmlElement('w:drawing'); nd.append(inl)
    return nd
def convert_par(src,p,unit):
    ppr=p.find(w('pPr')); ps=ppr.find(w('pStyle')).get(w('val')) if ppr is not None and ppr.find(w('pStyle')) is not None else src.default_pstyle
    eff=src.eff_p(ps,ppr)
    np_=OxmlElement('w:p'); npr=OxmlElement('w:pPr'); np_.append(npr)
    marker=None; mind=None
    npe=eff.get(w('numPr'))
    if npe is not None and npe.find(w('numId')) is not None and npe.find(w('numId')).get(w('val'))!='0':
        il=npe.find(w('ilvl')); il=int(il.get(w('val'))) if il is not None else 0
        marker,mind=src.num_marker(npe.find(w('numId')).get(w('val')),il)
    for t in ('keepNext','keepLines'):
        e=eff.get(w(t))
        if e is not None and e.get(w('val')) not in('0','false'): npr.append(mkel(t))
    tabs=eff.get(w('tabs'))
    if tabs is not None: npr.append(copy.deepcopy(tabs))
    sp=eff.get(w('spacing'))
    if sp is not None:
        n=mkel('spacing')
        for k in ('before','after','line','lineRule'):
            if sp.get(w(k)) is not None: n.set(w(k),sp.get(w(k)))
        if n.get(w('before')) is not None: n.set(w('before'),str(min(int(n.get(w('before'))),120)))
        if n.get(w('after')) is not None: n.set(w('after'),str(min(int(n.get(w('after'))),80)))
        if n.get(w('lineRule')) in (None,'auto') and n.get(w('line')) is not None and int(n.get(w('line')))>276: n.set(w('line'),'264')
        npr.append(n)
    ind=eff.get(w('ind'))
    if marker is not None and mind is not None:
        ind=mind
    if ind is not None:
        n=mkel('ind')
        for k in ('left','hanging','firstLine','start','end'):
            if ind.get(w(k)) is not None: n.set(w(k),ind.get(w(k)))
        if n.get(w('left')) is not None and int(n.get(w('left')))>1500: n.set(w('left'),'1500')
        if n.get(w('start')) is not None and int(n.get(w('start')))>1500: n.set(w('start'),'1500')
        npr.append(n)
    jc=eff.get(w('jc'))
    if jc is not None: npr.append(copy.deepcopy(jc))
    runs=[]
    def collect(parent):
        for ch in parent:
            if ch.tag==w('r'): runs.append(ch)
            elif ch.tag in (w('hyperlink'),w('ins'),w('smartTag'),w('sdt')): 
                collect(ch.find(w('sdtContent')) if ch.tag==w('sdt') else ch)
            elif ch.tag=='{http://schemas.openxmlformats.org/markup-compatibility/2006}AlternateContent': pass
    collect(p)
    has_img=False
    if marker is not None:
        r=OxmlElement('w:r'); r.append(build_rpr(src,ps,OxmlElement('w:r'))); t=OxmlElement('w:t'); t.text=marker; t.set(qn('xml:space'),'preserve'); r.append(t)
        tb=OxmlElement('w:r'); tb.append(OxmlElement('w:tab')); np_.append(r); np_.append(tb)
    for r in runs:
        nr=OxmlElement('w:r'); nr.append(build_rpr(src,ps,r))
        for ch in r:
            tg=ch.tag
            if tg==w('t'):
                t=OxmlElement('w:t'); t.text=ch.text; t.set(qn('xml:space'),'preserve'); nr.append(t)
            elif tg==w('tab'): nr.append(OxmlElement('w:tab'))
            elif tg==w('br'):
                if ch.get(w('type')) not in ('page','column'): nr.append(OxmlElement('w:br'))
            elif tg==w('cr'): nr.append(OxmlElement('w:br'))
            elif tg==w('noBreakHyphen'): nr.append(OxmlElement('w:noBreakHyphen'))
            elif tg==w('drawing'):
                dd_=convert_drawing(src,ch,unit)
                if dd_ is not None: nr.append(dd_); has_img=True
            elif tg=='{http://schemas.openxmlformats.org/markup-compatibility/2006}AlternateContent':
                ch2=ch.find('.//'+w('drawing'))
                if ch2 is not None:
                    dd_=convert_drawing(src,ch2,unit)
                    if dd_ is not None: nr.append(dd_); has_img=True
        if len(nr)>1: np_.append(nr)
    if has_img:
        for t in npr.findall(w('jc')): npr.remove(t)
        npr.append(mkel('jc',val='center'))
        for t in npr.findall(w('ind')): npr.remove(t)
    return np_,has_img
def dosubs(p):
    dosubs0(p)
    ts=[t for t in p.iter(w('t'))]
    j=''.join(t.text or '' for t in ts)
    if re.search(r'(?i)varshini|amirth',j):
        n=j
        for a,b in [('ajayvarshini693@gmail.com','ajay.i2007@gmail.com'),('Ajayvarshini','Ajay'),('ajayvarshini','ajay'),('Varshini','I'),('varshini','')]: n=n.replace(a,b)
        if n!=j and ts:
            ts[0].text=n
            for t in ts[1:]: t.text=''
def dosubs0(p):
    for t in p.iter(w('t')):
        if t.text:
            x=t.text
            for a,b in SUBS: x=x.replace(a,b)
            t.text=x
def unit_elems(src,start,end):
    out=[]
    for e in list(src.body)[start:end]:
        if e.tag==w('p'): out.append(e)
        elif e.tag==w('tbl'): out.append(e)
    return out
def parse_title(tbl):
    rows=tbl.findall(w('tr')); c0=[' '.join(ptext(p_).strip() for p_ in tc.findall(w('p'))).strip() for tc in rows[0].findall(w('tc'))]; c1=[ptext(tc).strip() for tc in rows[1].findall(w('tc'))]
    no=c0[0].replace('EX.NO.','').strip(); title=re.sub(r'\s+',' ',c0[1]).strip(' \xa0'); date=c1[0].replace('DATE:','').strip()
    return no,date,title
def units_of(src):
    ch=list(src.body); idx=[i for i,e in enumerate(ch) if e.tag==w('tbl')]
    res=[]
    for k,i in enumerate(idx):
        end=idx[k+1] if k+1<len(idx) else len(ch)
        res.append((parse_title(ch[i]),i+1,end))
    return res
FIX={'JAVASCRIPT DOM ( HOVER AND CLICKS ON A TABL)':'JAVASCRIPT DOM ( HOVER AND CLICKS ON A TABLE )','REACT JS: COMPONENTS, PROPS STATE AND EVENTS':'REACT JS: COMPONENTS, PROPS, STATE AND EVENTS','REACT JS: COMPONENT STATE AND INPUT HANDLING':'REACT JS: COMPONENT STATE AND INPUT HANDLING'}
def emit(src,no,date,title,s,e,unitkey):
    title=FIX.get(title,title); title=re.sub(r'\( ',r'(',title); title=re.sub(r' \)',r')',title)
    if not date and no=='2': date='06.07.26'
    pbp=para(after=0,ls=1.0); pbp.paragraph_format.page_break_before=True
    pPr=pbp._p.get_or_add_pPr(); rp=OxmlElement('w:rPr'); rp.append(mkel('sz',val=2)); pPr.append(rp)
    title_table(no,date,title); para(after=2,ls=1.0,size=4)
    elems=[x for x in unit_elems(src,s,e)]
    conv=[]
    for x in elems:
        if x.tag!=w('p'): continue
        if x.find('.//'+w('sectPr')) is not None and not ptext(x).strip() and x.find('.//'+w('drawing')) is None: continue
        np_,hi=convert_par(src,x,unitkey); dosubs(np_)
        txt=ptext(np_).strip(); conv.append((np_,txt,hi))
    CODE=re.compile(r'^(<|\}|\{|//|/\*|\*|@|[\w.#:\[\]>*,\s-]+\{\s*$|[\w-]+\s*:\s*[^:]+;?\s*$|(let|const|var|function|if|else|for|while|console|app\.|import|export|return|class|try|catch|async|await|\)|\]|\w+\(|\w+\.\w+)\b)')
    def codey(t): return bool(CODE.match(t.strip())) if t.strip() else False
    # drop empties between code-like lines
    keep=[]
    for i,(np_,txt,hi) in enumerate(conv):
        if not txt and not hi and 0<i<len(conv)-1:
            j=i-1
            while j>=0 and not conv[j][1] and not conv[j][2]: j-=1
            k=i+1
            while k<len(conv) and not conv[k][1] and not conv[k][2]: k+=1
            if j>=0 and k<len(conv) and codey(conv[j][1]) and codey(conv[k][1]): continue
        keep.append((np_,txt,hi))
    conv=keep
    # collapse empties
    clean=[]; 
    for np_,txt,hi in conv:
        empty=(not txt) and not hi
        if empty:
            if not clean or clean[-1][1]=='' and not clean[-1][2]: continue
        clean.append((np_,txt,hi))
    while clean and not clean[-1][1] and not clean[-1][2]: clean.pop()
    for np_,txt,hi in clean:
        if not txt and not hi:
            npr=np_.find(w('pPr'))
            for t in npr.findall(w('spacing')): npr.remove(t)
            npr.append(mkel('spacing',before=0,after=0,line=240,lineRule='auto'))
            rp=OxmlElement('w:rPr'); rp.append(mkel('sz',val=12)); npr.append(rp)
    # result block
    ri=None if os.environ.get('NORES') else next((i for i,(a,t,h) in enumerate(clean) if re.match(r'(?i)^result\b',t)),None)
    for i,(np_,txt,hi) in enumerate(clean):
        npr=np_.find(w('pPr'))
        if txt and len(txt)<=45 and (txt.endswith(':') or txt.isupper()) and i+1<len(clean):
            j=i+1
            while j<len(clean) and not clean[j][1] and not clean[j][2]: j+=1
            if j<len(clean) and (clean[j][2] or True) and (clean[j][2] or j==i+1):
                for q in range(i,j):
                    qp=clean[q][0].find(w('pPr'))
                    if qp.find(w('keepNext')) is None: qp.insert(0,mkel('keepNext'))
        if ri is not None and i>=ri:
            if npr.find(w('keepNext')) is None and i<len(clean)-1: npr.insert(0,mkel('keepNext'))
            if npr.find(w('keepLines')) is None: npr.insert(0,mkel('keepLines'))
        if ri is not None and i==ri:
            sp=npr.find(w('spacing'))
            if sp is None: sp=mkel('spacing'); npr.append(sp)
            sp.set(w('before'),str(int(float(SPC.get(unitkey,0))*20)))
        dbody.insert(list(dbody).index(sect),np_)
    return len(clean),ri is not None
if __name__=='__main__':
    log=[]
    for tag,path in (('e1','e1.docx'),('e2','e2.docx'),('e3','e3.docx'),('e4','e4.docx')):
        s=Src(path); s.tag=tag
        (no,date,title),a,b=units_of(s)[0]
        n,r=emit(s,no,date,title,a,b,no); log.append((no,title,n,r))
    s=Src('ref.docx'); s.tag='ref'
    for (no,date,title),a,b in units_of(s)[4:]:
        n,r=emit(s,no,date,title,a,b,no); log.append((no,title,n,r))
    for l in log: print(l)
    ORD=['pStyle','keepNext','keepLines','pageBreakBefore','framePr','widowControl','numPr','suppressLineNumbers','pBdr','shd','tabs','suppressAutoHyphens','kinsoku','wordWrap','overflowPunct','topLinePunct','autoSpaceDE','autoSpaceDN','bidi','adjustRightInd','snapToGrid','spacing','ind','contextualSpacing','mirrorIndents','suppressOverlap','jc','textDirection','textAlignment','textboxTightWrap','outlineLvl','divId','cnfStyle','rPr','sectPr','pPrChange']
    def fixp(root):
        for pp in root.iter(w('pPr')):
            kids=list(pp); seen={}
            for k in kids:
                seen.setdefault(k.tag,k)
            kids2=sorted(seen.values(),key=lambda e:ORD.index(e.tag.split('}')[1]) if e.tag.split('}')[1] in ORD else 99)
            for k in kids: pp.remove(k)
            for k in kids2: pp.append(k)
    if FRONT:
        for sp_ in dest.element.iter(w('sectPr')):
            for pb_ in sp_.findall(w('pgBorders')):
                for sd in list(pb_): 
                    sd.set(qn('w:val'),'single'); sd.set(qn('w:sz'),'12'); sd.set(qn('w:space'),'28')
    fixp(dest.element); fixp(sec.header._element); fixp(sec.footer._element)
    dest.save(OUT); print('saved',OUT)
