import re,os,io,sys
sys.path.insert(0,'../wtlab')
from wtlib import *
from docx.shared import Pt,Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH as AL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image
import exp3
A='{http://schemas.openxmlformats.org/drawingml/2006/main}'
R='{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed'
def img_by_name(src,name):
    for r in src.doc.part.rels.values():
        if 'image' in r.reltype and os.path.basename(str(r.target_part.partname))==name: return r.target_part.blob
def T(src,j): return ' '.join(ptext(src.body[j]).split())
def steps(src,a,b): return [T(src,j) for j in range(a,b+1) if T(src,j)]
def save(blob,name):
    os.makedirs('geni',exist_ok=True); fp='geni/'+name+'.jpg'; Image.open(io.BytesIO(blob)).convert('RGB').save(fp,quality=88); return fp
def CODE(lines):
    return lines
TXF=lambda x:x
import json
SHRINK=json.load(open('shrink_mcp.json')) if os.path.exists('shrink_mcp.json') else {}
def reindent(lines):
    out=[];d=0
    for ln in lines:
        t=ln.strip()
        if t.startswith('}'): d=max(0,d-1)
        out.append('  '*d+t)
        if t.endswith('{') : d+=1
        elif t=='{': d+=1
    return out
def exp7(ref):
    ts=lambda a,b:[TXF(T(ref,j)) for j in range(a,b+1) if T(ref,j)]
    code=lambda a,b:reindent([T(ref,j) for j in range(a,b+1) if T(ref,j)])
    return [
     dict(head='(a) To Display ASCII Number in a Port',aim=TXF(T(ref,1361)),steps=ts(1366,1377),code=code(1380,1389),imgs=[('image71.png','Output'),('image72.png','Output')],side=True),
     dict(head='(b) I/O Port Programming with Delay',aim=None,steps=ts(1401,1410),code=code(1413,1431),imgs=[('image73.png','Output'),('image74.png','Output'),('image75.png','Output')],side=True),
     dict(head='(c) 8051 I/O Port Programming to Toggle Between Two or More Ports',aim=None,steps=ts(1455,1464),code=code(1467,1486),imgs=[('image76.png','Output'),('image77.png','Output'),('image78.png','Output')],side=True,result=TXF(T(ref,1517)))]
def exp_parts(src):
    P={}
    P['4']=[
     dict(head='A) Transfer the Block of Data between two Memory Locations',aim=T(src,1084),steps=steps(src,1088,1096),
          code=exp3.asm_lines(' '.join(T(src,j) for j in (1100,1101,1102,1103)).replace(' 50 ',' ').replace('DJNZ R2,1','DJNZ R2,L1')),
          imgs=[('image63.jpeg','Memory window before execution'),('image64.jpeg','Memory window after execution')],side=False,
          result=T(src,1151)),
     dict(head='B) Exchange (Swap) the Block of Data Between Two Memory Locations',aim=T(src,1156),steps=steps(src,1159,1170),
          code=exp3.asm_lines(' '.join(T(src,j) for j in (1172,1173,1174,1175)).replace('PROGRAM:','')),
          imgs=[('image65.jpeg','Memory window before execution'),('image66.jpeg','Memory window after execution')],side=False,
          result=T(src,1215))]
    sm=[x.replace('largest','smallest').replace('larger','smaller') for x in steps(src,1246,1256)]
    P['5']=[
     dict(head='(a) Largest Number in an Array',aim=T(src,1220),steps=steps(src,1222,1232),code=exp3.asm_lines(' '.join(T(src,j) for j in range(1235,1239))),imgs=[('image67.jpeg','Output')],side=True),
     dict(head='(b) Smallest Number in an Array',aim='To write and execute an 8051 Assembly Language program to find the smallest number in an array stored in the specified memory location.',steps=sm,code=exp3.asm_lines(' '.join(T(src,j) for j in range(1259,1263))),imgs=[('image68.jpeg','Output'),('image69.jpeg','Memory window'),('image70.png','Memory window')],side=True,result=T(src,1308))]
    P['6']=[
     dict(head='(a) 8051 Timer Mode 1 with Delay',aim=T(src,1312),steps=steps(src,1315,1323),
          code=['#include <reg51.h>','sbit test = P1^0;','void timer_delay()','{','  TH0 = 0xFC;','  TL0 = 0x74;','  TR0 = 1;','  while(TF0 == 0);','  TR0 = 0;','  TF0 = 0;','}','void main()','{','  TMOD = 0x01;','  while(1)','  {','    test = ~test;','    timer_delay();','  }','}'],
          imgs=[('image71.jpeg','Output')],side=True),
     dict(head='(b) Timer-Based Interrupt to Toggle Output Pin',aim=T(src,1350),steps=steps(src,1353,1363),
          code=['#include <reg51.h>','sbit test = P1^0;','void Timer_init(void)','{','  TMOD = 0x01;','  TH0 = 0x4C;','  TL0 = 0x00;','  TR0 = 1;','}','void Timer0_ISR(void) interrupt 1','{','  test = ~test;','  TH0 = 0x4C;','  TL0 = 0x00;','}','int main(void)','{','  EA = 1;','  ET0 = 1;','  Timer_init();','  while(1);','}'],
          imgs=[('image72.jpeg','Output')],side=True),
     dict(head='(c) UART Programming Using 8051',aim=T(src,1390),steps=steps(src,1392,1400),
          code=['#include <reg51.h>','void main(void)','{','  TMOD = 0x20;','  TH1 = 0xFA;','  SCON = 0x50;','  TR1 = 1;','  while(1)','  {',"    SBUF = 'B';",'    while(TI == 0);','    TI = 0;','  }','}'],
          imgs=[('image73.jpeg','Serial window setup'),('image74.jpeg','Output')],side=True,result=T(src,1457))]
    return P
def emit_gen(B,src,no,date,title,parts,final):
    d=B.dest; para=B.para; runx=B.runx
    pbp=para(after=0,ls=1.0); pbp.paragraph_format.page_break_before=True
    ppr=pbp._p.get_or_add_pPr(); rp=OxmlElement('w:rPr'); rp.append(B.mkel('sz',val=2)); ppr.append(rp)
    B.title_table(no,date,title); para(after=2,ls=1.0,size=4)
    def hd(t,before=6): return para(t,bold=True,before=before,after=3,align=AL.LEFT,keep=True)
    for pi,p in enumerate(parts):
        hp=para(p['head'],bold=True,size=12,align=AL.LEFT,before=8,after=3,keep=True,left=0.0)
        if pi and no!='4': hp.paragraph_format.page_break_before=True
        if p['aim']: hd('AIM:',3); para(p['aim'],left=0.3,after=3)
        hd('ALGORITHM:',3)
        for i,s_ in enumerate(p['steps'],1):
            q=para(left=0.55,first=-0.3,after=1.5,ls=1.08); q.paragraph_format.tab_stops.add_tab_stop(Inches(0.55)); runx(q,f'{i}.\t',bold=True); runx(q,s_)
        imgs=[(save(img_by_name(src,n),f'{no}_{pi}_{k}'),c) for k,(n,c) in enumerate(p['imgs'])]
        sz=[Image.open(f).size for f,_ in imgs]
        hd('PROGRAM:' if True else '',6)
        if p['side']:
            tb=d.add_table(rows=2,cols=2); tb.alignment=docx.enum.table.WD_TABLE_ALIGNMENT.CENTER; tb.autofit=False; exp3.borders(tb)
            w1,w2=3.0,3.5
            for r in tb.rows:
                r._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit')); exp3.cellfmt(r.cells[0],w1); exp3.cellfmt(r.cells[1],w2)
            for i,x in enumerate(('Program','Output')):
                q=tb.cell(0,i).paragraphs[0]; q.alignment=AL.CENTER; q.paragraph_format.keep_with_next=True; q.paragraph_format.space_after=Pt(1); runx(q,x,11,True)
            c=tb.cell(1,0); first=True
            for ln in p['code']:
                q=c.paragraphs[0] if first else c.add_paragraph(); first=False
                q.alignment=AL.LEFT; q.paragraph_format.space_after=Pt(0); q.paragraph_format.left_indent=Inches(0.1)
                r=q.add_run(ln.replace(' ',' ') if ln.startswith(' ') else ln); r.font.name='Courier New'; r.font.size=Pt(9.5); r._r.get_or_add_rPr().rFonts.set(qn('w:hAnsi'),'Courier New')
            c=tb.cell(1,1); first=True
            n=len(imgs); hmax=(5.2 if n==1 else (2.2 if n==2 else 1.9))*SHRINK.get(no,1.0)
            for (f,cap),(W,H) in zip(imgs,sz):
                wi=min(3.3 if W>500 else 2.5,hmax*W/H)
                q=c.paragraphs[0] if first else c.add_paragraph(); first=False
                q.alignment=AL.CENTER; q.paragraph_format.space_before=Pt(3); q.paragraph_format.space_after=Pt(3)
                q.add_run().add_picture(f,width=Inches(wi))
        else:
            tb=d.add_table(rows=1,cols=1); tb.alignment=docx.enum.table.WD_TABLE_ALIGNMENT.CENTER; tb.autofit=False; exp3.borders(tb)
            exp3.cellfmt(tb.cell(0,0),6.5); c=tb.cell(0,0); first=True
            for ln in p['code']:
                q=c.paragraphs[0] if first else c.add_paragraph(); first=False
                q.paragraph_format.space_after=Pt(0); q.paragraph_format.left_indent=Inches(0.15)
                r=q.add_run(ln); r.font.name='Courier New'; r.font.size=Pt(10.5); r._r.get_or_add_rPr().rFonts.set(qn('w:hAnsi'),'Courier New')
            hd('OUTPUT:',8)
            for (f,cap),(W,H) in zip(imgs,sz):
                q=para(cap,bold=False,size=10.5,align=AL.CENTER,before=3,after=1,keep=True)
                q=para(align=AL.CENTER,before=0,after=4,ls=1.0); q.add_run().add_picture(f,width=Inches(min(6.2,2.0*W/H)))
        if p.get('result') and pi<len(parts)-1:
            h=hd('RESULT:',6); para(p['result'],left=0.3,after=2)
    return None

def emit_vl(B,src,no,date,title,s,e,ranges_start=None):
    import wtbuild as W_
    ch=list(src.body); d=B.dest; para=B.para; runx=B.runx
    pbp=para(after=0,ls=1.0); pbp.paragraph_format.page_break_before=True
    ppr=pbp._p.get_or_add_pPr(); rp=OxmlElement('w:rPr'); rp.append(B.mkel('sz',val=2)); ppr.append(rp)
    B.title_table(no,date,title); para(after=2,ls=1.0,size=4)
    idx=list(range(s,e))
    sim=next(j for j in idx if re.match(r'(?i)^\s*simulations?',ptext(ch[j])))
    scen=[j for j in idx if re.match(r'(?i)^\s*scenario',ptext(ch[j])) and j>=sim]
    # in some units SIMULATIONS and SCENARIO-1 share a paragraph
    if not scen or scen[0]>sim+0 and re.match(r'(?i)^\s*simulations?\s*scenario',ptext(ch[sim]).replace('\n','')): scen=[sim]+[j for j in scen if j!=sim]
    res=max(j for j in idx if re.match(r'(?i)^\s*result:?\s*$',ptext(ch[j])))
    # intro (skip title paragraph at s-?)
    for j in range(s,sim):
        x=ch[j]
        if x.tag!=W_.w('p'): continue
        if x.find('.//'+W_.w('txbxContent')) is not None or x.find('.//'+W_.w('pict')) is not None and not ptext(x).strip(): continue
        if x.find('.//'+W_.w('sectPr')) is not None and not ptext(x).strip() and x.find('.//'+W_.w('drawing')) is None: continue
        np_,hi=W_.convert_par(src,x,no); W_.dosubs(np_)
        if not ptext(np_).strip() and not hi: continue
        pPr=np_.find(W_.w('pPr'))
        for t in pPr.findall(W_.w('ind')):
            if int(t.get(W_.w('left'),'0'))>1500: t.set(W_.w('left'),'1500')
        W_.dbody.insert(list(W_.dbody).index(W_.sect),np_)
    para('SIMULATIONS:',bold=True,before=8,after=3,align=AL.LEFT,keep=True)
    bounds=scen+[res]
    for si,(a,b) in enumerate(zip(bounds,bounds[1:])):
        imgs=[]
        for j in range(a,b):
            for bl in ch[j].iter(A+'blip'):
                if bl.get(R):
                    blob=src.rels[bl.get(R)].target_part.blob; imgs.append(save(blob,f'vl{no}_{si}_{len(imgs)}'))
        name=re.sub(r'(?i)^\s*simulations?:?\s*','',ptext(ch[a]).replace('\n',' ')).strip().rstrip(':')
        hp=para(name.upper().replace('–','-').replace('-','–') if False else name,bold=True,size=12,before=8,after=3,align=AL.LEFT,keep=True)
        para('INPUT PARAMETERS:',bold=True,size=11,before=0,after=2,align=AL.LEFT,left=0.2,keep=True)
        if not imgs: continue
        sz=[Image.open(f).size for f in imgs]
        q=para(align=AL.CENTER,before=0,after=4,ls=1.0,keep=True); q.add_run().add_picture(imgs[0],width=Inches(min(5.6,1.15*sz[0][0]/sz[0][1])))
        rest=imgs[1:]
        if rest:
            tb=d.add_table(rows=1,cols=len(rest)); tb.alignment=docx.enum.table.WD_TABLE_ALIGNMENT.CENTER; tb.autofit=False
            wcol=6.4/len(rest)
            for r in tb.rows: r._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
            for i,f in enumerate(rest):
                c=tb.cell(0,i); exp3.cellfmt(c,wcol); W0,H0=Image.open(f).size
                wi=min(wcol-0.15,1.3*W0/H0) if len(rest)>1 else min(3.5,1.4*W0/H0)
                qq=c.paragraphs[0]; qq.alignment=AL.CENTER; qq.paragraph_format.space_after=Pt(2); qq.add_run().add_picture(f,width=Inches(wi))
        para(after=2,size=4,ls=1.0)
    return ptext(ch[res+1]).strip() if ptext(ch[res+1]).strip() else ptext(ch[res+2]).strip()
