import re,os,io,sys
sys.path.insert(0,'../wtlab')
from wtlib import *
from docx.shared import Pt,Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH as AL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image
A='{http://schemas.openxmlformats.org/drawingml/2006/main}'
R='{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed'
OPS='ORG|MOV|ADDC|ADD|SUBB|MUL|DIV|ANL|ORL|XRL|CPL|RLC|RRC|RL|RR|CLR|SETB|END|INC|DEC|DJNZ|CJNE|JC|JNC|JZ|JNZ|SJMP|SWAP|XCH'
def asm_lines(text):
    t=re.sub(r'\s+',' ',text).strip()
    t=re.sub(r'(?<=[0-9A-F]H)(?=(?:'+OPS+r')\b)',' ',t)
    t=re.sub(r'(?<=[,A-Z0-9])(?=(?:'+OPS+r')\s)',' ',t) if False else t
    return [m.group(0).strip() for m in re.finditer(r'(?:L\d+:\s*)?(?:'+OPS+r')\b(?:\s+[^\s]+)?',t)]
def parse(src,a,b):
    ch=list(src.body)
    heads=[j for j in range(a,b) if 'Objective:' in ptext(ch[j])]
    end=next(j for j in range(a,b) if re.match(r'(?i)^\s*result:?\s*$',ptext(ch[j])) and j>heads[-1]+20)
    blocks=[]
    for i,h in enumerate(heads):
        e=heads[i+1] if i+1<len(heads) else end
        title=ptext(ch[h]).replace('Objective:','').strip()
        obj=None;lab=None;d={'Input':[],'Calculation':[],'Result':[],'Program':[]};img=None;seen_out=False
        # group label detection
        grp=None
        for j in range(h+1,e):
            t=ptext(ch[j]).strip()
            blips=[bl for bl in ch[j].iter(A+'blip') if bl.get(R)]
            if (seen_out or re.match(r'^Output:?',t)) and blips and img is None:
                img=(src.rels[blips[0].get(R)].target_part.blob)
            m=re.match(r'^(Input|Calculation|Result|Program|Output):?\s*(.*)$',t)
            if m:
                lab=m.group(1); rest=m.group(2).strip()
                if lab=='Output': seen_out=True; lab=None
                elif rest:
                    if lab=='Result' and 'Program:' in rest:
                        r1,p1=rest.split('Program:',1); d['Result'].append(r1.strip()); d['Program'].append(p1.strip()); lab='Program'
                    else: d[lab].append(rest)
                continue
            if t and lab=='Result' and 'Program:' in t and not seen_out:
                r1,p1=t.split('Program:',1); d['Result'].append(r1.strip()); d['Program'].append(p1.strip()); lab='Program'
            elif t and lab and not seen_out: d[lab].append(t)
            elif t and not lab and not seen_out and obj is None: obj=t
            if t=='' and lab=='Calculation' and not seen_out: d['Calculation'].append('---')
        if obj is None: raise RuntimeError(title)
        blocks.append((title,obj,d,img))
    return blocks
def cellfmt(c,w):
    tcPr=c._tc.get_or_add_tcPr()
    for e in tcPr.findall(qn('w:tcW')): tcPr.remove(e)
    tw=OxmlElement('w:tcW'); tw.set(qn('w:w'),str(int(w*1440))); tw.set(qn('w:type'),'dxa'); tcPr.insert(0,tw)
def borders(tbl):
    tblPr=tbl._tbl.tblPr
    b=OxmlElement('w:tblBorders')
    for e in('top','left','bottom','right','insideH','insideV'):
        x=OxmlElement('w:'+e);x.set(qn('w:val'),'single');x.set(qn('w:sz'),'6');x.set(qn('w:space'),'0');x.set(qn('w:color'),'000000');b.append(x)
    nxt=None
    for tag in ('shd','tblLayout','tblCellMar','tblLook'):
        nxt=tblPr.find(qn('w:'+tag))
        if nxt is not None: break
    if nxt is not None: nxt.addprevious(b)
    else: tblPr.append(b)
def emit3(B,src,a,b,date,title):
    d=B.dest; para=B.para; runx=B.runx
    pbp=para(after=0,ls=1.0); pbp.paragraph_format.page_break_before=True
    ppr=pbp._p.get_or_add_pPr(); rp=OxmlElement('w:rPr'); rp.append(B.mkel('sz',val=2)); ppr.append(rp)
    B.title_table('3',date,title); para(after=2,ls=1.0,size=4)
    ch=list(src.body)
    def hd(t): return para(t,bold=True,before=6,after=3,align=AL.LEFT,keep=True)
    hd('AIM:'); para(ptext(ch[505]).strip(),left=0.3,after=4)
    hd('DESCRIPTION:')
    for j in (507,508): para(ptext(ch[j]).strip(),left=0.3,after=4)
    blocks=parse(src,a,b)
    ar,lg=0,0
    for k,(t,obj,dd,img) in enumerate(blocks):
        if k==0 or k==8:
            para('ARITHMETIC OPERATIONS' if k==0 else 'LOGICAL OPERATIONS',bold=True,size=12,align=AL.CENTER,before=10,after=4,keep=True)
        name=re.sub(r'^\s*\d*\s*\(?([ivx]+)\)?[:)]?\s*',r'(\1) ',t).strip()
        grp=(k//2+1) if k<8 else 5
        sub=f'{grp} {name}'
        h=para(sub,bold=True,before=8,after=2,align=AL.LEFT,keep=True,left=0.1)
        p=para(align=AL.LEFT,after=4,left=0.3,keep=True); runx(p,'Objective: ',bold=True); runx(p,obj)
        # IRC table
        has=[x for x in ('Input','Calculation','Result') if dd[x]]
        if has:
            tb=d.add_table(rows=2,cols=len(has)); tb.alignment=docx.enum.table.WD_TABLE_ALIGNMENT.CENTER; tb.autofit=False; borders(tb)
            wid=6.2/len(has)
            for i,x in enumerate(has):
                c0=tb.cell(0,i); c1=tb.cell(1,i); cellfmt(c0,wid); cellfmt(c1,wid)
                q=c0.paragraphs[0]; q.alignment=AL.CENTER; q.paragraph_format.space_after=Pt(1); runx(q,x,11,True)
                lines=dd[x]; first=True
                for ln in lines:
                    if ln=='---': ln='─'*10
                    q=c1.paragraphs[0] if first else c1.add_paragraph(); first=False
                    q.alignment=AL.CENTER; q.paragraph_format.space_after=Pt(0); q.paragraph_format.space_before=Pt(0); q.paragraph_format.keep_with_next=True
                    runx(q,ln,11)
            for r in tb.rows:
                r._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
                for c in r.cells:
                    for q in c.paragraphs: q.paragraph_format.keep_with_next=True
        # program | output
        pt=d.add_table(rows=2,cols=2); pt.alignment=docx.enum.table.WD_TABLE_ALIGNMENT.CENTER; pt.autofit=False; borders(pt)
        w1,w2=1.9,4.5
        for r in pt.rows:
            r._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
            cellfmt(r.cells[0],w1); cellfmt(r.cells[1],w2)
        for i,x in enumerate(('Program','Output')):
            q=pt.cell(0,i).paragraphs[0]; q.alignment=AL.CENTER; q.paragraph_format.space_after=Pt(1); q.paragraph_format.keep_with_next=True; runx(q,x,11,True)
        code=asm_lines(' '.join(dd['Program']))
        c=pt.cell(1,0); first=True
        for ln in code:
            q=c.paragraphs[0] if first else c.add_paragraph(); first=False
            q.alignment=AL.LEFT; q.paragraph_format.space_after=Pt(0); q.paragraph_format.left_indent=Inches(0.15)
            r=q.add_run(ln.replace(' ',' ',1) if False else ln); r.font.name='Courier New'; r.font.size=Pt(10.5); r._r.get_or_add_rPr().rFonts.set(qn('w:hAnsi'),'Courier New')
        if img:
            os.makedirs('e3img',exist_ok=True); fp=f'e3img/s{k}.jpg'
            Image.open(io.BytesIO(img)).convert('RGB').save(fp,quality=88)
            W,H=Image.open(fp).size
            wi=min(4.3,1.55*W/H)
            q=pt.cell(1,1).paragraphs[0]; q.alignment=AL.CENTER; q.paragraph_format.space_before=Pt(2); q.paragraph_format.space_after=Pt(2)
            q.add_run().add_picture(fp,width=Inches(wi))
        para(after=0,size=4,ls=1.0)
    # final result
    end=max(j for j in range(a,b) if re.match(r'(?i)^\s*result:?\s*$',ptext(ch[j])))
    return ptext(ch[end+1]).strip() if end+1<b else ptext(ch[end+1]).strip(),
