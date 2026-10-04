import copy, json, os, math
from docx import Document
from docx.shared import Pt, Inches, Emu, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH as AL, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph
from PIL import Image, ImageOps

TITLE = 'ADAPTIVE AR WORK INSTRUCTIONS FOR MACHINE ASSEMBLY LINES'
PAGES = json.load(open('pages.json')) if os.path.exists('pages.json') else [4,5,8,11,15,18]
FONT = 'Times New Roman'
d = Document('vr_template.docx')
body = d.element.body
sect_final = body.find(qn('w:sectPr'))
# --- discard the template's flowing content: it is rebuilt cleanly below (fixes its alignment problems) ---
for e in list(body):
    if e is not sect_final: body.remove(e)
# keep the header reference of the first template section, use one section for the 3 front pages
hdr_rid = 'rId6'
for hr in sect_final.findall(qn('w:headerReference')): sect_final.remove(hr)
for ch in list(sect_final):
    if ch.tag != qn('w:headerReference'): sect_final.remove(ch)
def mk(tag, **kw):
    e = OxmlElement('w:' + tag)
    for k, v in kw.items(): e.set(qn('w:' + k), str(v))
    return e
hr = mk('headerReference', type='default'); hr.set(qn('r:id'), hdr_rid); sect_final.append(hr)
sect_final.append(mk('pgSz', w=11906, h=16838))
sect_final.append(mk('pgMar', top=2900, right=900, bottom=1000, left=900, header=500, footer=450, gutter=0))
pb = mk('pgBorders', offsetFrom='page')
for s in ('top', 'left', 'bottom', 'right'): pb.append(mk(s, val='single', sz=12, space=20, color='000000'))
sect_final.append(pb)
sect_final.append(mk('cols', space=720))
sect = sect_final
students = [('ABINAYA S', '71812401002'), ('AISHWARYA U', '71812401006'), ('AJAY IYANRAJ', '71812401007'), ('AKHILESH RAJ P', '71812401009')]
BLUE = '000000'

def para(text='', bold=False, size=13, align=AL.JUSTIFY, italic=False, space_after=6, keep=False, font=FONT, indent=None, before=0):
    p = d.add_paragraph(); p.alignment = align
    pf = p.paragraph_format; pf.space_after = Pt(space_after); pf.space_before = Pt(before); pf.line_spacing = 1.15
    if keep: pf.keep_with_next = True
    if indent: pf.left_indent = Inches(indent)
    if text:
        r = p.add_run(text); r.bold = bold; r.italic = italic; r.font.size = Pt(size); r.font.name = font
        r._r.rPr.rFonts.set(qn('w:eastAsia'), font)
    return p
C = AL.CENTER
# ---------- page 1 : cover ----------
para('DEPARTMENT OF COMPUTER SCIENCE AND ENGINEERING', True, 14, C, space_after=44)
para('ACTIVE LEARNING METHODOLOGY REPORT', True, 14, C, space_after=44)
para(TITLE, True, 15, C, space_after=40)
for n, r_ in students: para(f'{n} ({r_})', False, 13, C, space_after=4)
para('', space_after=26)
para('THIRD YEAR B.E. CSE \u2013 V SEM', True, 13, C, space_after=6)
para('Academic Year 2026-2027', True, 13, C, space_after=40)
para('20CS2E51 & VIRTUAL REALITY/AUGMENTED REALITY', True, 13, C, space_after=120)
para('COURSE COORDINATOR', True, 12, AL.RIGHT, space_after=4)
para('Mrs.P.Sugantha Priyadharshini, AP(Sl.Gr.)/CSE', True, 12, AL.RIGHT, space_after=0)
# ---------- page 2 : index ----------
def tbl_borders(t, sz=6, color=BLUE):
    tblPr = t._tbl.tblPr
    for old in tblPr.findall(qn('w:tblBorders')): tblPr.remove(old)
    b = OxmlElement('w:tblBorders')
    for e in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        x = OxmlElement('w:' + e)
        for k, v in (('val', 'single'), ('sz', str(sz)), ('space', '0'), ('color', color)): x.set(qn('w:' + k), v)
        b.append(x)
    anchor = next((c for c in tblPr if c.tag in [qn('w:' + n) for n in ('shd', 'tblLayout', 'tblCellMar', 'tblLook', 'tblCaption')]), None)
    if anchor is not None: anchor.addprevious(b)
    else: tblPr.append(b)
def cell_text(cell, text, bold=False, size=12, align=C, vertical=False, fill=None, w=None):
    if w: cell.width = Inches(w)
    p = cell.paragraphs[0]; p.alignment = align
    p.paragraph_format.space_before = Pt(2); p.paragraph_format.space_after = Pt(2)
    for i, ln in enumerate(text.split('\n')):
        if i: p.add_run().add_break()
        r = p.add_run(ln); r.bold = bold; r.font.size = Pt(size); r.font.name = FONT
    tcPr = cell._tc.get_or_add_tcPr()
    if fill:
        s = OxmlElement('w:shd'); s.set(qn('w:val'), 'clear'); s.set(qn('w:color'), 'auto'); s.set(qn('w:fill'), fill); tcPr.append(s)
    if vertical:
        td = OxmlElement('w:textDirection'); td.set(qn('w:val'), 'btLr'); tcPr.append(td)
    va = OxmlElement('w:vAlign'); va.set(qn('w:val'), 'center'); tcPr.append(va)
p = para('DEPARTMENT OF COMPUTER SCIENCE AND ENGINEERING', True, 14, C, space_after=14); p.paragraph_format.page_break_before = True
para('20CS2E51 & VIRTUAL REALITY/AUGMENTED REALITY', True, 13, C, space_after=12)
para(TITLE, True, 13, C, space_after=14)
para('Assessment', True, 13, C, space_after=14)
para('INDEX', True, 14, C, space_after=10)
cols = [('S.\nNo.', 0.6, False), ('Name & Roll No.', 2.5, False), ('Introduction\n(2)', 0.75, True), ('Working Steps\n(4)', 0.8, True), ('Presentation\n(2)', 0.8, True), ('Report\n(2)', 0.7, True), ('Total\n(10)', 0.75, False)]
t = d.add_table(rows=1 + len(students), cols=7); t.autofit = False; t.alignment = WD_TABLE_ALIGNMENT.CENTER; tbl_borders(t)
for i, (h, w, v) in enumerate(cols):
    cell_text(t.rows[0].cells[i], h, True, 11.5, C, v, None, w)
t.rows[0].height = Inches(1.35)
from docx.enum.table import WD_ROW_HEIGHT_RULE
t.rows[0].height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
for ri, (n, r_) in enumerate(students, 1):
    row = t.rows[ri]; row.height = Inches(0.8); row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
    cell_text(row.cells[0], f'{ri}.', True, 12, C, False, None, cols[0][1])
    c = row.cells[1]; cell_text(c, n, True, 11.5, AL.LEFT, False, None, cols[1][1])
    p2 = c.add_paragraph(); p2.paragraph_format.space_after = Pt(2); r2 = p2.add_run(r_); r2.font.size = Pt(11.5); r2.font.name = FONT
    for k in range(2, 7): cell_text(row.cells[k], '', False, 12, C, False, None, cols[k][1])
for i, (_, w, _) in enumerate(cols): t.columns[i].width = Inches(w)
para('', space_after=40)
para('SIGNATURE OF COURSE COORDINATOR', True, 12, AL.RIGHT, space_after=4)
para('Mrs.P.Sugantha Priyadharshini, AP(Sl.Gr.)/CSE', True, 12, AL.RIGHT, space_after=0)
# ---------- page 3 : contents (table centred on the page) ----------
sp = para('', space_after=0); sp.paragraph_format.page_break_before = True
sp.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY; sp.paragraph_format.line_spacing = Pt(190)
para('CONTENTS', True, 14, C, space_after=14)
rows = [TITLE, 'INTRODUCTION', 'DESCRIPTION', 'WORKING STEPS', 'DIAGRAMMATIC REPRESENTATION', 'CONCLUSION']
t = d.add_table(rows=1 + len(rows), cols=3); t.autofit = False; t.alignment = WD_TABLE_ALIGNMENT.CENTER; tbl_borders(t)
for i, (h, w) in enumerate((('S. No.', 0.8), ('TITLE', 4.6), ('PAGE\nNUMBER', 1.1))): cell_text(t.rows[0].cells[i], h, True, 12, C, False, None, w)
for ri, name in enumerate(rows, 1):
    row = t.rows[ri]; row.height = Inches(0.42); row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
    cell_text(row.cells[0], '' if ri == 1 else f'{ri-1}.', False, 12, C, False, None, 0.8)
    cell_text(row.cells[1], name, ri == 1, 12, AL.LEFT, False, None, 4.6)
    cell_text(row.cells[2], str(PAGES[ri-1]), False, 12, C, False, None, 1.1)
for i, w in enumerate((0.8, 4.6, 1.1)): t.columns[i].width = Inches(w)
# section break: header stays on pages 1-3 only
endp = para('', space_after=0, size=2)
endp._p.get_or_add_pPr().append(copy.deepcopy(sect))
for old in list(sect.findall(qn('w:headerReference'))): sect.remove(old)
d.sections[-1].header.is_linked_to_previous = False
for p_ in d.sections[-1].header.paragraphs:
    for r_ in list(p_.runs): r_._r.getparent().remove(r_._r)
d.sections[-1].top_margin = Inches(0.8); d.sections[-1].header_distance = Inches(0.3)

# ---------- helpers for new content ----------

def para(text='', bold=False, size=13, align=AL.JUSTIFY, italic=False, space_after=6, keep=False, font=FONT, indent=None):
    p = d.add_paragraph()   # appended before sectPr by python-docx
    p.alignment = align
    pf = p.paragraph_format
    pf.space_after = Pt(space_after); pf.space_before = Pt(0); pf.line_spacing = 1.15
    if keep: pf.keep_with_next = True
    if indent: pf.left_indent = Inches(indent)
    if text:
        r = p.add_run(text); r.bold = bold; r.italic = italic; r.font.size = Pt(size); r.font.name = font
        r._r.rPr.rFonts.set(qn('w:eastAsia'), font)
    return p

def rich(parts, size=13, align=AL.JUSTIFY, space_after=4, indent=None, hanging=None):
    p = para('', align=align, space_after=space_after, indent=indent)
    for txt, b in parts:
        r = p.add_run(txt); r.bold = b; r.font.size = Pt(size); r.font.name = FONT
    if hanging: p.paragraph_format.first_line_indent = Inches(-hanging)
    return p

def heading(num, text, page_break=False):
    p = para((f'{num}. ' if num else '') + text.upper(), bold=True, size=14, align=AL.CENTER, space_after=8, keep=True)
    p.paragraph_format.space_before = Pt(10)
    if page_break: p.paragraph_format.page_break_before = True
    # bottom border
    pPr = p._p.get_or_add_pPr(); b = OxmlElement('w:pBdr'); bt = OxmlElement('w:bottom')
    for k, v in (('val', 'single'), ('sz', '8'), ('space', '1'), ('color', '2F5496')): bt.set(qn('w:' + k), v)
    b.append(bt)
    anc = next((c for c in pPr if c.tag in [qn('w:' + n) for n in ('shd', 'tabs', 'suppressAutoHyphens', 'spacing', 'ind', 'jc', 'rPr')]), None)
    if anc is not None: anc.addprevious(b)
    else: pPr.append(b)
    return p

def sub(text, page_break=False):
    p = para(text, bold=True, size=12, align=AL.LEFT, space_after=4, keep=True); p.paragraph_format.space_before = Pt(6)
    if page_break: p.paragraph_format.page_break_before = True
    return p

def bullet(text, boldlead=None):
    parts = ([(boldlead, True)] if boldlead else []) + [(text, False)]
    return rich([('•  ', False)] + parts, indent=0.3, hanging=0.2)

def step(n, text, boldlead=None):
    size = 13
    p = para('', align=AL.JUSTIFY, space_after=5)
    pf = p.paragraph_format
    pf.left_indent = Inches(1.35); pf.first_line_indent = Inches(-0.95)
    pf.tab_stops.add_tab_stop(Inches(1.35))
    r = p.add_run(f'Step {n}:'); r.bold = True; r.font.size = Pt(13); r.font.name = FONT
    r = p.add_run('\t'); r.font.size = Pt(12)
    if boldlead:
        r = p.add_run(boldlead); r.bold = True; r.font.size = Pt(13); r.font.name = FONT
    r = p.add_run(text); r.font.size = Pt(13); r.font.name = FONT
    return p

def set_borders(tbl, sz=6):
    tblPr = tbl._tbl.tblPr
    for old in tblPr.findall(qn('w:tblBorders')): tblPr.remove(old)
    b = OxmlElement('w:tblBorders')
    for e in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        x = OxmlElement('w:' + e)
        for k, v in (('val', 'single'), ('sz', str(sz)), ('space', '0'), ('color', '002060')): x.set(qn('w:' + k), v)
        b.append(x)
    anchor = next((c for c in tblPr if c.tag in [qn('w:' + n) for n in ('shd', 'tblLayout', 'tblCellMar', 'tblLook', 'tblCaption')]), None)
    if anchor is not None: anchor.addprevious(b)
    else: tblPr.append(b)

def shade(cell, color):
    tcPr = cell._tc.get_or_add_tcPr(); s = OxmlElement('w:shd')
    s.set(qn('w:val'), 'clear'); s.set(qn('w:color'), 'auto'); s.set(qn('w:fill'), color); tcPr.append(s)

def table(headers, rows, widths, size=10.5):
    t = d.add_table(rows=1, cols=len(headers)); t.autofit = False
    set_borders(t)
    def fill(cell, text, bold=False, w=None, center=False, white=False):
        cell.width = Inches(w); p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(2); p.paragraph_format.space_before = Pt(2)
        p.alignment = AL.CENTER if center else AL.LEFT
        r = p.add_run(text); r.bold = bold; r.font.size = Pt(size); r.font.name = FONT
        if white: r.font.color.rgb = RGBColor(255, 255, 255)
    for i, h in enumerate(headers):
        fill(t.rows[0].cells[i], h, True, widths[i], True, True); shade(t.rows[0].cells[i], '6666A0')
    trPr = t.rows[0]._tr.get_or_add_trPr(); th = OxmlElement('w:tblHeader'); trPr.append(th)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row): fill(cells[i], v, i == 0, widths[i])
    for i, w in enumerate(widths): t.columns[i].width = Inches(w)
    for r in t.rows[:-1]:
        for c in r.cells:
            for p_ in c.paragraphs: p_.paragraph_format.keep_with_next = True
    for r in t.rows:
        trPr = r._tr.get_or_add_trPr(); cs = OxmlElement('w:cantSplit'); trPr.append(cs)
    para('', space_after=4)
    return t

def code_block(lines, size=8):
    t = d.add_table(rows=1, cols=1); t.autofit = False
    set_borders(t, 8)
    cell = t.rows[0].cells[0]; cell.width = Inches(7.2); shade(cell, 'EEF3FA')
    first = True
    for ln in lines:
        p = cell.paragraphs[0] if first else cell.add_paragraph(); first = False
        pf = p.paragraph_format; pf.space_after = Pt(0); pf.space_before = Pt(0); pf.line_spacing = 1.0
        p.alignment = AL.LEFT
        r = p.add_run(ln if ln else ' '); r.font.name = 'Courier New'; r.font.size = Pt(size)
        r._r.rPr.rFonts.set(qn('w:hAnsi'), 'Courier New'); r._r.rPr.rFonts.set(qn('w:cs'), 'Courier New')
    return t

def figure(path, title, caption, width=6.4, maxh=8.8, new_page=False):
    im = Image.open(path).convert('RGB'); width = min(width, maxh * im.width / im.height); im = ImageOps.expand(im, border=2, fill=(0, 0, 0))
    out = path.replace('.png', '_b.png'); im.save(out)
    if title:
        p = para(title, bold=True, size=12, align=AL.LEFT, space_after=4, keep=True)
        p.paragraph_format.space_before = Pt(6)
        if new_page: p.paragraph_format.page_break_before = True
    p = para('', align=AL.CENTER, space_after=2, keep=True)
    p.add_run().add_picture(out, width=Inches(width))
    para(caption, italic=True, size=10.5, align=AL.CENTER, space_after=6)

