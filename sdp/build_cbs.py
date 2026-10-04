import docx,json,re,copy,glob,sys,os
from docx.shared import Pt,Inches,Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH as AL, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from PIL import Image
S='/tmp/claude-0/-home-user-ashwinsathishkumar846-wq/cd0e2464-2e83-5cb0-be37-3d5a128c853d/scratchpad/sdp'
OUT=sys.argv[1] if len(sys.argv)>1 else 'AJAY_I_-_71812401007_SOFTWARE_DEVELOPMENT_PROCESS.docx'
FONT='Times New Roman'
d=docx.Document('base.docx'); body=d.element.body
exps=json.load(open('struct_cbs.json'))
# ---- keep first header/footer refs, drop everything else, rebuild one section ----
sects=body.findall('.//'+qn('w:sectPr'))
refs=[copy.deepcopy(e) for e in sects[0] if e.tag in (qn('w:headerReference'),qn('w:footerReference'))]
for e in list(body): body.remove(e)
sp=OxmlElement('w:sectPr')
for r in refs: sp.append(r)
def mk(tag,**kw):
    e=OxmlElement('w:'+tag)
    for k,v in kw.items(): e.set(qn('w:'+k),str(v))
    return e
sp.append(mk('pgSz',w=11910,h=16840)); sp.append(mk('pgMar',top=1040,right=708,bottom=1700,left=708,header=725,footer=1300,gutter=0))
pb=mk('pgBorders',offsetFrom='page')
for s in('top','left','bottom','right'): pb.append(mk(s,val='single',sz=12,space=24,color='000000'))
sp.append(pb); sp.append(mk('pgNumType',start=71)); sp.append(mk('cols',space=720)); body.append(sp)
def para(text='',bold=False,size=12,align=AL.JUSTIFY,before=0,after=4,left=0,first=0,keep=False,ls=1.12,italic=False):
    p=d.add_paragraph(); pf=p.paragraph_format; p.alignment=align
    pf.space_before=Pt(before); pf.space_after=Pt(after); pf.line_spacing=ls
    if left: pf.left_indent=Inches(left)
    if first: pf.first_line_indent=Inches(first)
    if keep: pf.keep_with_next=True
    if text:
        r=p.add_run(text); r.bold=bold; r.italic=italic; r.font.size=Pt(size); r.font.name=FONT; r._r.rPr.rFonts.set(qn('w:eastAsia'),FONT)
    return p
def run(p,text,bold=False,size=12):
    r=p.add_run(text); r.bold=bold; r.font.size=Pt(size); r.font.name=FONT; r._r.rPr.rFonts.set(qn('w:eastAsia'),FONT); return r
def img_path(n):
    f=glob.glob(f'rebuilt/image{n}.*')
    return f[0] if f else glob.glob(f'{S}/x/word/media/image{n}.*')[0]
CUR=[0]
def add_image(n):
    path=img_path(n); W,H=Image.open(path).size
    w=6.4 if W>=1300 else (5.2 if W>=900 else max(2.6,W/200))
    w=min(w,3.45*W/H)*{7:0.8}.get(CUR[0],1)
    p=para(align=AL.CENTER,before=2,after=6,keep=False,ls=1.0)
    p.add_run().add_picture(path,width=Inches(w)); p.paragraph_format.keep_together=True
    return p
def set_cell_border(cell,**kw):
    tcPr=cell._tc.get_or_add_tcPr(); b=OxmlElement('w:tcBorders')
    for e in('top','left','bottom','right'):
        x=OxmlElement('w:'+e); x.set(qn('w:val'),'single'); x.set(qn('w:sz'),'8'); x.set(qn('w:space'),'0'); x.set(qn('w:color'),'000000'); b.append(x)
    tcPr.append(b)
def title_table(no,date,title):
    t=d.add_table(rows=2,cols=2); t.autofit=False; t.alignment=docx.enum.table.WD_TABLE_ALIGNMENT.CENTER
    widths=(1.55,5.45)
    for r in t.rows:
        r.height=Inches(0.42)
        for i,c in enumerate(r.cells): c.width=Inches(widths[i]); set_cell_border(c)
    for i,w in enumerate(widths): t.columns[i].width=Inches(w)
    def put(c,text,align,bold=True):
        p=c.paragraphs[0]; p.alignment=align; p.paragraph_format.space_before=Pt(4); p.paragraph_format.space_after=Pt(4)
        run(p,text,bold,12)
        tcPr=c._tc.get_or_add_tcPr(); v=OxmlElement('w:vAlign'); v.set(qn('w:val'),'center'); tcPr.append(v)
    put(t.cell(0,0),f'EX.NO.{no}',AL.LEFT); put(t.cell(1,0),f'DATE: {date}',AL.LEFT)
    m=t.cell(0,1).merge(t.cell(1,1)); put(m,title,AL.CENTER)
    for r in t.rows: 
        trPr=r._tr.get_or_add_trPr(); trPr.append(OxmlElement('w:cantSplit'))
    return t
def frame(p):
    pPr=p._p.get_or_add_pPr(); fp=OxmlElement('w:framePr')
    for k,v in dict(w=10494,hRule='auto',wrap='around',vAnchor='margin',hAnchor='margin',xAlign='left',yAlign='bottom').items(): fp.set(qn('w:'+k),str(v))
    pPr.insert_element_before(fp,'w:widowControl','w:numPr','w:suppressLineNumbers','w:pBdr','w:shd','w:tabs','w:suppressAutoHyphens','w:kinsoku','w:wordWrap','w:overflowPunct','w:topLinePunct','w:autoSpaceDE','w:autoSpaceDN','w:bidi','w:adjustRightInd','w:snapToGrid','w:spacing','w:ind','w:contextualSpacing','w:mirrorIndents','w:suppressOverlap','w:jc','w:textDirection','w:textAlignment','w:textboxTightWrap','w:outlineLvl','w:divId','w:cnfStyle','w:rPr','w:sectPr','w:pPrChange')
def head(text): return para(text,bold=True,before=8,after=3,align=AL.LEFT,keep=True)
def sub(text,center=False): return para(text,bold=True,before=6,after=3,align=AL.CENTER if center else AL.LEFT,keep=True)
def li(mark,text,term=None):
    p=para(left=0.5,first=-0.3,after=2.5,ls=1.1)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(0.5))
    run(p,mark+'\t',False)
    if term: run(p,term+' ',True)
    run(p,text); return p
import json as _j,os as _o
SPC=_j.load(open('spacers_cbs.json')) if _o.path.exists('spacers_cbs.json') else {}
# ---------------- normalise blocks ----------------
def norm(e):
    out=[]; blocks=e['b']; ctx=''; num=0; kind='bul'
    ALLCAP=lambda s:s.upper()==s and len(s)<45
    i=0
    while i<len(blocks):
        b=dict(blocks[i]); t=b.get('text','')
        if b['t']=='p' and out and out[-1]['t']=='p' and re.match(r'^(Review and Done\.)$',t): out[-1]['text']+=' '+t; i+=1; continue
        if b['t']=='p' and t.endswith(':') and ALLCAP(t): b['t']='sh'
        if b['t']=='li' and (ALLCAP(t) or t in('Creating Users','Creating Groups','Editing and Deleting Groups')): b['t']='sh'
        if b['t']=='p' and t in('Applications and Uses of JIRA','JIRA for Project Management Teams'): b['t']='sh'
        if b['t']=='p' and re.match(r'^\d+\.\s',t): b['t']='li'; b['text']=re.sub(r'^\d+\.\s*','',t)
        if b['t']=='li' and e['no']==8 and out and out[-1]['t']=='sh' and out[-1]['text'] in('ADDING A USER','EDITING A USER','DELETING A USER'): b['t']='p'
        if b['t']=='li' and re.match(r'^[A-Za-z ]+:$',t) and i+1<len(blocks) and blocks[i+1]['t']=='p': b['t']='term'
        out.append(b); i+=1
    return out
def render(e):
    bl=norm(e); ctx=''; n=0; prev=None; last_li=False
    ctxh=''
    for b in bl:
        t=b.get('text','')
        if b['t']=='h':
            ctx=t.upper(); head(t); n=0
        elif b['t']=='sh':
            ctxh=t; sub(t); n=0
        elif b['t']=='cap':
            para(t,bold=True,before=6,after=3,align=AL.LEFT,keep=True); n=0
        elif b['t']=='term':
            p=para(t,bold=True,before=3,after=1,align=AL.LEFT,keep=True,left=0.2)
        elif b['t']=='p':
            p=para(t,after=4); 
            if prev in('term',): p.paragraph_format.left_indent=Inches(0.2)
        elif b['t']=='li':
            if prev!='li': n=0
            n+=1
            numbered=('PROCEDURE' in ctx or 'ALGORITHM' in ctx or re.match(r'^[A-C]\.|^Creating|^Editing',ctxh)) and not 'OUTPUT' in ctx
            if e['no']==1 and 'DESCRIPTION' in ctx and not any(x['t']=='cap' for x in bl[:bl.index(b)]): li('✓',t)
            elif numbered: li(f'{n}.',t)
            else: li('•',t)
        elif b['t']=='img':
            add_image(b['n'])
        prev=b['t']
    # result handled separately
for e in exps:
    # page break paragraph, tiny
    pbp=para(after=0,ls=1.0,size=2); pbp.paragraph_format.page_break_before=True
    for r in pbp.runs: r.font.size=Pt(2)
    pPr=pbp._p.get_or_add_pPr(); 
    rp=OxmlElement('w:rPr'); sz=OxmlElement('w:sz'); sz.set(qn('w:val'),'2'); rp.append(sz); pPr.append(rp)
    para('CONTENT BEYOND SYLLABUS',bold=True,align=AL.CENTER,before=0,after=8)
    title_table(e['no'],e['date'],e['title'])
    para(after=2,ls=1.0,size=4)
    # split result off
    ri=max(i for i,b in enumerate(e['b']) if b['t']=='h' and b['text'].startswith('RESULT'))
    res=e['b'][ri:]; e2=dict(e); e2['b']=e['b'][:ri]
    CUR[0]=e['no']; render(e2)
    rtext=' '.join(b.get('text','') for b in res[1:]).strip()
    sp_=SPC.get(str(e['no']),0)
    p=head('RESULT:'); p.paragraph_format.keep_with_next=True; p.paragraph_format.space_before=Pt(sp_)
    q=para(rtext,left=0.3,after=0)
# move the sectPr to the end
body.remove(sp); body.append(sp)
d.core_properties.title='Software Development Process Laboratory - AJAY.I - 71812401007'
d.save(OUT); print('saved',OUT)
