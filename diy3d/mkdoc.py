import pymupdf as fitz,glob
from docx import Document
from docx.shared import Mm
from PIL import Image
fs=sorted(glob.glob('pg/p*.jpg'));doc=fitz.open();D=Document();first=True
for f in fs:
    im=Image.open(f);L=im.width>im.height;w,h=(842,595) if L else (595,842)
    p=doc.new_page(width=w,height=h);p.insert_image(p.rect,filename=f)
    sec=D.sections[0] if first else D.add_section();first=False
    sec.page_width,sec.page_height=(Mm(297),Mm(210)) if L else (Mm(210),Mm(297))
    for a in ('left_margin','right_margin','top_margin','bottom_margin'):setattr(sec,a,Mm(0))
    sec.header_distance=sec.footer_distance=Mm(0)
    par=D.add_paragraph();par.paragraph_format.space_before=par.paragraph_format.space_after=0
    par.add_run().add_picture(f,width=sec.page_width-Mm(.5))
doc.save('Evolution_of_CS_DIY_Cutout_Kit_STUDIO.pdf',deflate=True);D.save('Evolution_of_CS_DIY_Cutout_Kit_STUDIO.docx')
print(len(fs))
