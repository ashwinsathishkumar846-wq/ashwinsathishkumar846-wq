import re,sys,json
from docx import Document
from docx.shared import Pt,Cm,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH as AL, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from pygments import lex
from pygments.lexers import HtmlLexer,JavascriptLexer,CssLexer,BashLexer
from pygments.lexers.javascript import JavascriptLexer
from pygments.token import Token
try:
    from pygments.lexers import JsxLexer
except Exception:
    from pygments.lexers.jsx import JsxLexer
PAGES={}
BODY='Calibri';MONO='Consolas';W=9.2
NAVY='000000';BLUE='000000'
d=Document()
st=d.styles['Normal'];st.font.name=BODY;st.font.size=Pt(8.5)
st.element.rPr.rFonts.set(qn('w:eastAsia'),BODY)
st.paragraph_format.space_after=Pt(0)
s=d.sections[0]
s.page_width=Cm(21);s.page_height=Cm(29.7);s.left_margin=s.right_margin=Cm(1.1);s.top_margin=Cm(1.2);s.bottom_margin=Cm(1.4);s.footer_distance=Cm(0.45)
sp=s._sectPr
pb=OxmlElement('w:pgBorders');pb.set(qn('w:offsetFrom'),'page')
for k in ('top','left','bottom','right'):
    e=OxmlElement('w:'+k);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'10');e.set(qn('w:space'),'14');e.set(qn('w:color'),'000000');pb.append(e)
sp.find(qn('w:pgMar')).addnext(pb)
def fnt(r,size=None,bold=None,italic=None,color=None,name=None):
    nm=name or BODY;r.font.name=nm
    rf=r._r.get_or_add_rPr().find(qn('w:rFonts'))
    for a in ('w:ascii','w:hAnsi','w:cs','w:eastAsia'): rf.set(qn(a),nm)
    if size: r.font.size=Pt(size)
    if bold is not None: r.bold=bold
    if italic is not None: r.italic=italic
    if color: r.font.color.rgb=RGBColor.from_string(color)
def shade_p(p,fill):
    pPr=p._p.get_or_add_pPr();sh=OxmlElement('w:shd');sh.set(qn('w:val'),'clear');sh.set(qn('w:color'),'auto');sh.set(qn('w:fill'),fill);pPr.append(sh)
def shade_c(c,fill):
    tcPr=c._tc.get_or_add_tcPr();sh=OxmlElement('w:shd');sh.set(qn('w:val'),'clear');sh.set(qn('w:color'),'auto');sh.set(qn('w:fill'),fill);tcPr.append(sh)
def cell_borders(c,**kw):
    tcPr=c._tc.get_or_add_tcPr();b=OxmlElement('w:tcBorders')
    for k in ('top','left','bottom','right'):
        v=kw.get(k)
        e=OxmlElement('w:'+k)
        if v: e.set(qn('w:val'),'single');e.set(qn('w:sz'),str(v[0]));e.set(qn('w:space'),'0');e.set(qn('w:color'),v[1])
        else: e.set(qn('w:val'),'nil')
        b.append(e)
    tcPr.append(b)
def cell_mar(c,top=60,bottom=60,left=140,right=140):
    tcPr=c._tc.get_or_add_tcPr();m=OxmlElement('w:tcMar')
    for k,v in (('top',top),('left',left),('bottom',bottom),('right',right)):
        e=OxmlElement('w:'+k);e.set(qn('w:w'),str(v));e.set(qn('w:type'),'dxa');m.append(e)
    tcPr.append(m)
def tbl_fixed(t,widths):
    t.autofit=False;pr=t._tbl.tblPr;l=OxmlElement('w:tblLayout');l.set(qn('w:type'),'fixed');pr.append(l)
    for row in t.rows:
        for i,c in enumerate(row.cells): c.width=Cm(widths[i])
    for i,g in enumerate(t._tbl.tblGrid.findall(qn('w:gridCol'))): g.set(qn('w:w'),str(int(widths[i]*567)))
def no_split(row):
    row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
EM={'❌':'✗','✅':'✔','↙️':'','↘️':''}
def clean(t):
    for k,v in EM.items(): t=t.replace(k,v)
    return t
def runs(p,text,size=9.8,bold=False,italic=False,color=None):
    text=clean(text)
    parts=re.split(r'(\*\*.+?\*\*|`[^`]+`|⭐+)',text)
    for x in parts:
        if not x: continue
        if x.startswith('**'): fnt(p.add_run(x[2:-2]),size,True,italic,color)
        elif x.startswith('`'):
            r=p.add_run(' '+x[1:-1]+' ');fnt(r,size-0.8,False,False,'000000',MONO)
            rp=r._r.get_or_add_rPr();sh=OxmlElement('w:shd');sh.set(qn('w:val'),'clear');sh.set(qn('w:color'),'auto');sh.set(qn('w:fill'),'F1F3F5');rp.append(sh)
        elif x.startswith('⭐'): fnt(p.add_run('★'*len(x)),size,False,False,'E0A800',name='DejaVu Sans')
        else: fnt(p.add_run(x),size,bold,italic,color)
def para(text='',size=9.8,bold=False,italic=False,align=AL.LEFT,after=2,before=0,keep=False,left=0,color=None,line=1.0):
    p=d.add_paragraph();p.alignment=align;pf=p.paragraph_format;pf.space_after=Pt(after);pf.space_before=Pt(before);pf.keep_with_next=keep;pf.line_spacing=line
    if left: pf.left_indent=Cm(left)
    if text: runs(p,text,size,bold,italic,color)
    return p
def pbreak(): d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
def bar(text,fill=NAVY,size=15,color='FFFFFF',after=8,pbb=False):
    size=size*0.7;after=3;pbb=False
    t=d.add_table(rows=1,cols=1);t.alignment=WD_TABLE_ALIGNMENT.CENTER;tbl_fixed(t,[W]);c=t.rows[0].cells[0];shade_c(c,fill);cell_mar(c,45,45,110,110)
    cell_borders(c);p=c.paragraphs[0];p.paragraph_format.keep_with_next=True
    if pbb: p.paragraph_format.page_break_before=True
    runs(p,text,size,True,False,color);no_split(t.rows[0])
    sp=d.add_paragraph();sp.paragraph_format.space_after=Pt(0);sp.paragraph_format.line_spacing=Pt(3);sp.paragraph_format.keep_with_next=True
def h1(t):
    p=para(t,10.2,True,after=3,before=6,keep=True,color=NAVY)
    pPr=p._p.get_or_add_pPr();b=OxmlElement('w:pBdr');e=OxmlElement('w:bottom')
    for k,v in (('val','single'),('sz','8'),('space','2'),('color',BLUE)): e.set(qn('w:'+k),v)
    b.append(e);pPr.append(b);return p
def h2(t): return para(t,9.3,True,after=2,before=5,keep=True,color=BLUE)
def h3(t): return para(t,8.7,True,after=1.5,before=3,keep=True,color='333333')
def bullet(t,lvl=0):
    p=para('',9.8,after=1,left=0.5+lvl*0.4,line=1.0);p.paragraph_format.first_line_indent=Cm(-0.32)
    fnt(p.add_run('•\u00A0'),9.8,True,color=BLUE);runs(p,t,9.8);return p
def numitem(n,t):
    p=para('',9.8,after=1,left=0.55,line=1.0);p.paragraph_format.first_line_indent=Cm(-0.4)
    fnt(p.add_run(f'{n}.\u00A0'),9.8,True,color=BLUE);runs(p,t,9.8);return p
def quote(lines,green=False):
    t=d.add_table(rows=1,cols=1);t.alignment=WD_TABLE_ALIGNMENT.CENTER;tbl_fixed(t,[W]);c=t.rows[0].cells[0]
    fill='E8F5EC' if green else 'EAF2FC';edge='2E9B57' if green else BLUE
    shade_c(c,fill);cell_mar(c,40,40,100,90);cell_borders(c,left=(30,edge))
    first=True
    for ln in lines:
        p=c.paragraphs[0] if first else c.add_paragraph();first=False
        p.paragraph_format.space_after=Pt(1);p.paragraph_format.line_spacing=1.0
        if ln=='': continue
        runs(p,ln,9.8,False,False,'1B1B1B')
    no_split(t.rows[0]);sp=d.add_paragraph();sp.paragraph_format.space_after=Pt(0);sp.paragraph_format.line_spacing=Pt(3)
def mdtable(rows):
    hdr=rows[0];body=rows[2:]
    n=len(hdr)
    lens=[max(len(clean(r[i])) if i<len(r) else 0 for r in [hdr]+body) for i in range(n)]
    tot=sum(max(l,6) for l in lens);ws=[max(1.3,W*max(l,6)/tot) for l in lens];k=W/sum(ws);ws=[w*k for w in ws]
    t=d.add_table(rows=1+len(body),cols=n);t.alignment=WD_TABLE_ALIGNMENT.CENTER;tbl_fixed(t,ws)
    pr=t._tbl.tblPr;b=OxmlElement('w:tblBorders')
    for kk in ('top','left','bottom','right','insideH','insideV'):
        e=OxmlElement('w:'+kk);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'6');e.set(qn('w:space'),'0');e.set(qn('w:color'),'000000');b.append(e)
    pr.append(b)
    for i,h in enumerate(hdr):
        c=t.rows[0].cells[i];shade_c(c,'D9D9D9');cell_mar(c,30,30,70,70);p=c.paragraphs[0];p.alignment=AL.CENTER;runs(p,h,9,True,False,'000000')
    t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    for r,row in enumerate(body):
        no_split(t.rows[r+1])
        for i in range(n):
            c=t.rows[r+1].cells[i];cell_mar(c,25,25,70,70)
            p=c.paragraphs[0];p.paragraph_format.line_spacing=1.0;runs(p,row[i] if i<len(row) else '',9)
    sp=d.add_paragraph();sp.paragraph_format.space_after=Pt(0);sp.paragraph_format.line_spacing=Pt(3)
# ---------- code
BG='F3F3F3';FG='1A1A1A'
COL={Token.Keyword:'9B2D8B',Token.Name.Tag:'2F4B8F',Token.Name.Attribute:'7A5A12',Token.Literal.String:'2E7D32',Token.Literal.Number:'9C4A6B',
Token.Comment:'7A7F87',Token.Name.Function:'6A3FB5',Token.Name.Class:'6A3FB5',Token.Name.Builtin:'6A3FB5',Token.Operator.Word:'9B2D8B',Token.Name.Decorator:'6A3FB5',Token.Name.Namespace:'6A3FB5',Token.Name.Constant:'9C4A6B'}
def colour(tt):
    while tt is not None:
        if tt in COL: return COL[tt]
        tt=tt.parent
    return FG
LEX={'html':HtmlLexer(),'javascript':JavascriptLexer(),'jsx':JsxLexer(),'css':CssLexer(),'bash':BashLexer()}
LABEL={'html':'HTML','javascript':'JavaScript','jsx':'JavaScript (JSX)','css':'CSS','bash':'Bash','text':''}
def codeblock(lang,lines):
    while lines and not lines[-1].strip(): lines.pop()
    while lines and not lines[0].strip(): lines.pop(0)
    plain=(lang=='text')
    t=d.add_table(rows=(1 if plain else 2),cols=1);t.alignment=WD_TABLE_ALIGNMENT.CENTER;tbl_fixed(t,[W])
    ri=0
    if not plain:
        c=t.rows[0].cells[0];shade_c(c,BG);cell_mar(c,30,10,110,80);cell_borders(c)
        p=c.paragraphs[0];p.paragraph_format.keep_with_next=True
        fnt(p.add_run('</>  '),6.8,True,color='222222',name=MONO);fnt(p.add_run(LABEL[lang]),7.6,True,color='111111');no_split(t.rows[0])
    c=t.rows[ri+(0 if plain else 1)].cells[0];shade_c(c,BG);cell_mar(c,20,50,110,80);cell_borders(c)
    if len(lines)<=28: no_split(t.rows[-1])
    comp=[]
    for x in lines:
        if not x.strip() and comp and not comp[-1].strip(): continue
        comp.append(x)
    lines=comp
    src='\n'.join(lines)
    toks=list(lex(src,LEX[lang])) if not plain else [(Token.Text,src)]
    # split tokens into lines
    ln=[[]]
    for tt,val in toks:
        parts=val.split('\n')
        for i,pt in enumerate(parts):
            if i>0: ln.append([])
            if pt:
                c_=colour(tt)
                if c_==FG and tt in Token.Name and re.match(r'^([A-Z][A-Za-z]+|use[A-Z]\w+)$',pt) and lang in ('jsx','javascript'): c_='6A3FB5'
                ln[-1].append((c_,pt))
    if ln and not ln[-1]: ln.pop()
    first=True
    for L in ln:
        p=c.paragraphs[0] if first else c.add_paragraph();first=False
        pf=p.paragraph_format;pf.space_after=Pt(0);pf.line_spacing=Pt(7.6)
        if not L:
            pf.line_spacing=Pt(3.4);r=p.add_run('\u00A0');fnt(r,3,color=FG,name=MONO);continue
        # leading whitespace -> nbsp
        for col,txt in L:
            txt=txt.replace(' ',' ').replace('\t',' '*4)
            fnt(p.add_run(txt),6.3,False,color=col,name=MONO)
    sp=d.add_paragraph();sp.paragraph_format.space_after=Pt(0);sp.paragraph_format.line_spacing=Pt(3)
# ---------- footer
s.footer.is_linked_to_previous=False
fp=s.footer.paragraphs[0];fp.alignment=AL.CENTER
def fld(p,instr):
    r=p.add_run();fnt(r,10,color='555555')
    for k,tx in (('begin',None),('instr',instr),('separate',None),('text','1'),('end',None)):
        if k=='instr': e=OxmlElement('w:instrText');e.set(qn('xml:space'),'preserve');e.text=tx;r._r.append(e)
        elif k=='text': e=OxmlElement('w:t');e.text=tx;r._r.append(e)
        else: e=OxmlElement('w:fldChar');e.set(qn('w:fldCharType'),k);r._r.append(e)
fld(fp,' PAGE ')


for nm in ('NAVY','BLUE'): globals()[nm]='000000'
S0=9.8
def codebox(lines):
    t=d.add_table(rows=1,cols=1);t.alignment=WD_TABLE_ALIGNMENT.CENTER;tbl_fixed(t,[W]);c=t.rows[0].cells[0];shade_c(c,'F2F2F2');cell_mar(c,40,40,100,80);cell_borders(c,top=(4,'7F7F7F'),left=(4,'7F7F7F'),bottom=(4,'7F7F7F'),right=(4,'7F7F7F'))
    first=True
    for l in lines:
        p=c.paragraphs[0] if first else c.add_paragraph();first=False
        p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=Pt(9.4)
        fnt(p.add_run((l.replace(' ',' ') if l.strip() else ' ')),7.8,False,color='000000',name=MONO)
    no_split(t.rows[0]);sp=d.add_paragraph();sp.paragraph_format.space_after=Pt(0);sp.paragraph_format.line_spacing=Pt(3)
sec=d.sections[0]
cols=sec._sectPr.find(qn('w:cols'))
if cols is None:
    cols=OxmlElement('w:cols');sec._sectPr.append(cols)
cols.set(qn('w:num'),'2');cols.set(qn('w:space'),'340');cols.set(qn('w:sep'),'1')
def strip_emoji(t): return re.sub(r'[📘⚡]\s*','',t).replace('\\.','.').strip()
L=open('notes.txt',encoding='utf-8').read().split('\n')
i=0
while i<len(L):
    ln=L[i];s_=ln.rstrip()
    if not s_.strip() or s_.strip() in ('....','---'): i+=1;continue
    if s_.startswith('```'):
        i+=1;buf=[]
        while i<len(L) and not L[i].startswith('```'): buf.append(L[i]);i+=1
        i+=1;codebox(buf);continue
    if s_.startswith('|'):
        rows=[]
        while i<len(L) and L[i].startswith('|'):
            rows.append([x.strip() for x in L[i].strip().strip('|').split('|')]);i+=1
        mdtable(rows);continue
    m=re.match(r'^(#{1,3}) (.*)',s_)
    if m:
        lv=len(m.group(1));tx=strip_emoji(m.group(2))
        if lv==1: bar(re.sub(r'\s+',' ',tx),'000000',15)
        elif lv==2:
            p=para(tx,10.8,True,after=2,before=5,keep=True)
            pPr=p._p.get_or_add_pPr();b=OxmlElement('w:pBdr')
            for k in ('top','bottom'):
                e=OxmlElement('w:'+k)
                for kk,v in (('val','single'),('sz','8'),('space','2'),('color','000000')): e.set(qn('w:'+kk),v)
                b.append(e)
            pPr.append(b)
        else: para(tx,10,True,after=1.5,before=3,keep=True)
        i+=1;continue
    mb=re.match(r'^\s*[-*] (.*)',s_)
    if mb: bullet(mb.group(1));i+=1;continue
    mn=re.match(r'^(\d+)(?:\\)?\. (.*)',s_)
    if mn: numitem(mn.group(1),mn.group(2));i+=1;continue
    buf=[s_.strip()];i+=1
    while i<len(L) and L[i].strip() and not re.match(r'^(#|```|\||[-*] |\d+\\?\. )',L[i]):
        buf.append(L[i].strip());i+=1
    txt=' '.join(buf)
    # bold labels like "Example:", "Conclusion:", "Remember:", "Answer:" and memory tricks
    txt=re.sub(r'^(Example|Conclusion|Remember|Memory trick|Answer|Explanation|Benefits|Working|Steps|Full question|Suitable choice|Workflow|Example commands|Example Dockerfile):',r'**\1:**',txt)
    para(txt,S0,after=2)
d.core_properties.title='DevOps Ten-Mark and Two-Mark Questions with Answers';d.core_properties.author=''
d.save('DevOps_Questions_and_Answers.docx');print('saved')
