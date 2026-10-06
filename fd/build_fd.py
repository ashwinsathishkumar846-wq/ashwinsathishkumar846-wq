import json,re,sys,copy
from docx import Document
from docx.shared import Pt,Cm,RGBColor,Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH as AL, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
PM=json.load(open('parsed.json')); CODE=json.load(open('code.json'))
PAGES=json.load(open('pagemap.json')) if len(sys.argv)>1 and sys.argv[1]=='final' else {}
FONT='Times New Roman'
d=Document()
st=d.styles['Normal'];st.font.name=FONT;st.font.size=Pt(12)
st.element.rPr.rFonts.set(qn('w:eastAsia'),FONT);st.element.rPr.rFonts.set(qn('w:cs'),FONT)
st.paragraph_format.space_after=Pt(0);st.paragraph_format.space_before=Pt(0)
s0=d.sections[0]
for s in [s0]:
    s.page_width=Cm(21);s.page_height=Cm(29.7);s.left_margin=s.right_margin=Cm(2.54);s.top_margin=Cm(1.6);s.bottom_margin=Cm(1.8)
    s.header_distance=Cm(0.8);s.footer_distance=Cm(0.9)
def setfont(r,size=None,bold=None,italic=None,name=None,color=None):
    r.font.name=name or FONT
    rp=r._r.get_or_add_rPr();rf=rp.find(qn('w:rFonts'))
    for a in ('w:ascii','w:hAnsi','w:cs','w:eastAsia'): rf.set(qn(a),name or FONT)
    if size: r.font.size=Pt(size)
    if bold is not None: r.bold=bold
    if italic is not None: r.italic=italic
    if color: r.font.color.rgb=RGBColor.from_string(color)
def para(text='',size=12,bold=False,italic=False,align=AL.JUSTIFY,after=6,before=0,first=0,keep=False,line=1.12,left=0,pbb=False,name=None):
    p=d.add_paragraph();pf=p.paragraph_format;p.alignment=align;pf.space_after=Pt(after);pf.space_before=Pt(before)
    if first: pf.first_line_indent=Cm(first)
    if left: pf.left_indent=Cm(left)
    pf.line_spacing=line;pf.keep_with_next=keep;pf.page_break_before=pbb;pf.widow_control=True
    if text: add_runs(p,text,size,bold,italic,name)
    return p
BOLD_TERMS=[]
def add_runs(p,text,size=12,bold=False,italic=False,name=None):
    parts=re.split(r'(\*\*.+?\*\*)',text)
    for x in parts:
        if not x: continue
        if x.startswith('**'): r=p.add_run(x[2:-2]);setfont(r,size,True,italic,name)
        else: r=p.add_run(x);setfont(r,size,bold,italic,name)
def H1(t,pbb=True):
    return para(t,14,True,align=AL.CENTER,after=10,before=0,keep=True,pbb=pbb,line=1.0)
def H2(t,before=8):
    return para(t,12,True,align=AL.LEFT,after=4,before=before,keep=True,line=1.0)
def H3(t):
    return para(t,11.5,True,align=AL.LEFT,after=2,before=5,keep=True,line=1.0)
def body(t,first=1.0,after=6): return para(t,12,False,first=first,after=after)
def body11(t,first=1.0,after=5): return para(t,11.5,False,first=first,after=after,line=1.08)
def bullet(t,size=10.5):
    p=para('',size,align=AL.JUSTIFY,after=2,left=1.0,line=1.08);p.paragraph_format.first_line_indent=Cm(-0.5)
    add_runs(p,'•  '+t,size);return p
def shade(cell,fill):
    tcPr=cell._tc.get_or_add_tcPr();sh=OxmlElement('w:shd');sh.set(qn('w:val'),'clear');sh.set(qn('w:color'),'auto');sh.set(qn('w:fill'),fill);tcPr.append(sh)
def cell_text(cell,text,size=10.5,bold=False,align=AL.LEFT,italic=False):
    cell.text='';p=cell.paragraphs[0];p.alignment=align;p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.05
    add_runs(p,text,size,bold,italic)
    tcPr=cell._tc.get_or_add_tcPr();va=OxmlElement('w:vAlign');va.set(qn('w:val'),'center');tcPr.append(va)
def set_cell_margins(tbl,top=50,bottom=50,left=90,right=90):
    tblPr=tbl._tbl.tblPr;m=OxmlElement('w:tblCellMar')
    for k,v in (('top',top),('left',left),('bottom',bottom),('right',right)):
        e=OxmlElement('w:'+k);e.set(qn('w:w'),str(v));e.set(qn('w:type'),'dxa');m.append(e)
    tblPr.append(m)
def borders(tbl,sz=6):
    tblPr=tbl._tbl.tblPr;b=OxmlElement('w:tblBorders')
    for k in ('top','left','bottom','right','insideH','insideV'):
        e=OxmlElement('w:'+k);e.set(qn('w:val'),'single');e.set(qn('w:sz'),str(sz));e.set(qn('w:space'),'0');e.set(qn('w:color'),'000000');b.append(e)
    tblPr.append(b)
def fixed(tbl,widths):
    tbl.autofit=False;tblPr=tbl._tbl.tblPr
    lay=OxmlElement('w:tblLayout');lay.set(qn('w:type'),'fixed');tblPr.append(lay)
    for row in tbl.rows:
        for i,c in enumerate(row.cells): c.width=Cm(widths[i])
    grid=tbl._tbl.tblGrid
    for i,gc in enumerate(grid.findall(qn('w:gridCol'))): gc.set(qn('w:w'),str(int(widths[i]*567)))
def row_opts(row,header=False,cant=True,height=None):
    trPr=row._tr.get_or_add_trPr()
    if cant: trPr.append(OxmlElement('w:cantSplit'))
    if header: trPr.append(OxmlElement('w:tblHeader'))
    if height:
        h=OxmlElement('w:trHeight');h.set(qn('w:val'),str(int(height*567)));h.set(qn('w:hRule'),'atLeast');trPr.append(h)
def table(headers,rows,widths,aligns=None,size=10.5,hfill='BDD7EE',band='F2F7FC',after=8,center_first=True):
    t=d.add_table(rows=1+len(rows),cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER
    borders(t);set_cell_margins(t);fixed(t,widths)
    aligns=aligns or [AL.CENTER]+[AL.LEFT]*(len(headers)-1)
    for i,h in enumerate(headers):
        cell_text(t.rows[0].cells[i],h,size,True,AL.CENTER);shade(t.rows[0].cells[i],hfill)
    row_opts(t.rows[0],header=True)
    for r,row in enumerate(rows):
        for i,v in enumerate(row):
            cell_text(t.rows[r+1].cells[i],str(v),size,False,aligns[i])
            if band and r%2==1: shade(t.rows[r+1].cells[i],band)
        row_opts(t.rows[r+1])
    sp=para('',2,after=after,line=1.0)
    return t
FIG=[0]
def fig(path,w,caption,after=8,keep=True):
    p=para('',12,align=AL.CENTER,after=2,keep=True,line=1.0)
    p.add_run().add_picture(path,width=Cm(w))
    FIG[0]+=1
    para(f'Figure {FIG[0]}: {caption}',10,True,align=AL.CENTER,after=after,line=1.0)
def dashed_line():
    p=para('',6,after=8,line=1.0);pPr=p._p.get_or_add_pPr();pb=OxmlElement('w:pBdr');b=OxmlElement('w:bottom')
    for k,v in (('val','dashed'),('sz','8'),('space','1'),('color','000000')): b.set(qn('w:'+k),v)
    pb.append(b);pPr.append(pb)
def letterhead():
    p=para('',12,align=AL.CENTER,after=0,line=1.0);p.add_run().add_picture('letterhead.jpg',width=Cm(15.9));dashed_line()
def pagebreak():
    p=d.add_paragraph();p.add_run().add_break(WD_BREAK.PAGE)

# ================= FRONT MATTER =================
letterhead()
para('DEPARTMENT OF COMPUTER SCIENCE AND ENGINEERING',12,True,align=AL.CENTER,after=22,before=6)
para('ACTIVE LEARNING METHODOLOGY REPORT',12,True,align=AL.CENTER,after=22)
para('LPC2148 BASED MULTI-SENSOR FAULT DETECTION SYSTEM',12,True,align=AL.CENTER,after=0,line=1.0)
para('USING INTERRUPTS, TIMERS AND SAFE SHUTDOWN',12,True,align=AL.CENTER,after=34,line=1.0)
for n in ['BHAVANA S (71812401027)','CHANDRA M (71812401028)','CHRISTINA DAMARIS E (71812401029)','CHRISWIN JANUSZ A (71812401030)','DHAKSHITHA S (71812401031)']:
    para(n,12,False,align=AL.CENTER,after=0,line=1.15)
para('',12,after=24)
para('THIRD YEAR B.E. CSE – V SEM',12,True,align=AL.CENTER,after=8)
para('Academic Year 2026-2027',12,True,align=AL.CENTER,after=30)
para('20EC252 & MICROCONTROLLER',12,True,align=AL.CENTER,after=0,line=1.0)
para('AND PROGRAMMING',12,True,align=AL.CENTER,after=48,line=1.0)
para('COURSE COORDINATOR',12,True,align=AL.RIGHT,after=2)
para('Mrs. Amuthasurabi M, AP/CSE',12,True,align=AL.RIGHT,after=0)
pagebreak()
# page 2: index
letterhead()
para('DEPARTMENT OF COMPUTER SCIENCE AND ENGINEERING',12,True,align=AL.CENTER,after=22,before=6)
para('20EC252 & MICROCONTROLLER',12,True,align=AL.CENTER,after=0,line=1.0)
para('AND PROGRAMMING',12,True,align=AL.CENTER,after=22,line=1.0)
para('INDEX',12,True,align=AL.CENTER,after=14)
names=['BHAVANA S (71812401027)','CHANDRA M (71812401028)','CHRISTINA DAMARIS E (71812401029)','CHRISWIN JANUSZ A (71812401030)','DHAKSHITHA S (71812401031)']
t=d.add_table(rows=6,cols=7);t.alignment=WD_TABLE_ALIGNMENT.CENTER;borders(t,6);set_cell_margins(t);fixed(t,[1.3,6.9,1.9,1.9,1.9,1.9,0.1+0])
# 6 cols used: merge last tiny col out -> rebuild properly
d._body._body.remove(t._tbl)
t=d.add_table(rows=6,cols=6);t.alignment=WD_TABLE_ALIGNMENT.CENTER;borders(t,6);set_cell_margins(t);fixed(t,[1.5,6.8,1.9,1.9,1.9,1.9])
hd=['S.\nNo.','Name & Roll No.','Tool Explanation (2)','Implementation (5)','Presentation (3)','Total\n(10)']
for i,h in enumerate(hd):
    c=t.rows[0].cells[i];cell_text(c,h.replace('\n',' '),11,True,AL.CENTER)
    if i in (2,3,4):
        tcPr=c._tc.get_or_add_tcPr();td=OxmlElement('w:textDirection');td.set(qn('w:val'),'btLr');tcPr.append(td)
row_opts(t.rows[0],height=4.3)
for r,n in enumerate(names):
    cell_text(t.rows[r+1].cells[0],f'{r+1}.',11,True,AL.LEFT);cell_text(t.rows[r+1].cells[1],n,11.5)
    for i in range(2,6): cell_text(t.rows[r+1].cells[i],'',11)
    row_opts(t.rows[r+1],height=1.5)
para('',12,after=40)
para('SIGNATURE OF COURSE COORDINATOR',11.5,True,align=AL.RIGHT,after=2)
para('Mrs. Amuthasurabi M, AP/CSE',12,True,align=AL.RIGHT)
pagebreak()
# page 3: contents
letterhead()
para('',12,after=36)
para('CONTENTS',14,True,align=AL.CENTER,after=14)
TOC=[('INTRODUCTION','intro'),('WORKFLOW DESCRIPTION','workflow'),('IMPLEMENTATION STEPS','steps'),('CODE IMPLEMENTATION','code'),('OUTPUT','output'),('DIAGRAMMATIC REPRESENTATION','diagram'),('CONCLUSION','conclusion')]
t=d.add_table(rows=1+len(TOC),cols=3);t.alignment=WD_TABLE_ALIGNMENT.CENTER;borders(t,6);set_cell_margins(t,70,70,100,100);fixed(t,[2.1,11.5,2.3])
for i,h in enumerate(['S. No.','TITLE','PAGE NUMBER']): cell_text(t.rows[0].cells[i],h,11,True,AL.CENTER)
row_opts(t.rows[0],height=1.1)
for i,(n,k) in enumerate(TOC):
    cell_text(t.rows[i+1].cells[0],f'{i+1}.',12,False,AL.CENTER);cell_text(t.rows[i+1].cells[1],n,12);cell_text(t.rows[i+1].cells[2],str(PAGES.get(k,'#')),12,False,AL.CENTER);row_opts(t.rows[i+1],height=1.1)

# ---- new section (content) ----
sec=d.add_section(WD_SECTION.NEW_PAGE)
sec.left_margin=sec.right_margin=Cm(2.54);sec.top_margin=Cm(2.2);sec.bottom_margin=Cm(2.3);sec.footer_distance=Cm(1.2)
# page numbering restart + borders
sp=sec._sectPr
pg=OxmlElement('w:pgNumType');pg.set(qn('w:start'),'1')
pb=OxmlElement('w:pgBorders');pb.set(qn('w:offsetFrom'),'page')
for k in ('top','left','bottom','right'):
    e=OxmlElement('w:'+k);e.set(qn('w:val'),'double');e.set(qn('w:sz'),'6');e.set(qn('w:space'),'24');e.set(qn('w:color'),'000000');pb.append(e)
pm=sp.find(qn('w:pgMar'));pm.addnext(pb);pm.addnext(pg) if False else None
# order: pgSz, pgMar, paperSrc, pgBorders, lnNumType, pgNumType, cols
pb.addnext(pg)
sec.footer.is_linked_to_previous=False
fp=sec.footer.paragraphs[0];fp.alignment=AL.CENTER
def fld(p,instr):
    r=p.add_run();setfont(r,12)
    for t_,txt in (('begin',None),('instr',instr),('separate',None),('text','1'),('end',None)):
        if t_=='instr':
            e=OxmlElement('w:instrText');e.set(qn('xml:space'),'preserve');e.text=txt;r._r.append(e)
        elif t_=='text':
            e=OxmlElement('w:t');e.text=txt;r._r.append(e)
        else:
            e=OxmlElement('w:fldChar');e.set(qn('w:fldCharType'),t_);r._r.append(e)
fld(fp,' PAGE ')
# first section footer empty, header none
# ================= INTRODUCTION =================
H1('INTRODUCTION',pbb=False)
BT=['LM35','MQ-2','ACS712','flame sensor','emergency-stop push button','external interrupts','Timer0 interrupts every 100 ms','safe shutdown routine']
def bolden(t):
    for b in BT: t=t.replace(b,'**'+b+'**',1) if '**'+b not in t else t
    return t
for t_,x in PM['4'][1:]: body(bolden(x))
# ================= WORKFLOW =================
H1('WORKFLOW DESCRIPTION')
fig('img/../d_workflow.png',15.8,'Workflow of the LPC2148 Based Multi-Sensor Fault Detection System',after=8)
items=PM['5'][1:]+PM['6']
for t_,x in items:
    if t_=='H': H2(x)
    elif x.startswith('Power ON/Reset'): para(x,12,True,after=6)
    else: body(x)
# ================= STEPS =================
H1('IMPLEMENTATION STEPS')
for t_,x in PM['7'][1:]+PM['8']:
    if t_=='H': H3(x)
    else: body11(x)
H2('SUMMARY OF IMPLEMENTATION STEPS',before=12)
table(['Step','Activity','Tool / Menu Used'],[
 ['1','Open Keil µVision 5 and start a new project','Project → New µVision Project'],
 ['2','Save the project in a separate folder','LPC2148_Fault_Detection.uvprojx'],
 ['3','Select the LPC2148 device and add Startup.s','NXP → LPC2000 → LPC2148'],
 ['4','Add a new C source file','Add New Item to Group → main.c'],
 ['5','Write and save the embedded C program','Keil editor (Ctrl + S)'],
 ['6','Enter the crystal frequency (12.000 MHz)','Options for Target → Target'],
 ['7','Build the project and check for errors','Build Target (F7)'],
 ['8','Enable Create HEX File and rebuild','Options for Target → Output'],
 ['9','Draw the circuit and load the HEX file','Proteus ISIS → Program File'],
 ['10','Run the simulation','Play button'],
 ['11','Apply test inputs and observe the outputs','Sensors, switches, LCD, Virtual Terminal'],
 ['12','Verify all test cases','Test case table']],[1.4,8.5,6.0],[AL.CENTER,AL.LEFT,AL.LEFT],10.5)
H2('COMPONENTS USED IN THE PROTEUS SIMULATION',before=6)
table(['S. No.','Component','Value / Part','Quantity','Purpose'],[
 ['1','LPC2148','ARM7TDMI-S','1','Main controller'],
 ['2','Crystal with capacitors','12 MHz, 2 × 33 pF','1','Clock source for the PLL'],
 ['3','LM35','10 mV/°C','1','Temperature sensing'],
 ['4','Potentiometers','1 kΩ (RV1, RV2)','2','Simulate MQ-2 and ACS712 outputs'],
 ['5','Push buttons / logic toggle','10 kΩ pull-ups','3','Emergency stop, reset, flame input'],
 ['6','LM016L LCD','16 × 2, 4-bit mode','1','Status display'],
 ['7','LEDs','Red, yellow, green + 330 Ω','3','Critical, warning, normal indication'],
 ['8','Buzzer','1 kΩ series resistor','1','Audible alarm'],
 ['9','2N2222 + 1N4007 + relay','12 V relay','1','Load switching and flyback protection'],
 ['10','Lamp','12 V','1','Protected load'],
 ['11','Virtual Terminal','9600 baud, 8N1','1','UART fault log']],[1.3,4.0,3.7,1.8,5.1],[AL.CENTER,AL.LEFT,AL.LEFT,AL.CENTER,AL.LEFT],10)
# ================= CODE =================
H1('CODE IMPLEMENTATION')
for t_,x in PM['9'][:0]: pass
body(PM['9'][0][1]);body(PM['9'][1][1])
H2('PIN CONFIGURATION:')
pins=[['1.','UART0 TxD / RxD','P0.0 / P0.1'],['2.','Flame Sensor / EINT1','P0.3'],['3.','Fault Reset Key','P0.4'],['4.','Buzzer','P0.10'],['5.','Red LED (Critical)','P0.11'],['6.','Yellow LED (Warning)','P0.12'],['7.','Green LED (Normal)','P0.13'],['8.','Load Relay','P0.15'],['9.','Emergency Stop / EINT0','P0.16'],['10.','LM35 Temperature / AD0.1','P0.28'],['11.','MQ-2 Gas / AD0.2','P0.29'],['12.','ACS712 Current / AD0.3','P0.30'],['13.','LCD RS','P1.16'],['14.','LCD EN','P1.17'],['15.','LCD D4 – D7','P1.18 – P1.21']]
table(['S. No.','Function','LPC2148 Pin'],pins,[2.0,8.4,5.5],[AL.CENTER,AL.LEFT,AL.LEFT],10.5,hfill='C6D9F1')
H2('THRESHOLD VALUES:',before=6)
table(['Sensor','Warning Level','Critical Level (3 samples)'],[['LM35 Temperature','≥ 50 °C','≥ 70 °C'],['MQ-2 Gas','≥ 600 ADC counts','≥ 800 ADC counts'],['ACS712 Current','≥ 700 ADC counts','≥ 850 ADC counts'],['Flame Sensor','–','Immediate (EINT1)'],['Emergency Stop','–','Immediate (EINT0)']],[5.3,5.3,5.3],[AL.LEFT,AL.LEFT,AL.LEFT],10.5,hfill='C6D9F1')
body(PM['10'][0][1],after=6)
H2('BUILD SUMMARY:',before=6)
table(['Parameter','Value'],[['Target device','NXP LPC2148 (ARM7TDMI-S)'],['Crystal frequency / CCLK','12.000 MHz / 60 MHz'],['Program size','Code = 4168 bytes, RO-data = 388 bytes, RW-data = 64 bytes, ZI-data = 1260 bytes'],['Build result','0 Error(s), 0 Warning(s)'],['Output file','LPC2148_Fault_Detection.hex (HEX-80 format)']],[5.2,10.7],[AL.LEFT,AL.LEFT],10.5)
H2('PROGRAM:',before=4)
code=list(CODE)
code=[l for l in code]
fix=[]
for l in code:
    if l.startswith('const char *fault_text[]'):
        fix+=['const char *fault_text[] = { "NONE", "OVER TEMP", "GAS LEAK",','                            "OVER CURRENT", "FLAME", "E-STOP" };']
    else: fix.append(l)
for l in fix:
    p=d.add_paragraph();pf=p.paragraph_format;pf.space_after=Pt(0);pf.line_spacing=Pt(10.6);pf.left_indent=Cm(0.25);pf.right_indent=Cm(0.1)
    pPr=p._p.get_or_add_pPr();sh=OxmlElement('w:shd');sh.set(qn('w:val'),'clear');sh.set(qn('w:color'),'auto');sh.set(qn('w:fill'),'F3F5F8');pPr.append(sh)
    r=p.add_run(l.replace(' ',' ') if l.strip() else ' ');setfont(r,8,name='Courier New')
H2('FUNCTION DESCRIPTION:',before=12)
table(['Function','Purpose'],[
 ['pll_init()','Raises the 12 MHz crystal clock to CCLK = PCLK = 60 MHz using PLL0'],
 ['delay_ms()','Software delay used by the LCD driver and the start-up sequence'],
 ['lcd_init(), lcd_cmd(), lcd_char(), lcd_string(), lcd_num()','4-bit HD44780 driver on P1.16–P1.21 for the 16×2 LCD'],
 ['uart0_init(), uart0_char(), uart0_string(), uart0_num()','UART0 at 9600 baud, 8N1 for the fault log'],
 ['adc_init(), adc_read()','Configures P0.28–P0.30 as AD0.1–AD0.3 and returns a 10-bit conversion result'],
 ['safe_shutdown()','Relay OFF → red LED and buzzer ON → latch fault code → request UART log'],
 ['TIMER0_ISR()','100 ms tick: sets check_due and disp_due, blinks the warning LED and buzzer'],
 ['EINT0_ISR(), EINT1_ISR()','Emergency stop and flame interrupts: call safe_shutdown() immediately'],
 ['timer0_init(), interrupt_init()','Timer0 match at 100 ms; EINT0/EINT1 falling edge; VIC slots 0, 1 and 2'],
 ['wdt_init(), wdt_feed()','Watchdog set for about 2 s and fed once in every main-loop pass'],
 ['confirm()','Three-sample confirmation filter against false alarms'],
 ['scan_sensors()','Reads LM35, MQ-2 and ACS712, applies the thresholds and updates the state'],
 ['show_status(), log_fault()','LCD status display and UART fault record'],
 ['main()','Initialisation, self test, interrupt set-up and the super-loop']],[6.4,9.5],[AL.LEFT,AL.LEFT],10)
# ================= OUTPUTS =================
H1('OUTPUTS')
def kfig(title,img,cap,desc,first=False,pbb=False):
    H2(title,before=0 if first else 6).paragraph_format.page_break_before=pbb
    fig(img,15.6,cap,after=3)
    para(desc,11,False,italic=True,align=AL.JUSTIFY,after=8,line=1.05)
kfig('CREATING A NEW µVISION PROJECT:','img/k2.jpg','Create New Project dialog (LPC2148_Fault_Detection)','The project is saved as LPC2148_Fault_Detection in its own folder so that the source, object and HEX files stay together.',first=True)
kfig('SELECTING LPC2148 DEVICE:','img/k3.jpg','Select Device for Target – NXP LPC2148','LPC2148 is selected under NXP → LPC2000; the description pane confirms 512 kB flash, 32 kB RAM, UART, ADC and timers.')
kfig('ADDING MAIN.C SOURCE FILE:','img/k4.jpg','Add New Item to Group \'Source Group 1\' – C File (main.c)','A C file named main is added to Source Group 1; the project tree then shows Startup.s and main.c.',pbb=True)
kfig('EMBEDDED C PROGRAM IN KEIL:','img/k5.jpg','µVision editor showing the fault detection program (main.c)','The program with pin definitions, thresholds, system states and fault codes is shown in the editor with syntax colouring.')
kfig('TARGET CONFIGURATION:','img/k6.jpg','Options for Target – Xtal (MHz) = 12.0','The crystal frequency is set to 12.0 MHz; IROM1 (0x0, 0x80000) and IRAM1 (0x40000000, 0x8000) match the LPC2148 memory map.',pbb=True)
kfig('SUCCESSFUL BUILD OUTPUT:','img/k7.jpg','Build Output window – 0 Error(s), 0 Warning(s)','The program compiles and links successfully; the Build Output window reports the program size and zero errors and warnings.')
kfig('HEX FILE GENERATION:','img/k8.jpg','Options for Target → Output with Create HEX File ticked','Create HEX File is enabled with the HEX-80 format so that Keil produces the file loaded by the Proteus model.',pbb=True)
kfig('HEX FILE CREATED AFTER REBUILD:','img/k9.jpg','Build Output after rebuild – hex file created','After rebuilding, Keil converts the .axf file into LPC2148_Fault_Detection.hex with 0 errors and 0 warnings.')
# Proteus
H2('PROTEUS CIRCUIT DIAGRAM:',before=0).paragraph_format.page_break_before=True
fig('img/p_design.jpg',15.8,'ISIS schematic of the LPC2148 based fault detection system (simulation not started)',after=3)
para('The circuit contains the LPC2148 with a 12 MHz crystal, the LM35 and two potentiometers for the analog inputs, the emergency-stop, reset and flame inputs, the status LEDs, buzzer, relay-driven lamp, the LM016L LCD and the Virtual Terminal.',11,False,italic=True,after=8,line=1.05)
H2('LOADING THE HEX FILE IN PROTEUS:',before=4)
fig('img/p_prog.jpg',15.8,'Edit Component dialog of U1 – Program File = LPC2148_Fault_Detection.hex, clock 12 MHz',after=3)
para('The HEX file generated by Keil is selected as the Program File of the LPC2148 and the crystal frequency is set to 12 MHz before the simulation is started.',11,False,italic=True,after=6,line=1.05)
EXTRA={
'POWER-ON SELF TEST:':(['32','210','310','Released','High','Self test'],['The relay output P0.15 stays low until the self test passes, so the load never starts in an unsafe condition.','The LCD, UART0 and ADC are initialised before the first sensor scan.','Timer0, the external interrupts and the watchdog are enabled only after the three self-test scans.']),
'NORMAL CONDITION:':(['32','210','310','Released','High','NORMAL'],['P0.15 and P0.13 are high, so the lamp and the green LED are ON.','The first LCD line shows the live temperature and gas readings; the second line shows STATUS: NORMAL.','The Virtual Terminal has printed “System started: all sensors normal”.']),
'WARNING CONDITION (TEMPERATURE 56 °C):':(['56','240','310','Released','High','WARNING'],['Temperature ≥ 50 °C moves the state machine from NORMAL to WARNING.','Timer0 toggles the yellow LED and the buzzer every 400 ms; the load keeps running.','The system returns to NORMAL when all values fall below the warning limits.']),
'CRITICAL FAULT – OVER TEMPERATURE (74 °C):':(['74','250','330','Released','High','SHUTDOWN'],['The reading is confirmed for three consecutive 100 ms scans before shutdown.','The relay is de-energised first, then the red LED and the buzzer are switched ON.','The fault record is sent over UART0 and the state remains latched until reset.']),
'CRITICAL FAULT – GAS LEAK:':(['33','820','310','Released','High','SHUTDOWN'],['MQ-2 output ≥ 800 counts for three samples is treated as a gas leak.','The lamp (load) is OFF and the red LED and buzzer are ON.','Only the first fault cause is latched in fault_code.']),
'CRITICAL FAULT – OVER CURRENT:':(['34','215','870','Released','High','SHUTDOWN'],['ACS712 output ≥ 850 counts for three samples is treated as over current.','The load is disconnected to protect the equipment and wiring.','The LCD shows the latched cause until the reset key is pressed.']),
'EMERGENCY STOP ACTIVATED:':(['32','210','310','Pressed (low)','High','SHUTDOWN'],['EINT0 has the highest VIC priority (slot 0), so the response is immediate.','The ISR clears EXTINT and writes VICVectAddr = 0 before returning.','The load stays OFF even after the button is released, until the reset key is pressed.']),
'FLAME DETECTED:':(['35','220','310','Released','Low','SHUTDOWN'],['EINT1 is configured for falling-edge detection on P0.3.','safe_shutdown(F_FLAME) is called from EINT1_ISR (VIC slot 1).','After the flame input returns high, the reset key and a healthy re-scan restart the load.']),
}
def scen(title,lcdimg,pimg,lcdcap,pcap,desc,vals,pbb=True,inp=None,pts=None):
    h=H2(title,before=0);h.paragraph_format.page_break_before=pbb
    fig(lcdimg,12.4,lcdcap,after=3)
    fig(pimg,15.8,pcap,after=4)
    table(['Load Lamp','Green LED','Yellow LED','Red LED','Buzzer','LCD – line 2','UART'],[vals],[2.2,2.1,2.3,1.9,2.0,3.3,2.1],[AL.CENTER]*7,9.5,band=None,after=4)
    inp,pts=EXTRA.get(title,(None,None))
    if inp:
        table(['LM35 (°C)','MQ-2 (counts)','ACS712 (counts)','E-Stop (P0.16)','Flame (P0.3)','System State'],[inp],[2.5,2.7,2.9,2.7,2.4,2.7],[AL.CENTER]*6,9.5,band=None,after=4)
    para(desc,11,False,italic=False,align=AL.JUSTIFY,after=4,line=1.05)
    for x in (pts or []): bullet(x,11)
scen('POWER-ON SELF TEST:','img/lcd_boot.png','img/p_boot.jpg','LCD module at start-up','Proteus – self test running, load still OFF','After reset the load relay is held OFF, the LCD shows FAULT DETECTION / SELF TEST... and all sensors are scanned three times.',['OFF','OFF','OFF','OFF','OFF','SELF TEST...','–'],pbb=True)
scen('NORMAL CONDITION:','img/lcd_normal.png','img/p_normal.jpg','LCD module – T:32C G:210 / STATUS: NORMAL','Proteus – load lamp ON, green LED ON','All readings are below their warning levels, so the load runs and the system stays in the NORMAL state.',['ON','ON','OFF','OFF','OFF','STATUS: NORMAL','Started'])
scen('WARNING CONDITION (TEMPERATURE 56 °C):','img/lcd_warn.png','img/p_warn.jpg','LCD module – T:56C G:240 / STATUS: WARNING','Proteus – load ON, yellow LED blinking, buzzer beeping','The temperature is above the 50 °C warning limit but below the critical limit, so the load keeps running while the yellow LED blinks and the buzzer beeps.',['ON','OFF','Blinking','OFF','Beeping','STATUS: WARNING','–'])
scen('CRITICAL FAULT – OVER TEMPERATURE (74 °C):','img/lcd_temp.png','img/p_temp.jpg','LCD module – T:74C G:250 / FLT:OVER TEMP','Proteus – load lamp OFF, red LED and buzzer ON','The temperature stays above 70 °C for three consecutive scans (300 ms), so safe_shutdown() switches the load OFF, lights the red LED and sounds the buzzer.',['OFF','OFF','OFF','ON','Continuous','FLT:OVER TEMP','Fault log'])
scen('CRITICAL FAULT – GAS LEAK:','img/lcd_gas.png','img/p_gas.jpg','LCD module – T:33C G:820 / FLT:GAS LEAK','Proteus – load OFF after the gas potentiometer exceeds 800 counts','The MQ-2 output crosses 800 ADC counts for three samples, so the load is cut off and the cause is latched as GAS LEAK.',['OFF','OFF','OFF','ON','Continuous','FLT:GAS LEAK','Fault log'])
scen('CRITICAL FAULT – OVER CURRENT:','img/lcd_curr.png','img/p_curr.jpg','LCD module – T:34C G:215 / FLT:OVER CURRENT','Proteus – load OFF after the current potentiometer exceeds 850 counts','The ACS712 output crosses 850 ADC counts, so the load is switched OFF and OVER CURRENT is latched.',['OFF','OFF','OFF','ON','Continuous','FLT:OVER CURRENT','Fault log'])
scen('EMERGENCY STOP ACTIVATED:','img/lcd_estop.png','img/p_estop.jpg','LCD module – FLT:E-STOP','Proteus – E-stop pressed, load OFF immediately, LCD shows FLT:E-STOP','Pressing the emergency-stop button generates a falling edge on EINT0; the interrupt service routine calls safe_shutdown() at once.',['OFF','OFF','OFF','ON','Continuous','FLT:E-STOP','Fault log'])
scen('FLAME DETECTED:','img/lcd_flame.png','img/p_flame.jpg','LCD module – FLT:FLAME','Proteus – flame input pulled low, load OFF immediately','When the flame input is pulled low, EINT1 interrupts the main program and the load is switched OFF without waiting for the periodic scan.',['OFF','OFF','OFF','ON','Continuous','FLT:FLAME','Fault log'])
H2('EXPECTED VIRTUAL TERMINAL OUTPUT:',before=0).paragraph_format.page_break_before=True
fig('img/vt_full.png',15.4,'Virtual Terminal (9600 baud, 8N1) – fault log sent over UART0',after=4)
para('The terminal first shows the start-up message after a successful self test, then the fault record with the cause, temperature, gas value and current value when a shutdown occurs, and finally the message printed after the reset key is pressed and the re-scan is healthy.',11,False,after=8,line=1.05)
H2('TEST CASES:',before=6)
table(['No.','Input Condition','Expected Output','Result'],[
 ['1','All sensors normal (T = 32 °C, gas = 210)','Load ON, green LED, “STATUS: NORMAL”','Pass'],
 ['2','Temperature raised to 56 °C','Load ON, yellow LED blinks, intermittent beep, “STATUS: WARNING”','Pass'],
 ['3','Temperature 74 °C for more than 300 ms','Load OFF, red LED, continuous buzzer, “FLT:OVER TEMP”, UART log','Pass'],
 ['4','Gas value above 800 counts','Load OFF, “FLT:GAS LEAK”','Pass'],
 ['5','Current value above 850 counts','Load OFF, “FLT:OVER CURRENT”','Pass'],
 ['6','Emergency stop pressed','Immediate load OFF via EINT0, “FLT:E-STOP”','Pass'],
 ['7','Flame input pulled low','Immediate load OFF via EINT1, “FLT:FLAME”','Pass'],
 ['8','Reset pressed while fault still present','System stays in SHUTDOWN','Pass'],
 ['9','Reset pressed after values return to normal','Fault cleared, load restarted','Pass']],[1.2,5.6,7.2,1.9],[AL.CENTER,AL.LEFT,AL.LEFT,AL.CENTER],10.5,hfill='C6D9F1')
H2('RESULT:',before=6)
body11('All nine test cases give the expected behaviour in the Proteus simulation. The load, LEDs, buzzer, LCD and Virtual Terminal respond correctly to normal operation, warning conditions, sensor faults, the emergency stop and the flame input.',first=1.0)
# ================= DIAGRAMS =================
H1('DIAGRAMMATIC REPRESENTATION')
H2('BLOCK DIAGRAM:',before=0)
fig('d_block.png',13.8,'Block diagram of the LPC2148 based multi-sensor fault detection system',after=6)
for t_,x in PM['22'][2:]: body(x)
H2('SIGNAL INTERFACING TABLE:',before=6)
table(['Block','Signal','LPC2148 Pin','Interface'],[
 ['LM35 temperature sensor','Analog, 10 mV/°C','P0.28 / AD0.1','ADC0 channel 1'],
 ['MQ-2 gas / smoke sensor','Analog','P0.29 / AD0.2','ADC0 channel 2'],
 ['ACS712 current sensor','Analog','P0.30 / AD0.3','ADC0 channel 3'],
 ['Flame sensor','Digital, active low','P0.3 / EINT1','External interrupt'],
 ['Emergency stop button','Digital, active low','P0.16 / EINT0','External interrupt'],
 ['Relay driver, buzzer, LEDs','Digital outputs','P0.15, P0.10, P0.11–P0.13','GPIO'],
 ['16×2 LCD','4-bit data + RS, EN','P1.16–P1.21','GPIO'],
 ['PC / Virtual Terminal','UART0, 9600 baud','P0.0 / P0.1','Serial']],[4.6,3.6,4.0,3.7],[AL.LEFT,AL.LEFT,AL.LEFT,AL.LEFT],10)
H2('SOFTWARE STRUCTURE:',before=0).paragraph_format.page_break_before=True
fig('d_software.png',9.3,'Layered software structure of the fault detection firmware',after=6)
body(PM['23'][1][1])
table(['Layer','Contents','Role'],[
 ['1 – Interrupt','TIMER0_ISR, EINT0_ISR, EINT1_ISR','Short handlers that react or set flags'],
 ['2 – Application','scan_sensors(), state machine, show_status(), log_fault()','Main super-loop decisions and outputs'],
 ['3 – Safety','safe_shutdown(), watchdog timer','Always drives the system to a safe state'],
 ['4 – Driver','pll_init(), adc_*(), lcd_*(), uart0_*(), timer0_init(), interrupt_init()','Low-level hardware access']],[3.2,7.6,5.1],[AL.LEFT,AL.LEFT,AL.LEFT],10)
H2('INTERRUPT FLOW:',before=0).paragraph_format.page_break_before=True
fig('d_interrupt.png',15.4,'Interrupt flow in the LPC2148 based fault detection system',after=6)
body(PM['24'][1][1])
H2('INTERRUPT SEQUENCE:',before=4)
para(PM['24'][3][1],12,True,after=6)
H2('VIC INTERRUPT ASSIGNMENT:',before=6)
table(['VIC Slot','Source','Channel','Handler','Trigger'],[
 ['0 (highest)','EINT0 – emergency stop','14','EINT0_ISR','Falling edge on P0.16'],
 ['1','EINT1 – flame sensor','15','EINT1_ISR','Falling edge on P0.3'],
 ['2','Timer0 match MR0','4','TIMER0_ISR','Every 100 ms']],[2.6,4.6,2.0,3.0,3.7],[AL.CENTER,AL.LEFT,AL.CENTER,AL.LEFT,AL.LEFT],10)
H2('STATE MACHINE:',before=0).paragraph_format.page_break_before=True
fig('d_state.png',15.4,'System state machine (NORMAL, WARNING, SHUTDOWN)',after=6)
body(PM['25'][1][1])
H2('STATE TRANSITION TABLE:',before=6)
table(['Present State','Condition','Next State','Outputs'],[
 ['Power ON','Self test OK','NORMAL','Load ON, green LED ON'],
 ['NORMAL','Any value ≥ warning limit','WARNING','Yellow LED blinks, buzzer beeps'],
 ['WARNING','All values below warning limit','NORMAL','Yellow LED and buzzer OFF, green LED ON'],
 ['WARNING','Critical value for 3 samples','SHUTDOWN','Load OFF, red LED and buzzer ON'],
 ['NORMAL / WARNING','E-Stop or flame interrupt','SHUTDOWN','Load OFF immediately, fault latched'],
 ['SHUTDOWN','Reset key pressed and healthy re-scan','NORMAL','Fault cleared, load restarted']],[3.4,5.3,2.6,4.6],[AL.LEFT,AL.LEFT,AL.CENTER,AL.LEFT],10)
H2('SYSTEM WORKFLOW:',before=0).paragraph_format.page_break_before=True
fig('d_flow.png',11.2,'Flowchart of the LPC2148 based fault detection system',after=2)
# ================= CONCLUSION =================
H1('CONCLUSION')
for t_,x in PM['27'][1:]: para(x,11,False,first=1.0,after=3,line=1.05)
H2('APPLICATIONS:',before=6)
for a in ['Motor and machine protection in industrial plants','Industrial electrical safety panels','Server-room and laboratory equipment monitoring','Kitchen, boiler and furnace safety systems']: bullet(a)
H2('FUTURE ENHANCEMENTS:',before=6)
for a in ['Storing fault history in the on-chip flash or an external EEPROM','Sending SMS alerts through a GSM module','Adding vibration and voltage sensors','Connecting the system to an IoT dashboard for remote monitoring']: bullet(a)
H2('REFERENCES:',before=6)
for a in ['NXP Semiconductors, “UM10139 – LPC214x User Manual”.','Keil / Arm, “µVision User’s Guide and MDK-ARM Compiler Documentation”.','Labcenter Electronics, “Proteus VSM ISIS Simulation Guide”.','Texas Instruments, “LM35 Precision Centigrade Temperature Sensors” data sheet.']: bullet(a)
d.core_properties.title='LPC2148 Based Multi-Sensor Fault Detection System';d.core_properties.author=''
out='MCP_ALM_Fault_Detection_System_FINAL.docx' if PAGES else 'tmp_fd.docx'
d.save(out);print('saved',out)
