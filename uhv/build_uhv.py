import re,sys,json
from docx import Document
from docx.shared import Pt,Cm,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH as AL, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from pygments import lex
from pygments.lexers import HtmlLexer,JavascriptLexer,CssLexer,BashLexer
from pygments.lexers.javascript import JavascriptLexer
from pygments.token import Token
try:
    from pygments.lexers import JsxLexer
except Exception:
    from pygments.lexers.jsx import JsxLexer
PAGES=json.load(open('pm.json')) if len(sys.argv)>1 else {}
BODY='Calibri';MONO='Consolas';W=9.2
NAVY='1F3A5F';BLUE='2E6DB4'
d=Document()
st=d.styles['Normal'];st.font.name=BODY;st.font.size=Pt(8.3)
st.element.rPr.rFonts.set(qn('w:eastAsia'),BODY)
st.paragraph_format.space_after=Pt(0)
s=d.sections[0]
s.page_width=Cm(21);s.page_height=Cm(29.7);s.left_margin=s.right_margin=Cm(1.1);s.top_margin=Cm(1.2);s.bottom_margin=Cm(1.4);s.footer_distance=Cm(0.45)
sp=s._sectPr
pb=OxmlElement('w:pgBorders');pb.set(qn('w:offsetFrom'),'page')
for k in ('top','left','bottom','right'):
    e=OxmlElement('w:'+k);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'10');e.set(qn('w:space'),'14');e.set(qn('w:color'),'000000');pb.append(e)
sp.find(qn('w:pgMar')).addnext(pb)
def fnt(r,size=None,bold=None,italic=None,color=None,name=None):
    nm=name or BODY;r.font.name=nm
    rf=r._r.get_or_add_rPr().find(qn('w:rFonts'))
    for a in ('w:ascii','w:hAnsi','w:cs','w:eastAsia'): rf.set(qn(a),nm)
    if size: r.font.size=Pt(size)
    if bold is not None: r.bold=bold
    if italic is not None: r.italic=italic
    if color: r.font.color.rgb=RGBColor.from_string(color)
def shade_p(p,fill):
    pPr=p._p.get_or_add_pPr();sh=OxmlElement('w:shd');sh.set(qn('w:val'),'clear');sh.set(qn('w:color'),'auto');sh.set(qn('w:fill'),fill);pPr.append(sh)
def shade_c(c,fill):
    tcPr=c._tc.get_or_add_tcPr();sh=OxmlElement('w:shd');sh.set(qn('w:val'),'clear');sh.set(qn('w:color'),'auto');sh.set(qn('w:fill'),fill);tcPr.append(sh)
def cell_borders(c,**kw):
    tcPr=c._tc.get_or_add_tcPr();b=OxmlElement('w:tcBorders')
    for k in ('top','left','bottom','right'):
        v=kw.get(k)
        e=OxmlElement('w:'+k)
        if v: e.set(qn('w:val'),'single');e.set(qn('w:sz'),str(v[0]));e.set(qn('w:space'),'0');e.set(qn('w:color'),v[1])
        else: e.set(qn('w:val'),'nil')
        b.append(e)
    tcPr.append(b)
def cell_mar(c,top=60,bottom=60,left=140,right=140):
    tcPr=c._tc.get_or_add_tcPr();m=OxmlElement('w:tcMar')
    for k,v in (('top',top),('left',left),('bottom',bottom),('right',right)):
        e=OxmlElement('w:'+k);e.set(qn('w:w'),str(v));e.set(qn('w:type'),'dxa');m.append(e)
    tcPr.append(m)
def tbl_fixed(t,widths):
    t.autofit=False;pr=t._tbl.tblPr;l=OxmlElement('w:tblLayout');l.set(qn('w:type'),'fixed');pr.append(l)
    for row in t.rows:
        for i,c in enumerate(row.cells): c.width=Cm(widths[i])
    for i,g in enumerate(t._tbl.tblGrid.findall(qn('w:gridCol'))): g.set(qn('w:w'),str(int(widths[i]*567)))
def no_split(row):
    row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
EM={'❌':'✗','✅':'✔','↙️':'','↘️':''}
def clean(t):
    for k,v in EM.items(): t=t.replace(k,v)
    return t
def runs(p,text,size=8.3,bold=False,italic=False,color=None):
    text=clean(text)
    parts=re.split(r'(\*\*.+?\*\*|`[^`]+`|⭐+)',text)
    for x in parts:
        if not x: continue
        if x.startswith('**'): fnt(p.add_run(x[2:-2]),size,True,italic,color)
        elif x.startswith('`'):
            r=p.add_run(' '+x[1:-1]+' ');fnt(r,size-0.8,False,False,'B03060',MONO)
            rp=r._r.get_or_add_rPr();sh=OxmlElement('w:shd');sh.set(qn('w:val'),'clear');sh.set(qn('w:color'),'auto');sh.set(qn('w:fill'),'F1F3F5');rp.append(sh)
        elif x.startswith('⭐'): fnt(p.add_run('★'*len(x)),size,False,False,'E0A800',name='DejaVu Sans')
        else: fnt(p.add_run(x),size,bold,italic,color)
def para(text='',size=8.3,bold=False,italic=False,align=AL.LEFT,after=2,before=0,keep=False,left=0,color=None,line=1.0):
    p=d.add_paragraph();p.alignment=align;pf=p.paragraph_format;pf.space_after=Pt(after);pf.space_before=Pt(before);pf.keep_with_next=keep;pf.line_spacing=line
    if left: pf.left_indent=Cm(left)
    if text: runs(p,text,size,bold,italic,color)
    return p
def pbreak(): d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
def bar(text,fill=NAVY,size=15,color='FFFFFF',after=8,pbb=False):
    size=size*0.62;after=3;pbb=False
    t=d.add_table(rows=1,cols=1);t.alignment=WD_TABLE_ALIGNMENT.CENTER;tbl_fixed(t,[W]);c=t.rows[0].cells[0];shade_c(c,fill);cell_mar(c,45,45,110,110)
    cell_borders(c);p=c.paragraphs[0];p.paragraph_format.keep_with_next=True
    if pbb: p.paragraph_format.page_break_before=True
    runs(p,text,size,True,False,color);no_split(t.rows[0])
    sp=d.add_paragraph();sp.paragraph_format.space_after=Pt(0);sp.paragraph_format.line_spacing=Pt(3);sp.paragraph_format.keep_with_next=True
def h1(t):
    p=para(t,10.2,True,after=3,before=6,keep=True,color=NAVY)
    pPr=p._p.get_or_add_pPr();b=OxmlElement('w:pBdr');e=OxmlElement('w:bottom')
    for k,v in (('val','single'),('sz','8'),('space','2'),('color',BLUE)): e.set(qn('w:'+k),v)
    b.append(e);pPr.append(b);return p
def h2(t): return para(t,9.3,True,after=2,before=5,keep=True,color=BLUE)
def h3(t): return para(t,8.7,True,after=1.5,before=3,keep=True,color='333333')
def bullet(t,lvl=0):
    p=para('',8.3,after=1,left=0.5+lvl*0.4,line=1.0);p.paragraph_format.first_line_indent=Cm(-0.32)
    fnt(p.add_run('•\u00A0'),8.3,True,color=BLUE);runs(p,t,8.3);return p
def numitem(n,t):
    p=para('',8.3,after=1,left=0.55,line=1.0);p.paragraph_format.first_line_indent=Cm(-0.4)
    fnt(p.add_run(f'{n}.\u00A0'),8.3,True,color=BLUE);runs(p,t,8.3);return p
def quote(lines,green=False):
    t=d.add_table(rows=1,cols=1);t.alignment=WD_TABLE_ALIGNMENT.CENTER;tbl_fixed(t,[W]);c=t.rows[0].cells[0]
    fill='E8F5EC' if green else 'EAF2FC';edge='2E9B57' if green else BLUE
    shade_c(c,fill);cell_mar(c,40,40,100,90);cell_borders(c,left=(30,edge))
    first=True
    for ln in lines:
        p=c.paragraphs[0] if first else c.add_paragraph();first=False
        p.paragraph_format.space_after=Pt(1);p.paragraph_format.line_spacing=1.0
        if ln=='': continue
        runs(p,ln,8.3,False,False,'1B1B1B')
    no_split(t.rows[0]);sp=d.add_paragraph();sp.paragraph_format.space_after=Pt(0);sp.paragraph_format.line_spacing=Pt(3)
def mdtable(rows):
    hdr=rows[0];body=rows[2:]
    n=len(hdr)
    lens=[max(len(clean(r[i])) if i<len(r) else 0 for r in [hdr]+body) for i in range(n)]
    tot=sum(max(l,6) for l in lens);ws=[max(1.3,W*max(l,6)/tot) for l in lens];k=W/sum(ws);ws=[w*k for w in ws]
    t=d.add_table(rows=1+len(body),cols=n);t.alignment=WD_TABLE_ALIGNMENT.CENTER;tbl_fixed(t,ws)
    pr=t._tbl.tblPr;b=OxmlElement('w:tblBorders')
    for kk in ('top','left','bottom','right','insideH','insideV'):
        e=OxmlElement('w:'+kk);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'6');e.set(qn('w:space'),'0');e.set(qn('w:color'),'8FA9C8');b.append(e)
    pr.append(b)
    for i,h in enumerate(hdr):
        c=t.rows[0].cells[i];shade_c(c,NAVY);cell_mar(c,30,30,70,70);p=c.paragraphs[0];p.alignment=AL.CENTER;runs(p,h,7.8,True,False,'FFFFFF')
    t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    for r,row in enumerate(body):
        no_split(t.rows[r+1])
        for i in range(n):
            c=t.rows[r+1].cells[i];cell_mar(c,25,25,70,70)
            if r%2==1: shade_c(c,'EEF3FA')
            p=c.paragraphs[0];p.paragraph_format.line_spacing=1.0;runs(p,row[i] if i<len(row) else '',7.8)
    sp=d.add_paragraph();sp.paragraph_format.space_after=Pt(0);sp.paragraph_format.line_spacing=Pt(3)
# ---------- code
BG='F3F3F3';FG='1A1A1A'
COL={Token.Keyword:'9B2D8B',Token.Name.Tag:'2F4B8F',Token.Name.Attribute:'7A5A12',Token.Literal.String:'2E7D32',Token.Literal.Number:'9C4A6B',
Token.Comment:'7A7F87',Token.Name.Function:'6A3FB5',Token.Name.Class:'6A3FB5',Token.Name.Builtin:'6A3FB5',Token.Operator.Word:'9B2D8B',Token.Name.Decorator:'6A3FB5',Token.Name.Namespace:'6A3FB5',Token.Name.Constant:'9C4A6B'}
def colour(tt):
    while tt is not None:
        if tt in COL: return COL[tt]
        tt=tt.parent
    return FG
LEX={'html':HtmlLexer(),'javascript':JavascriptLexer(),'jsx':JsxLexer(),'css':CssLexer(),'bash':BashLexer()}
LABEL={'html':'HTML','javascript':'JavaScript','jsx':'JavaScript (JSX)','css':'CSS','bash':'Bash','text':''}
def codeblock(lang,lines):
    while lines and not lines[-1].strip(): lines.pop()
    while lines and not lines[0].strip(): lines.pop(0)
    plain=(lang=='text')
    t=d.add_table(rows=(1 if plain else 2),cols=1);t.alignment=WD_TABLE_ALIGNMENT.CENTER;tbl_fixed(t,[W])
    ri=0
    if not plain:
        c=t.rows[0].cells[0];shade_c(c,BG);cell_mar(c,30,10,110,80);cell_borders(c)
        p=c.paragraphs[0];p.paragraph_format.keep_with_next=True
        fnt(p.add_run('</>  '),6.8,True,color='222222',name=MONO);fnt(p.add_run(LABEL[lang]),7.6,True,color='111111');no_split(t.rows[0])
    c=t.rows[ri+(0 if plain else 1)].cells[0];shade_c(c,BG);cell_mar(c,20,50,110,80);cell_borders(c)
    if len(lines)<=28: no_split(t.rows[-1])
    comp=[]
    for x in lines:
        if not x.strip() and comp and not comp[-1].strip(): continue
        comp.append(x)
    lines=comp
    src='\n'.join(lines)
    toks=list(lex(src,LEX[lang])) if not plain else [(Token.Text,src)]
    # split tokens into lines
    ln=[[]]
    for tt,val in toks:
        parts=val.split('\n')
        for i,pt in enumerate(parts):
            if i>0: ln.append([])
            if pt:
                c_=colour(tt)
                if c_==FG and tt in Token.Name and re.match(r'^([A-Z][A-Za-z]+|use[A-Z]\w+)$',pt) and lang in ('jsx','javascript'): c_='6A3FB5'
                ln[-1].append((c_,pt))
    if ln and not ln[-1]: ln.pop()
    first=True
    for L in ln:
        p=c.paragraphs[0] if first else c.add_paragraph();first=False
        pf=p.paragraph_format;pf.space_after=Pt(0);pf.line_spacing=Pt(7.6)
        if not L:
            pf.line_spacing=Pt(3.4);r=p.add_run('\u00A0');fnt(r,3,color=FG,name=MONO);continue
        # leading whitespace -> nbsp
        for col,txt in L:
            txt=txt.replace(' ',' ').replace('\t',' '*4)
            fnt(p.add_run(txt),6.3,False,color=col,name=MONO)
    sp=d.add_paragraph();sp.paragraph_format.space_after=Pt(0);sp.paragraph_format.line_spacing=Pt(3)
# ---------- footer
s.footer.is_linked_to_previous=False
fp=s.footer.paragraphs[0];fp.alignment=AL.CENTER
def fld(p,instr):
    r=p.add_run();fnt(r,10,color='555555')
    for k,tx in (('begin',None),('instr',instr),('separate',None),('text','1'),('end',None)):
        if k=='instr': e=OxmlElement('w:instrText');e.set(qn('xml:space'),'preserve');e.text=tx;r._r.append(e)
        elif k=='text': e=OxmlElement('w:t');e.text=tx;r._r.append(e)
        else: e=OxmlElement('w:fldChar');e.set(qn('w:fldCharType'),k);r._r.append(e)
fld(fp,' PAGE ')

import os
SPECS=[];DN=[0]
def diagram(spec,width=9.0):
    DN[0]+=1;spec['id']='d%d'%DN[0];SPECS.append(spec)
    f='dg/%s.png'%spec['id']
    if os.path.exists(f):
        p=para('',8,align=AL.CENTER,after=3,line=1.0,keep=False);p.add_run().add_picture(f,width=Cm(width))
    else:
        para('[diagram %s]'%spec['id'],8)
def clean_node(t):
    t=t.strip();t=re.sub(r'\s{2,}$','',t)
    t=t.replace('**','').strip()
    t=re.sub(r'^[→├└─\s]+','',t)
    t=re.sub(r'^[├└]──\s*','',t)
    return t.strip()
# ---------- cover + contents
para('UNIVERSAL HUMAN VALUES (UHV)',24,True,AL.CENTER,after=2,color=NAVY,line=1.0)
para('EXAM STUDY MATERIAL',15,True,AL.CENTER,after=4,color=BLUE,line=1.0)
para('10-Mark Topics  •  2-Mark Topics  •  Visual flow diagrams for quick learning',10,False,AL.CENTER,after=14,color='555555')
TOC10=[('1','Globalisation','g'),('2','Computer Ethics','c'),('3','Safety and Risk','s'),('4','Engineering as Experimentation','e')]
TOC2=[('1','Definitiveness of Ethical Human Conduct','d1'),('2','Competence in Professional Ethics','d2'),('3','Role and Responsibilities of Engineers','d3'),('4','Professional Societies and their Codes of Ethics','d4'),('5','Constraints in Engineering','d5')]
def toc(title,items):
    t=d.add_table(rows=1,cols=1);t.alignment=WD_TABLE_ALIGNMENT.CENTER;tbl_fixed(t,[17.6]);c=t.rows[0].cells[0];shade_c(c,NAVY);cell_mar(c,60,60,140,140);cell_borders(c)
    runs(c.paragraphs[0],title,11,True,False,'FFFFFF')
    t=d.add_table(rows=len(items),cols=3);t.alignment=WD_TABLE_ALIGNMENT.CENTER;tbl_fixed(t,[1.4,14.0,2.2])
    for i,(n,tx,k) in enumerate(items):
        for j,v in enumerate([n,tx,str(PAGES.get(k,'#'))]):
            c=t.rows[i].cells[j];cell_mar(c,45,45,100,100);cell_borders(c,bottom=(4,'C9D3E0'))
            p=c.paragraphs[0];p.alignment=AL.CENTER if j!=1 else AL.LEFT;runs(p,v,10.5,j==0,False,NAVY if j==0 else None)
    q=d.add_paragraph();q.paragraph_format.space_after=Pt(6)
toc('PART A — 10-MARK QUESTIONS',TOC10)
toc('PART B — 2-MARK QUESTIONS',TOC2)
from docx.enum.section import WD_SECTION
sec2=d.add_section(WD_SECTION.CONTINUOUS)
cols=sec2._sectPr.find(qn('w:cols'))
if cols is None:
    cols=OxmlElement('w:cols');sec2._sectPr.append(cols)
cols.set(qn('w:num'),'2');cols.set(qn('w:space'),'340');cols.set(qn('w:sep'),'1')
OVER={'GLOBALISATION':('Globalisation',['Benefits of Globalisation','Transfer of Technology','Relative Values','International Rights','Morally Justifiable Measures','Business Process Outsourcing']),
'COMPUTER ETHICS':('Computer Ethics',['Computer Crimes','Categories of Computer Crimes','Commandments of Computer Ethics','Ethical Principles for Software','Mobile Etiquette Guidelines','Ethical Hackers']),
'SAFETY AND RISK':('Safety and Risk',['Underestimation of Risk','Overestimation of Risk','Designing for Safety']),
'ENGINEERING AS EXPERIMENTATION':('Engineering as Experimentation',['Similarities to Standard Experiments','Learning from the Past','Knowledge Gained','Framing the Problem'])}
# ---------- parse
L=open('notes.txt',encoding='utf-8').read().split('\n')
i=0;part=None;last_head=''
def blockat(i):
    j=i;buf=[]
    while j<len(L) and L[j].strip(): buf.append(L[j]);j+=1
    return buf,j
while i<len(L):
    ln=L[i];s_=ln.rstrip()
    if s_ in ('UHV 10 MARKS','UHV 2 MARKS'):
        part='A' if '10' in s_ else 'B'
        bar('PART A — 10-MARK QUESTIONS' if part=='A' else 'PART B — 2-MARK QUESTIONS',NAVY,20);i+=1;continue
    if not s_.strip() or s_.strip()=='---': i+=1;continue
    # special: fork
    if s_.strip()=='**Risk estimation**':
        diagram({'type':'fork','root':'Risk estimation','left':{'title':'Too Low','items':['Underestimation','Unsafe decisions','Higher possibility of accidents']},'right':{'title':'Too High','items':['Overestimation','Excessive cost / restrictions','Inefficient decisions']}})
        while i<len(L) and not L[i].startswith('Therefore, engineers should aim'): i+=1
        continue
    if s_.strip()=='**Computer Crimes**' and i+1<len(L) and L[i+1].strip().startswith('↓') and L[i+2].strip().startswith('├'):
        buf,j=blockat(i);items=[clean_node(x) for x in buf[2:]]
        diagram({'type':'tree','root':'Computer Crimes','items':items});i=j;continue
    if s_.startswith('```'):
        i+=1
        while i<len(L) and not L[i].startswith('```'): i+=1
        i+=1;continue
    if s_.startswith('|'):
        rows=[]
        while i<len(L) and L[i].startswith('|'):
            rows.append([x.strip() for x in L[i].strip().strip('|').split('|')]);i+=1
        mdtable(rows);continue
    if s_.startswith('>'):
        qs=[]
        while i<len(L) and L[i].startswith('>'):
            qs.append(L[i][1:].strip());i+=1
        quote([clean(x) for x in qs],green=True);continue
    m=re.match(r'^(#{1,3}) (.*)',s_)
    if m:
        lv=len(m.group(1));tx=m.group(2).strip();last_head=tx
        if lv==1 and re.match(r'^\d\)',tx):
            bar(clean(tx),BLUE,16)
            key=re.sub(r'^\d\)\s*','',tx);key=re.split(r'\s+—',key)[0]
            if part=='A' and key in OVER:
                nm,its=OVER[key];diagram({'type':'hub','title':'Topic map','center':nm,'items':its})
        elif lv==1: h1(clean(tx))
        elif lv==2: h2(clean(tx))
        else: h3(clean(tx))
        i+=1;continue
    # flow blocks
    buf,j=blockat(i)
    sb=[x.strip() for x in buf]
    if any(x.startswith('↓') for x in sb) and not any(x.startswith('-') for x in sb):
        nodes=[];pend=None;k=0
        raw=[x for x in sb]
        cur=[]
        for x in raw:
            if x.startswith('↓'): 
                nodes.append(' '.join(cur));cur=[]
            elif x=='+': cur.append('+')
            else: cur.append(clean_node(x))
        if cur: nodes.append(' '.join(cur))
        nodes=[re.sub(r'\s*\+\s*',' + ',n).strip() for n in nodes]
        if len(nodes)>=2: diagram({'type':'chain','nodes':nodes});i=j;continue
    if len(sb)>=2 and sb[0].startswith('**') and all(x.startswith('→') for x in sb[1:]):
        if sb[0].strip('*')=='Globalisation':
            diagram({'type':'hub','center':'Globalisation','items':[clean_node(x) for x in sb[1:]]})
        else:
            diagram({'type':'chain','nodes':[clean_node(sb[0])]+[clean_node(x) for x in sb[1:]]})
        i=j;continue
    mb=re.match(r'^(\s*)[-*] (.*)',s_)
    if mb: bullet(mb.group(2),1 if len(mb.group(1))>=2 else 0);i+=1;continue
    mn=re.match(r'^(\d+)\. (.*)',s_)
    if mn: numitem(mn.group(1),mn.group(2));i+=1;continue
    # paragraph
    buf=[ln.rstrip('\n')];i+=1
    while i<len(L) and L[i].strip() and not re.match(r'^(#|```|\||>|---|[-*] |\d+\. |\*\*Risk estimation)',L[i]) and not L[i].strip().startswith('↓') and not L[i].strip().startswith('→'):
        buf.append(L[i]);i+=1
    txt=''
    for k,x in enumerate(buf):
        txt+=x.strip()
        if k<len(buf)-1: txt+=('\n' if x.endswith('  ') else ' ')
    para(txt,8.3,after=2)
json.dump(SPECS,open('specs.json','w'))
d.core_properties.title='UHV Exam Study Material';d.core_properties.author=''
out='UHV_Study_Material_Compact.docx' if PAGES else 'tmp_uhv.docx'
d.save(out);print('saved',out,len(SPECS),'diagrams')
