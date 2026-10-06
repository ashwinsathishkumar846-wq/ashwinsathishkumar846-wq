import sys,docx
from docx.shared import Pt,Emu,RGBColor,Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH as AL, WD_TAB_ALIGNMENT as TA
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
BODY=float(sys.argv[1]) if len(sys.argv)>1 else 10.0
GAP=float(sys.argv[2]) if len(sys.argv)>2 else 1.5
F='Calibri'
d=docx.Document()
s=d.sections[0]; s.page_width=Emu(7569200); s.page_height=Emu(10706100)
s.left_margin=s.right_margin=Pt(30); s.top_margin=Pt(18); s.bottom_margin=Pt(10)
st=d.styles['Normal']; st.font.name=F; st.font.size=Pt(BODY); st.element.rPr.rFonts.set(qn('w:eastAsia'),F)
st.paragraph_format.space_after=Pt(0); st.paragraph_format.space_before=Pt(0); st.paragraph_format.line_spacing=1.04
W=Pt(595.5-60)
def run(p,t,b=False,i=False,sz=None,link=None,color=None):
    if link:
        rid=d.part.relate_to(link,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
        h=OxmlElement('w:hyperlink'); h.set(qn('r:id'),rid)
        r=OxmlElement('w:r'); rp=OxmlElement('w:rPr')
        rf=OxmlElement('w:rFonts'); rf.set(qn('w:ascii'),F); rf.set(qn('w:hAnsi'),F); rp.append(rf)
        if b: rp.append(OxmlElement('w:b'))
        if i: rp.append(OxmlElement('w:i'))
        c=OxmlElement('w:color'); c.set(qn('w:val'),'1F4E9C'); rp.append(c)
        z=OxmlElement('w:sz'); z.set(qn('w:val'),str(int((sz or BODY)*2))); rp.append(z)
        u=OxmlElement('w:u'); u.set(qn('w:val'),'single'); rp.append(u)
        r.append(rp); tt=OxmlElement('w:t'); tt.text=t; tt.set(qn('xml:space'),'preserve'); r.append(tt); h.append(r); p._p.append(h); return
    r=p.add_run(t); r.bold=b; r.italic=i; r.font.name=F
    if sz: r.font.size=Pt(sz)
    return r
def P(before=0,align=None,after=0):
    p=d.add_paragraph(); p.paragraph_format.space_before=Pt(before); p.paragraph_format.space_after=Pt(after)
    if align is not None: p.alignment=align
    return p
def H(t):
    p=P(before=6,after=2); p.paragraph_format.keep_with_next=True
    run(p,t,b=True,sz=12.6)
    pPr=p._p.get_or_add_pPr(); bd=OxmlElement('w:pBdr'); b=OxmlElement('w:bottom')
    for k,v in dict(val='single',sz='10',space='1',color='000000').items(): b.set(qn('w:'+k),v)
    bd.append(b); pPr.append(bd)
def row(left,right,sz=None,bl=True,br=True,before=0):
    p=P(before=before); p.paragraph_format.tab_stops.add_tab_stop(W,TA.RIGHT)
    for t,b in left: run(p,t,b=b,sz=sz)
    run(p,'\t'); 
    for t,b in right: run(p,t,b=b,sz=sz)
    return p
def B(parts,before=GAP):
    p=P(before=before); pf=p.paragraph_format; pf.left_indent=Pt(14); pf.first_line_indent=Pt(-14); pf.tab_stops.add_tab_stop(Pt(14))
    run(p,'→',b=True); run(p,'\t')
    for t,b in parts: run(p,t,b=b)
    return p
def TS(t):
    p=P(before=GAP); run(p,'Tech Stack: ',b=True)
    run(p,t,i=True)
# header
p=P(align=AL.CENTER); run(p,'Ajay I',b=True,sz=25); p.paragraph_format.line_spacing=1.0
p=P(before=2,align=AL.CENTER)
run(p,'ajayiyanraj.2401007@srec.ac.in',sz=BODY); run(p,'  |  ',sz=BODY)
run(p,'LinkedIn',link='https://www.linkedin.com/in/ajay-i-a86298329/',sz=BODY); run(p,'  |  ',sz=BODY)
run(p,'GitHub',link='https://github.com/Ajay-2007-bee',sz=BODY); run(p,'  |  ',sz=BODY)
run(p,'LeetCode',link='https://leetcode.com/u/ajay-2207/',sz=BODY)
H('Summary')
p=P(before=1); p.alignment=AL.LEFT
run(p,'Computer Science Engineering student skilled in predictive machine learning and building AI agents. Proficient in Python, Java, C and C++, with a proven track record in competitive hackathons and a research publication. Passionate about leveraging advanced ML models and robust architectures to develop scalable, data-driven solutions for real-world challenges.')
H('Education')
row([('Sri Ramakrishna Engineering College, Coimbatore',True)],[('2024 – 2028',True)],before=1)
row([('B.E. Computer Science and Engineering (Anna University) | ',False),('CGPA: 9.08 / 10',True),(' (up to 4th Sem)',False)],[('Coimbatore, IN',False)],sz=BODY-0.4)
row([('Githanjali Public School, Coimbatore (CBSE)',True)],[('2022 – 2024',True)],before=2.5)
p=P(); run(p,'HSC (2024) | SSLC (2022)',sz=BODY-0.4)
H('Technical Skills')
B([('Languages: ',True),('Java, C, C++, Python',False)],before=1)
B([('AI / ML & Data: ',True),('Predictive Machine Learning, NLP (Transformers), Multi-Agent AI, CatBoost, XGBoost, Isolation Forest, Pandas, NumPy, NetworkX, Data Analytics & Visualization, Geospatial Processing',False)])
B([('Web & Backend: ',True),('FastAPI, React.js, Tailwind CSS, JavaScript, REST APIs, GitHub / Slack / Twilio / WhatsApp APIs',False)])
B([('Databases & Core: ',True),('Relational DBs, SQLite, PostgreSQL, Firebase | Problem Solving, Full-Stack Development, Rapid Prototyping',False)])
H('Projects')
row([('Workforce Contribution Monitor',True)],[('January 2026',True)],before=1)
B([('Objective: ',True),('bridge the gap between employee activity volume and actual work impact through meaningful contribution metrics.',False)])
B([('Built a full-stack evaluation system integrating data from communication (',False),('Slack',True),(') and execution (',False),('GitHub',True),(') platforms; automated API-based ',False),('ingestion pipelines',True),(' fetch, process and normalize multi-source data.',False)])
B([('Designed a ',False),('custom scoring algorithm',True),(' separating high-activity from high-impact contributions, shown on an interactive, responsive dashboard of trends and performance indicators.',False)])
B([('Focused on modular backend architecture and efficient database handling for scalability, real-time processing and maintainability.',False)])
TS('React, Tailwind CSS, Python, FastAPI, SQLite, GitHub & Slack APIs')
row([('AI Customer Feedback Intelligence & Retention System',True)],[('',False)],before=4.5)
B([('Objective: ',True),('analyze customer sentiment, predict churn and automate personalized retention strategies.',False)])
B([('Analyzes multi-channel feedback using ',False),('Transformer-based NLP',True),(' sentiment analysis and behavioral modeling to identify at-risk users.',False)])
B([('Automated pipeline triggering personalized email and WhatsApp follow-ups based on risk levels.',False)])
B([('Integrated a ',False),('multi-agent decision framework',True),(' recommending refunds, escalations or retention offers, improving decision efficiency.',False)])
TS('Python, NLP (Transformers), FastAPI, React, Twilio API / WhatsApp API, PostgreSQL')
row([('Satellite-Based Water Quality Monitoring & Prediction System',True),(' (Mini Project)',False)],[('',False)],before=4.5)
B([('Objective: ',True),('enable large-scale monitoring and prediction of water quality using satellite data and machine learning.',False)])
B([('Processed satellite and geospatial data to predict water quality parameters with a ',False),('CatBoost Regressor',True),('.',False)])
B([('Built an automated ingestion, preprocessing and prediction pipeline with visualization of spatial and temporal trends.',False)])
TS('Python, CatBoost, Pandas, NumPy, Geospatial Processing, Data Visualization Tools')
row([('ChainWatch – Blockchain AML Intelligence System',True)],[('September 2026',True)],before=4.5)
B([('Objective: ',True),('detect suspicious cryptocurrency transactions using machine learning and behavioural analysis.',False)])
B([('Built a real-time system using ',False),('XGBoost, Isolation Forest',True),(', graph analytics and multi-agent pattern detection to identify high-risk wallet activity and generate explainable risk scores.',False)])
TS('Python, FastAPI, XGBoost, Isolation Forest, NetworkX, Firebase, JavaScript')
H('Publications')
U='https://ieeexplore.ieee.org/document/11654963'
p=P(before=1); run(p,'A High-Resolution Data Fusion Framework for Water Body Assessment Using Sentinel-2 and Gradient Boosting',b=True,link=U)
p=P(); run(p,'IEEE Xplore',link=U,sz=BODY-0.4); run(p,' | ',sz=BODY-0.4); run(p,'DOI',link=U,sz=BODY-0.4); run(p,' | ',sz=BODY-0.4); run(p,'ICCPCT 2026',i=True,sz=BODY-0.4); run(p,' | ',sz=BODY-0.4); run(p,'IEEE Kerala',i=True,sz=BODY-0.4); run(p,' | 2026',sz=BODY-0.4)
H('Hackathons & Achievements')
B([('First Prize',True),(' — 24-hr HackTIDE Hackathon, Amrita Vishwa Vidyapeetham, Coimbatore (08.01.2026) | ',False),('₹25,000',True)],before=1)
B([('Second Prize',True),(' — 12-hr Hackathon, Sri Ramakrishna College of Arts & Science (SRCAS), Coimbatore (17.02.2025) | ',False),('₹2,500',True)])
B([('Third Prize',True),(' — 24-hr VeloHack’26, Vel Tech University, Chennai (11.09.2026) | ',False),('₹16,000',True)])
H('Certifications')
B([('Introduction to Artificial Intelligence',True),(' — Infosys Springboard',False)],before=1)
B([('Cleared ',False),('Business English Certificate (BEC)',True),(' — Business Preliminary Examination',False)])
ORD=['pStyle','keepNext','keepLines','pageBreakBefore','widowControl','numPr','pBdr','shd','tabs','spacing','ind','contextualSpacing','jc','rPr']
for pp in d.element.iter(qn('w:pPr')):
    k=list(pp)
    for e in k: pp.remove(e)
    for e in sorted(k,key=lambda e:ORD.index(e.tag.split('}')[1]) if e.tag.split('}')[1] in ORD else 99): pp.append(e)
z=d.settings.element.find(qn('w:zoom'))
if z is not None: z.set(qn('w:percent'),'100')
d.core_properties.title='Ajay I - Resume'; d.core_properties.author='Ajay I'
d.save('Ajay_I_Resume.docx')
