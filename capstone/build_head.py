import copy, sys, json, os
from docx import Document
from docx.shared import Pt, Inches, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH as AL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from PIL import Image, ImageOps

PAGES = json.load(open('pages.json')) if os.path.exists('pages.json') else [4,4,6,8,20,25]
FONT = 'Times New Roman'
d = Document('cap_template.docx')
body = d.element.body
P = lambda i: docx_par(i)
from docx.text.paragraph import Paragraph
paras = [Paragraph(e, d) for e in body if e.tag == qn('w:p')]

def set_text(p, text):
    p.runs[0].text = text
    for r in p.runs[1:]: r._r.getparent().remove(r._r)

# ---------- cover page ----------
cover = {p.text: p for p in paras}
for p_ in paras:
    if p_.text == '<project title in caps>':
        set_text(p_, 'STUDENT COURSE REGISTRATION SYSTEM')
        for r_ in p_.runs: r_.bold = True
namep = cover['<Name (Rollno)>']
students = [('AISHWARYA U', '71812401006'), ('AJAY IYANRAJ', '71812401007'),
            ('AKASHVARMAN R', '71812401008'), ('AKHILESH RAJ P', '71812401009')]
set_text(namep, f'{students[0][0]} ({students[0][1]})')
prev = namep._p
for n, r_ in students[1:]:
    new = copy.deepcopy(namep._p); prev.addnext(new); prev = new
    Paragraph(new, d).runs[0].text = f'{n} ({r_})'
# remove a few spare blank paragraphs after names so the cover keeps one page
nxt = prev.getnext(); removed = 0
while nxt is not None and removed < 3 and nxt.tag == qn('w:p') and not ''.join(nxt.itertext()).strip():
    n2 = nxt.getnext(); nxt.getparent().remove(nxt); nxt = n2; removed += 1

# explicit page breaks instead of relying on blank filler paragraphs
def blank(e): return e.tag == qn('w:p') and not ''.join(e.itertext()).strip()
def drop_blanks_after(par_el, keep=0):
    n = par_el.getnext(); k = 0
    while n is not None and blank(n):
        nn = n.getnext()
        if k >= keep: n.getparent().remove(n)
        k += 1; n = nn
drop_blanks_after([p for p in paras if p.text.startswith('20CS279')][0]._p, keep=2)
last_cover = [p for p in paras if 'Mr.N.Manoj' in p.text][0]
drop_blanks_after(last_cover._p)
dept2 = [p for p in paras if p.text.startswith('DEPARTMENT')][1]
dept2.paragraph_format.page_break_before = True
sig2 = [p for p in paras if 'Mr.N.Manoj' in p.text][1]
drop_blanks_after(sig2._p)

_sp = copy.deepcopy(cover['CONTENTS']._p)
for r_ in _sp.findall(qn('w:r')): _sp.remove(r_)
cover['CONTENTS']._p.addprevious(_sp)
_spp = Paragraph(_sp, d); _spp.paragraph_format.page_break_before = False
from docx.enum.text import WD_LINE_SPACING
_spp.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY; _spp.paragraph_format.line_spacing = Pt(225)
_spp.paragraph_format.space_before = Pt(0); _spp.paragraph_format.space_after = Pt(0)

# ---------- index table ----------
t0, t1 = d.tables[0], d.tables[1]
for i, (n, r_) in enumerate(students, 1):
    cell = t0.rows[i].cells[1]
    p = cell.paragraphs[0]
    for r in list(p.runs): r._r.getparent().remove(r._r)
    p.paragraph_format.space_before = Pt(0); p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.first_line_indent = None
    run = p.add_run(n); run.bold = True; run.font.size = Pt(11); run.font.name = FONT
    p2 = cell.add_paragraph(); p2.paragraph_format.space_after = Pt(0)
    r2 = p2.add_run(r_); r2.font.size = Pt(11); r2.font.name = FONT
    p2.paragraph_format.left_indent = p.paragraph_format.left_indent

# ---------- contents table ----------
cp = t1.rows[1].cells[1].paragraphs[0]
set_text(cp, 'STUDENT COURSE REGISTRATION SYSTEM')
for i in range(1, 7):
    cell = t1.rows[i].cells[2]
    p = cell.paragraphs[0]
    p.alignment = AL.CENTER
    r = p.add_run(str(PAGES[i-1])); r.font.size = Pt(12); r.font.name = FONT

# ---------- helpers for new content ----------
sect = body.find(qn('w:sectPr'))
def add_el(el): sect.addprevious(el)

def para(text='', bold=False, size=12, align=AL.JUSTIFY, italic=False, space_after=6, keep=False, font=FONT, indent=None):
    p = d.add_paragraph()   # appended before sectPr by python-docx
    p.alignment = align
    pf = p.paragraph_format
    pf.space_after = Pt(space_after); pf.space_before = Pt(0); pf.line_spacing = 1.15
    if keep: pf.keep_with_next = True
    if indent: pf.left_indent = Inches(indent)
    if text:
        r = p.add_run(text); r.bold = bold; r.italic = italic; r.font.size = Pt(size); r.font.name = font
        r._r.rPr.rFonts.set(qn('w:eastAsia'), font)
    return p

def rich(parts, size=12, align=AL.JUSTIFY, space_after=4, indent=None, hanging=None):
    p = para('', align=align, space_after=space_after, indent=indent)
    for txt, b in parts:
        r = p.add_run(txt); r.bold = b; r.font.size = Pt(size); r.font.name = FONT
    if hanging: p.paragraph_format.first_line_indent = Inches(-hanging)
    return p

def heading(num, text, page_break=False):
    p = para((f'{num}. ' if num else '') + text.upper(), bold=True, size=14, align=AL.CENTER, space_after=8, keep=True)
    p.paragraph_format.space_before = Pt(10)
    if page_break: p.paragraph_format.page_break_before = True
    # bottom border
    pPr = p._p.get_or_add_pPr(); b = OxmlElement('w:pBdr'); bt = OxmlElement('w:bottom')
    for k, v in (('val', 'single'), ('sz', '8'), ('space', '1'), ('color', '000000')): bt.set(qn('w:' + k), v)
    b.append(bt)
    anc = next((c for c in pPr if c.tag in [qn('w:' + n) for n in ('shd', 'tabs', 'suppressAutoHyphens', 'spacing', 'ind', 'jc', 'rPr')]), None)
    if anc is not None: anc.addprevious(b)
    else: pPr.append(b)
    return p

def sub(text, page_break=False):
    p = para(text, bold=True, size=12, align=AL.LEFT, space_after=4, keep=True); p.paragraph_format.space_before = Pt(6)
    if page_break: p.paragraph_format.page_break_before = True
    return p

def bullet(text, boldlead=None):
    parts = ([(boldlead, True)] if boldlead else []) + [(text, False)]
    return rich([('•  ', False)] + parts, indent=0.3, hanging=0.2)

def step(n, text, boldlead=None):
    parts = [(f'Step {n}: ', True)] + ([(boldlead, True)] if boldlead else []) + [(text, False)]
    return rich(parts, indent=0.0)

def set_borders(tbl, sz=6):
    tblPr = tbl._tbl.tblPr
    for old in tblPr.findall(qn('w:tblBorders')): tblPr.remove(old)
    b = OxmlElement('w:tblBorders')
    for e in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        x = OxmlElement('w:' + e)
        for k, v in (('val', 'single'), ('sz', str(sz)), ('space', '0'), ('color', '000000')): x.set(qn('w:' + k), v)
        b.append(x)
    anchor = next((c for c in tblPr if c.tag in [qn('w:' + n) for n in ('shd', 'tblLayout', 'tblCellMar', 'tblLook', 'tblCaption')]), None)
    if anchor is not None: anchor.addprevious(b)
    else: tblPr.append(b)

def shade(cell, color):
    tcPr = cell._tc.get_or_add_tcPr(); s = OxmlElement('w:shd')
    s.set(qn('w:val'), 'clear'); s.set(qn('w:color'), 'auto'); s.set(qn('w:fill'), color); tcPr.append(s)

def table(headers, rows, widths, size=10.5):
    t = d.add_table(rows=1, cols=len(headers)); t.autofit = False
    set_borders(t)
    def fill(cell, text, bold=False, w=None, center=False):
        cell.width = Inches(w); p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(2); p.paragraph_format.space_before = Pt(2)
        p.alignment = AL.CENTER if center else AL.LEFT
        r = p.add_run(text); r.bold = bold; r.font.size = Pt(size); r.font.name = FONT
    for i, h in enumerate(headers):
        fill(t.rows[0].cells[i], h, True, widths[i], True); shade(t.rows[0].cells[i], 'D9E2D5')
    trPr = t.rows[0]._tr.get_or_add_trPr(); th = OxmlElement('w:tblHeader'); trPr.append(th)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row): fill(cells[i], v, i == 0, widths[i])
    for i, w in enumerate(widths): t.columns[i].width = Inches(w)
    for r in t.rows[:-1]:
        for c in r.cells:
            for p_ in c.paragraphs: p_.paragraph_format.keep_with_next = True
    for r in t.rows:
        trPr = r._tr.get_or_add_trPr(); cs = OxmlElement('w:cantSplit'); trPr.append(cs)
    para('', space_after=4)
    return t

def code_block(lines, size=8):
    t = d.add_table(rows=1, cols=1); t.autofit = False
    set_borders(t, 8)
    cell = t.rows[0].cells[0]; cell.width = Inches(6.4); shade(cell, 'F5F5F5')
    first = True
    for ln in lines:
        p = cell.paragraphs[0] if first else cell.add_paragraph(); first = False
        pf = p.paragraph_format; pf.space_after = Pt(0); pf.space_before = Pt(0); pf.line_spacing = 1.0
        p.alignment = AL.LEFT
        r = p.add_run(ln if ln else ' '); r.font.name = 'Courier New'; r.font.size = Pt(size)
        r._r.rPr.rFonts.set(qn('w:hAnsi'), 'Courier New'); r._r.rPr.rFonts.set(qn('w:cs'), 'Courier New')
    return t

def figure(path, title, caption, width=6.4, maxh=8.8):
    im = Image.open(path).convert('RGB'); width = min(width, maxh * im.width / im.height); im = ImageOps.expand(im, border=2, fill=(0, 0, 0))
    out = path.replace('.png', '_b.png'); im.save(out)
    if title:
        p = para(title, bold=True, size=12, align=AL.LEFT, space_after=4, keep=True)
        p.paragraph_format.space_before = Pt(6)
    p = para('', align=AL.CENTER, space_after=2, keep=True)
    p.add_run().add_picture(out, width=Inches(width))
    para(caption, italic=True, size=10.5, align=AL.CENTER, space_after=6)

