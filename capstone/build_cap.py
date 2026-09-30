import copy, sys, json, os
from docx import Document
from docx.shared import Pt, Inches, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH as AL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from PIL import Image, ImageOps

PAGES = json.load(open('pages.json')) if os.path.exists('pages.json') else [4,4,6,8,20,25]
FONT = 'Times New Roman'
d = Document('cap_template.docx')
body = d.element.body
P = lambda i: docx_par(i)
from docx.text.paragraph import Paragraph
paras = [Paragraph(e, d) for e in body if e.tag == qn('w:p')]

def set_text(p, text):
    p.runs[0].text = text
    for r in p.runs[1:]: r._r.getparent().remove(r._r)

# ---------- cover page ----------
cover = {p.text: p for p in paras}
for p_ in paras:
    if p_.text == '<project title in caps>':
        set_text(p_, 'STUDENT COURSE REGISTRATION SYSTEM')
        for r_ in p_.runs: r_.bold = True
namep = cover['<Name (Rollno)>']
students = [('AISHWARYA U', '71812401006'), ('AJAY IYANRAJ', '71812401007'),
            ('AKASHVARMAN R', '71812401008'), ('AKHILESH RAJ P', '71812401009')]
set_text(namep, f'{students[0][0]} ({students[0][1]})')
prev = namep._p
for n, r_ in students[1:]:
    new = copy.deepcopy(namep._p); prev.addnext(new); prev = new
    Paragraph(new, d).runs[0].text = f'{n} ({r_})'
# remove a few spare blank paragraphs after names so the cover keeps one page
nxt = prev.getnext(); removed = 0
while nxt is not None and removed < 3 and nxt.tag == qn('w:p') and not ''.join(nxt.itertext()).strip():
    n2 = nxt.getnext(); nxt.getparent().remove(nxt); nxt = n2; removed += 1

# explicit page breaks instead of relying on blank filler paragraphs
def blank(e): return e.tag == qn('w:p') and not ''.join(e.itertext()).strip()
def drop_blanks_after(par_el, keep=0):
    n = par_el.getnext(); k = 0
    while n is not None and blank(n):
        nn = n.getnext()
        if k >= keep: n.getparent().remove(n)
        k += 1; n = nn
drop_blanks_after([p for p in paras if p.text.startswith('20CS279')][0]._p, keep=2)
last_cover = [p for p in paras if 'Mr.N.Manoj' in p.text][0]
drop_blanks_after(last_cover._p)
dept2 = [p for p in paras if p.text.startswith('DEPARTMENT')][1]
dept2.paragraph_format.page_break_before = True
sig2 = [p for p in paras if 'Mr.N.Manoj' in p.text][1]
drop_blanks_after(sig2._p)

_sp = copy.deepcopy(cover['CONTENTS']._p)
for r_ in _sp.findall(qn('w:r')): _sp.remove(r_)
cover['CONTENTS']._p.addprevious(_sp)
_spp = Paragraph(_sp, d); _spp.paragraph_format.page_break_before = False
from docx.enum.text import WD_LINE_SPACING
_spp.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY; _spp.paragraph_format.line_spacing = Pt(225)
_spp.paragraph_format.space_before = Pt(0); _spp.paragraph_format.space_after = Pt(0)

# ---------- index table ----------
t0, t1 = d.tables[0], d.tables[1]
for i, (n, r_) in enumerate(students, 1):
    cell = t0.rows[i].cells[1]
    p = cell.paragraphs[0]
    for r in list(p.runs): r._r.getparent().remove(r._r)
    p.paragraph_format.space_before = Pt(0); p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.first_line_indent = None
    run = p.add_run(n); run.bold = True; run.font.size = Pt(11); run.font.name = FONT
    p2 = cell.add_paragraph(); p2.paragraph_format.space_after = Pt(0)
    r2 = p2.add_run(r_); r2.font.size = Pt(11); r2.font.name = FONT
    p2.paragraph_format.left_indent = p.paragraph_format.left_indent

# ---------- contents table ----------
cp = t1.rows[1].cells[1].paragraphs[0]
set_text(cp, 'STUDENT COURSE REGISTRATION SYSTEM')
for i in range(1, 7):
    cell = t1.rows[i].cells[2]
    p = cell.paragraphs[0]
    p.alignment = AL.CENTER
    r = p.add_run(str(PAGES[i-1])); r.font.size = Pt(12); r.font.name = FONT

# ---------- helpers for new content ----------
sect = body.find(qn('w:sectPr'))
def add_el(el): sect.addprevious(el)

def para(text='', bold=False, size=12, align=AL.JUSTIFY, italic=False, space_after=6, keep=False, font=FONT, indent=None):
    p = d.add_paragraph()   # appended before sectPr by python-docx
    p.alignment = align
    pf = p.paragraph_format
    pf.space_after = Pt(space_after); pf.space_before = Pt(0); pf.line_spacing = 1.15
    if keep: pf.keep_with_next = True
    if indent: pf.left_indent = Inches(indent)
    if text:
        r = p.add_run(text); r.bold = bold; r.italic = italic; r.font.size = Pt(size); r.font.name = font
        r._r.rPr.rFonts.set(qn('w:eastAsia'), font)
    return p

def rich(parts, size=12, align=AL.JUSTIFY, space_after=4, indent=None, hanging=None):
    p = para('', align=align, space_after=space_after, indent=indent)
    for txt, b in parts:
        r = p.add_run(txt); r.bold = b; r.font.size = Pt(size); r.font.name = FONT
    if hanging: p.paragraph_format.first_line_indent = Inches(-hanging)
    return p

def heading(num, text, page_break=False):
    p = para((f'{num}. ' if num else '') + text.upper(), bold=True, size=14, align=AL.CENTER, space_after=8, keep=True)
    p.paragraph_format.space_before = Pt(10)
    if page_break: p.paragraph_format.page_break_before = True
    # bottom border
    pPr = p._p.get_or_add_pPr(); b = OxmlElement('w:pBdr'); bt = OxmlElement('w:bottom')
    for k, v in (('val', 'single'), ('sz', '8'), ('space', '1'), ('color', '000000')): bt.set(qn('w:' + k), v)
    b.append(bt)
    anc = next((c for c in pPr if c.tag in [qn('w:' + n) for n in ('shd', 'tabs', 'suppressAutoHyphens', 'spacing', 'ind', 'jc', 'rPr')]), None)
    if anc is not None: anc.addprevious(b)
    else: pPr.append(b)
    return p

def sub(text, page_break=False):
    p = para(text, bold=True, size=12, align=AL.LEFT, space_after=4, keep=True); p.paragraph_format.space_before = Pt(6)
    if page_break: p.paragraph_format.page_break_before = True
    return p

def bullet(text, boldlead=None):
    parts = ([(boldlead, True)] if boldlead else []) + [(text, False)]
    return rich([('•  ', False)] + parts, indent=0.3, hanging=0.2)

def step(n, text, boldlead=None):
    parts = [(f'Step {n}: ', True)] + ([(boldlead, True)] if boldlead else []) + [(text, False)]
    return rich(parts, indent=0.0)

def set_borders(tbl, sz=6):
    tblPr = tbl._tbl.tblPr
    for old in tblPr.findall(qn('w:tblBorders')): tblPr.remove(old)
    b = OxmlElement('w:tblBorders')
    for e in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        x = OxmlElement('w:' + e)
        for k, v in (('val', 'single'), ('sz', str(sz)), ('space', '0'), ('color', '000000')): x.set(qn('w:' + k), v)
        b.append(x)
    anchor = next((c for c in tblPr if c.tag in [qn('w:' + n) for n in ('shd', 'tblLayout', 'tblCellMar', 'tblLook', 'tblCaption')]), None)
    if anchor is not None: anchor.addprevious(b)
    else: tblPr.append(b)

def shade(cell, color):
    tcPr = cell._tc.get_or_add_tcPr(); s = OxmlElement('w:shd')
    s.set(qn('w:val'), 'clear'); s.set(qn('w:color'), 'auto'); s.set(qn('w:fill'), color); tcPr.append(s)

def table(headers, rows, widths, size=10.5):
    t = d.add_table(rows=1, cols=len(headers)); t.autofit = False
    set_borders(t)
    def fill(cell, text, bold=False, w=None, center=False):
        cell.width = Inches(w); p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(2); p.paragraph_format.space_before = Pt(2)
        p.alignment = AL.CENTER if center else AL.LEFT
        r = p.add_run(text); r.bold = bold; r.font.size = Pt(size); r.font.name = FONT
    for i, h in enumerate(headers):
        fill(t.rows[0].cells[i], h, True, widths[i], True); shade(t.rows[0].cells[i], 'D9E2D5')
    trPr = t.rows[0]._tr.get_or_add_trPr(); th = OxmlElement('w:tblHeader'); trPr.append(th)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row): fill(cells[i], v, i == 0, widths[i])
    for i, w in enumerate(widths): t.columns[i].width = Inches(w)
    for r in t.rows[:-1]:
        for c in r.cells:
            for p_ in c.paragraphs: p_.paragraph_format.keep_with_next = True
    for r in t.rows:
        trPr = r._tr.get_or_add_trPr(); cs = OxmlElement('w:cantSplit'); trPr.append(cs)
    para('', space_after=4)
    return t

def code_block(lines, size=8):
    t = d.add_table(rows=1, cols=1); t.autofit = False
    set_borders(t, 8)
    cell = t.rows[0].cells[0]; cell.width = Inches(6.4); shade(cell, 'F5F5F5')
    first = True
    for ln in lines:
        p = cell.paragraphs[0] if first else cell.add_paragraph(); first = False
        pf = p.paragraph_format; pf.space_after = Pt(0); pf.space_before = Pt(0); pf.line_spacing = 1.0
        p.alignment = AL.LEFT
        r = p.add_run(ln if ln else ' '); r.font.name = 'Courier New'; r.font.size = Pt(size)
        r._r.rPr.rFonts.set(qn('w:hAnsi'), 'Courier New'); r._r.rPr.rFonts.set(qn('w:cs'), 'Courier New')
    return t

def figure(path, title, caption, width=6.4, maxh=8.8):
    im = Image.open(path).convert('RGB'); width = min(width, maxh * im.width / im.height); im = ImageOps.expand(im, border=2, fill=(0, 0, 0))
    out = path.replace('.png', '_b.png'); im.save(out)
    if title:
        p = para(title, bold=True, size=12, align=AL.LEFT, space_after=4, keep=True)
        p.paragraph_format.space_before = Pt(6)
    p = para('', align=AL.CENTER, space_after=2, keep=True)
    p.add_run().add_picture(out, width=Inches(width))
    para(caption, italic=True, size=10.5, align=AL.CENTER, space_after=6)

# ================= CONTENT =================
students = [('AISHWARYA U', '71812401006'), ('AJAY IYANRAJ', '71812401007'),
            ('AKASHVARMAN R', '71812401008'), ('AKHILESH RAJ P', '71812401009')]

def snippet(lines, size=9):
    return code_block(lines, size)

# ---------- TITLE / ABSTRACT PAGE ----------
heading(None, 'Student Course Registration System', page_break=True)
sub('Abstract')
para('The Student Course Registration System is a web application that allows the students of a college to log in with '
     'their roll number, browse the list of courses offered in the semester, register for the courses they want, drop '
     'a course, and view the registered courses in a weekly timetable. The system checks every registration against '
     'the available seats, the maximum credit limit and the time-slot clash rules, and it shows clear messages when a '
     'rule fails. The project is developed with HTML5 for the structure, CSS3 for the design and JavaScript for the '
     'logic, and the registrations are stored in the browser localStorage, so the complete application runs on the '
     'client side without any server.')
sub('Team members')
table(['S. No.', 'Name', 'Roll Number'], [[str(i), n, r] for i, (n, r) in enumerate(students, 1)], [0.8, 3.4, 2.2])
sub('Objectives')
for txt in ['To replace the manual, paper based course registration with a simple and fast web application.',
            'To allow only valid students to log in and to register for courses.',
            'To prevent over-booking of seats, crossing of the credit limit and clash of class timings.',
            'To show the registered courses, the total credits and the weekly timetable instantly.',
            'To keep the design responsive so that it works on desktop and mobile screens.']:
    bullet(txt)
sub('Technologies used')
table(['Technology / Tool', 'Purpose'], [
    ['HTML5', 'Page structure: login form, tabs, tables and containers'],
    ['CSS3', 'Layout, colours, badges, progress bar, timetable grid and responsive design'],
    ['JavaScript (ES6)', 'Login validation, course filtering, registration rules, timetable and localStorage'],
    ['Browser localStorage', 'Stores the registrations of all students without a database'],
    ['Visual Studio Code', 'Writing and formatting the source code'],
    ['Google Chrome / Playwright', 'Running the application and capturing the output screenshots']],
    [2.2, 4.2])

# ---------- 1. INTRODUCTION ----------
heading(1, 'Introduction')
sub('1.1 Background')
para('At the beginning of every semester a student has to choose the courses to study, such as the core subjects, '
     'the electives and the laboratory courses. In many colleges this is still done with paper forms or spreadsheets. '
     'The staff have to check the seats, the credits and the time-table of every student by hand, which takes a lot '
     'of time and often leads to mistakes such as two courses at the same hour or more students than seats.')
sub('1.2 Problem Statement')
para('Design and develop a Student Course Registration System in which a student can log in, view the courses with '
     'their credits, faculty, time slots and available seats, register or drop courses, and see the total credits and '
     'the timetable. The system must automatically reject a registration when the course is full, when the credit '
     'limit of 24 is crossed, or when the time slot clashes with an already registered course.')
sub('1.3 Scope of the project')
for txt in ['Student side functions: login, search / filter courses, register, drop, view credits and timetable.',
            'Twelve sample courses of five departments with limited seats and fixed weekly time slots.',
            'Four sample student accounts (the members of this team) protected by a password.',
            'A client side implementation only; a real deployment would store the data on a server database.']:
    bullet(txt)
sub('1.4 Modules of the system')
table(['Module', 'Description'], [
    ['Login', 'Checks the roll number format, the existence of the student and the password.'],
    ['Available Courses', 'Lists all courses with seats left; supports search by code, title, faculty and department filter.'],
    ['Registration engine', 'Applies the rules for seats, credits and time clash before saving a registration.'],
    ['My Registrations', 'Shows the registered courses, total credits with a progress bar, and a Drop button.'],
    ['Timetable', 'Builds a Monday to Friday, five period grid from the slots of the registered courses.'],
    ['Storage', 'Saves the registrations of every student in localStorage so they remain after a page refresh.']],
    [1.7, 4.7])
sub('1.5 Key features')
for lead, txt in [('Live seat count: ', 'the seats left of every course are updated immediately and shown in green, yellow or red badges.'),
                  ('Credit control: ', 'a progress bar shows the credits used out of the maximum of 24.'),
                  ('Clash detection: ', 'a course whose slot is already taken is rejected with the name of the clashing course.'),
                  ('Responsive design: ', 'the layout adapts to phone screens with a scrollable table.')]:
    bullet(txt, lead)
sub('1.6 Advantages')
for txt in ['Saves the time of both students and the staff.',
            'Avoids human errors in seat allotment and in time-table clashes.',
            'Works in any browser without installation, and the data is available offline.']:
    bullet(txt)

# ---------- 2. LOGIC BUILDING ----------
heading(2, 'Logic Building')
sub('2.1 Overall logic')
step(1, 'The student enters the roll number and the password; the login logic validates both.', 'Login: ')
step(2, 'After a successful login the application hides the login view and shows the main view with the student name.', 'Show portal: ')
step(3, 'The array of courses is displayed as table rows; the search text and department filter reduce the rows.', 'List courses: ')
step(4, 'When Register is clicked the three rules (seats, credits, clash) are checked one after another.', 'Validate: ')
step(5, 'If all rules pass the course code is added to the list of the student and saved in localStorage.', 'Save: ')
step(6, 'The seat badges, credit chip, progress bar, My Registrations table and the timetable are redrawn.', 'Refresh: ')
sub('2.2 Data used by the program')
para('Each course is stored as a JavaScript object. The registrations of all the students are stored in one object in '
     'localStorage, where the key is the roll number and the value is the list of registered course codes.')
snippet(['// one course object',
         "{ code: '20CS212', title: 'Web Technologies', dept: 'CSE', credits: 3,",
         "  faculty: 'Mr. N. Manoj', slots: ['Mon-1', 'Wed-2'], seats: 60, taken: 41 }",
         '',
         '// registrations saved in localStorage under the key "scrs"',
         "{ '71812401006': ['20CS212', '20CS202'], '71812401007': ['20CS205'] }",
         '',
         '// seats left = total seats - seats already taken - registrations of all students',
         'seatsLeft = course.seats - course.taken - countOfStudentsRegistered(course.code)'], 9)
para('', space_after=4)
sub('2.3 Business rules')
table(['Rule', 'Condition checked', 'Message shown'], [
    ['Valid roll number', '11 digits and present in the student list', 'Roll number must contain exactly 11 digits / not found'],
    ['Valid password', 'At least 8 characters and equal to the stored password', 'Incorrect password. Please try again.'],
    ['Seat availability', 'seats left > 0', '20CS206 is full. No seats left.'],
    ['Credit limit', 'registered credits + course credits <= 24', 'Credit limit of 24 exceeded.'],
    ['Time clash', 'no common slot with any registered course', 'Time clash with 20CS212 at Mon-1.'],
    ['Duplicate', 'a registered course shows a disabled button', 'Registered (button disabled)']],
    [1.5, 2.8, 2.1], size=11)
sub('2.4 Algorithm for registering a course')
snippet(['registerCourse(code):',
         '    course <- find course with this code',
         '    if seatsLeft(course) <= 0                        -> show "course is full"; stop',
         '    if totalCredits() + course.credits > 24          -> show "credit limit exceeded"; stop',
         '    if any registered course shares a slot           -> show "time clash"; stop',
         '    add code to registrations of the current student',
         '    save registrations in localStorage',
         '    show "Registered" message',
         '    redraw catalogue, my registrations and timetable'], 9)
para('', space_after=4)
sub('2.5 Flowchart')
figure('shots/flow.png', 'Flowchart of login and course registration', 'Fig. 1 - Logic of login and course registration.', width=4.2, maxh=6.3)
sub('2.6 Logic of credits and timetable')
para('The total credits are the sum of the credits of the registered courses; the width of the progress bar is '
     'total / 24 x 100 percent. For the timetable every slot is written as Day-Period (for example Wed-2). A map is '
     'created from slot to course code and the grid of five days and five periods is drawn by looking up each cell in '
     'this map, so a clash is impossible because a slot can hold only one course.')

# ---------- 3. IMPLEMENTATION STEPS ----------
heading(3, 'Implementation Steps')
sub('3.1 Steps followed')
for i, (lead, txt) in enumerate([
    ('Create the project folder ', 'with three files: index.html, style.css and script.js.'),
    ('Design the HTML: ', 'a login section, and a main section with a top bar, three tabs and three panels.'),
    ('Style with CSS: ', 'green theme, cards, badges, progress bar, timetable grid and a media query for phones.'),
    ('Store the data: ', 'define the student list, the constants and the array of twelve courses in script.js.'),
    ('Write the login logic: ', 'validate the roll number and the password and switch the views.'),
    ('Display the catalogue: ', 'build the table rows with template strings; add the search box and the department filter.'),
    ('Write the registration engine: ', 'functions seatsLeft, totalCredits, findClash, registerCourse and dropCourse.'),
    ('Build My Registrations and Timetable: ', 'render the list, the credit bar and the weekly grid.'),
    ('Save to localStorage ', 'after every register or drop so the data survives a refresh.'),
    ('Test the application ', 'with wrong login, search, filter, registration, clash, full course and drop cases and capture the screenshots.')], 1):
    step(i, txt, lead)
sub('3.2 Project files')
table(['File', 'Lines', 'Purpose'], [
    ['index.html', '84', 'Structure of the login page, the tabs, the tables and the toast message box'],
    ['style.css', '66', 'Appearance of every component and the responsive rules'],
    ['script.js', '171', 'Data, login, filtering, registration rules, timetable and storage']],
    [1.6, 0.9, 3.9])
sub('3.3 Main functions of script.js')
table(['Function', 'Purpose'], [
    ['seatsLeft(course)', 'Returns the seats still free after counting the registrations of all students'],
    ['totalCredits()', 'Adds the credits of the courses registered by the current student'],
    ['findClash(course)', 'Returns the registered course that shares a time slot with the given course'],
    ['registerCourse(code)', 'Applies seat, credit and clash rules and saves the registration'],
    ['dropCourse(code)', 'Removes a registered course and refreshes the screen'],
    ['renderCatalogue()', 'Draws the course table after applying the search text and department filter'],
    ['renderMine()', 'Draws the registered courses and the credit progress bar'],
    ['renderTimetable()', 'Draws the Monday to Friday timetable grid'],
    ['toast(msg, type)', 'Shows a green success or red error message for a short time']],
    [2.0, 4.4], size=11)
sub('3.4 Test cases')
table(['S. No.', 'Test case', 'Expected result', 'Result'], [
    ['1', 'Login with a wrong password', 'Error message "Incorrect password"', 'Pass'],
    ['2', 'Login with a valid roll number and password', 'Portal opens with the student name', 'Pass'],
    ['3', 'Filter by department AI&DS', 'Only the two AI&DS courses are listed', 'Pass'],
    ['4', 'Register 20CS212', 'Success message and credits become 3 / 24', 'Pass'],
    ['5', 'Register 20CS204 after 20CS212', 'Error: time clash at Mon-1', 'Pass'],
    ['6', 'Open a full course (20CS206)', 'Button shows Full and is disabled', 'Pass'],
    ['7', 'Open the Timetable tab', 'Registered slots are highlighted', 'Pass']],
    [0.6, 2.6, 2.6, 0.6], size=11)

# ---------- 4. CODE AND OUTPUT SCREENSHOTS ----------
heading(4, 'Code and Output Screenshots', page_break=True)
sub('4.1 index.html')
code_block(open('index.html').read().split('\n'), 8)
para('', space_after=4)
sub('4.2 style.css')
code_block(open('style.css').read().split('\n'), 8)
para('', space_after=4)
sub('4.3 script.js')
code_block(open('script.js').read().split('\n'), 8)
para('', space_after=4)
sub('4.4 Output screenshots')
para('The screenshots below were captured from the running application in a Chromium browser.')
mob = Image.open('shots/9_mobile.png'); mob.crop((0, 0, mob.width, 1250)).save('shots/9_mobile_crop.png')
figs = [
 ('shots/1_login.png', 'Output 1: Login page', 'Fig. 2 - Login page asking for the roll number and password.'),
 ('shots/2_login_error.png', 'Output 2: Login with a wrong password', 'Fig. 3 - Error message shown for an incorrect password.'),
 ('shots/3_catalogue.png', 'Output 3: Available courses after login', 'Fig. 4 - Course catalogue with credits, faculty, slots and seat badges (20CS206 is full).'),
 ('shots/4_filter.png', 'Output 4: Department filter', 'Fig. 5 - Filtering by the AI&DS department shows only two courses.'),
 ('shots/4b_search.png', 'Output 5: Search box', 'Fig. 6 - Searching "web" lists Web Technologies and Web Technologies Lab.'),
 ('shots/5_registered.png', 'Output 6: Successful registration', 'Fig. 7 - Registering 20CS212 shows a green message and the button changes to Registered.'),
 ('shots/6_clash.png', 'Output 7: Time clash error', 'Fig. 8 - Registering 20CS204 is rejected because it clashes with 20CS212 at Mon-1.'),
 ('shots/7_mine.png', 'Output 8: My Registrations', 'Fig. 9 - Four registered courses, 12 of 24 credits and a Drop button for each course.'),
 ('shots/8_timetable.png', 'Output 9: Timetable', 'Fig. 10 - Weekly timetable with the slots of the registered courses highlighted.'),
 ('shots/9_mobile_crop.png', 'Output 10: Mobile view', 'Fig. 11 - The portal on a 390 px wide phone screen.')]
for f, t, c in figs:
    if 'mobile' in f: figure(f, None, c, width=6.0, maxh=3.6)
    else: figure(f, None, c, width=4.8, maxh=3.4)

# ---------- 5. CONCLUSION ----------
heading(5, 'Conclusion')
para('The Student Course Registration System was designed and implemented with HTML5, CSS3 and JavaScript. It '
     'allows a student to log in, search and filter the courses, register and drop courses and view the credits and '
     'the timetable. The rules for seat availability, credit limit and time clash are applied automatically, which '
     'removes the manual work and the mistakes of the paper based process. The test cases and the screenshots show '
     'that the application works correctly on desktop and mobile screens.')
sub('Learning outcomes')
for txt in ['Built a multi-view single page application using only HTML, CSS and JavaScript.',
            'Applied array methods (filter, find, reduce, map) to search and process the course data.',
            'Used localStorage to keep data in the browser and events to make the page interactive.']:
    bullet(txt)
sub('Future enhancements')
for txt in ['Connect the application to a server and a database (Node.js / MySQL) for real data storage.',
            'Add an administrator login to create courses, change seats and approve registrations.',
            'Send the registration confirmation by email and allow the timetable to be downloaded as a PDF.',
            'Add prerequisites, elective groups and a waiting list for full courses.']:
    bullet(txt)
sub('References')
table(['S. No.', 'Reference'], [
    ['1', 'MDN Web Docs - HTML, CSS and JavaScript reference (developer.mozilla.org)'],
    ['2', 'MDN Web Docs - Window.localStorage and Web Storage API'],
    ['3', 'W3Schools - JavaScript array methods and CSS grid tutorials']],
    [0.8, 5.6])

# ---------- footer page numbers, margins, page border ----------
sec = d.sections[0]
sec.bottom_margin = Emu(int(0.6 * 914400)); sec.footer_distance = Emu(int(0.25 * 914400))
fp = sec.footer.paragraphs[0]; fp.alignment = AL.CENTER
def fld(p, instr):
    r = p.add_run(); r.font.size = Pt(10); r.font.name = FONT
    a = OxmlElement('w:fldChar'); a.set(qn('w:fldCharType'), 'begin'); r._r.append(a)
    it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve'); it.text = instr; r._r.append(it)
    c = OxmlElement('w:fldChar'); c.set(qn('w:fldCharType'), 'separate'); r._r.append(c)
    t = OxmlElement('w:t'); t.text = '1'; r._r.append(t)
    e = OxmlElement('w:fldChar'); e.set(qn('w:fldCharType'), 'end'); r._r.append(e)
fld(fp, 'PAGE')

sp = sect
pb = OxmlElement('w:pgBorders'); pb.set(qn('w:offsetFrom'), 'page')
for side in ('top', 'left', 'bottom', 'right'):
    e = OxmlElement('w:' + side)
    for k_, v in (('val', 'single'), ('sz', '12'), ('space', '20'), ('color', '000000')): e.set(qn('w:' + k_), v)
    pb.append(e)
sp.find(qn('w:pgMar')).addnext(pb)
# section break after page 2: header only on cover + index pages
sig2._p.get_or_add_pPr().append(copy.deepcopy(sp))
last = d.sections[-1]
for hr in last._sectPr.findall(qn('w:headerReference')): last._sectPr.remove(hr)
last.header.is_linked_to_previous = False
for p_ in last.header.paragraphs:
    for r_ in list(p_.runs): r_._r.getparent().remove(r_._r)
last.top_margin = Inches(0.8)
d.core_properties.title = 'Capstone Project Report - Student Course Registration System'
d.save('Capstone_Student_Course_Registration_System.docx')
