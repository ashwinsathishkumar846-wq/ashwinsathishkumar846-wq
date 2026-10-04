import copy, json, os, math
from docx import Document
from docx.shared import Pt, Inches, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH as AL, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph
from PIL import Image, ImageOps

TITLE = 'TEMPERATURE MONITORING SYSTEM USING LPC2148'
PAGES = json.load(open('pages.json')) if os.path.exists('pages.json') else [4,5,7,9,12,17,19,21]
FONT = 'Times New Roman'
d = Document('mcp_template.docx')
body = d.element.body
paras = [Paragraph(e, d) for e in body if e.tag == qn('w:p')]
sect = body.find(qn('w:sectPr'))

def set_text(p, text):
    p.runs[0].text = text
    for r in p.runs[1:]: r._r.getparent().remove(r._r)
def blank(e): return e.tag == qn('w:p') and not ''.join(e.itertext()).strip()
def drop_blanks_after(par_el, keep=0):
    n = par_el.getnext(); k = 0
    while n is not None and blank(n):
        nn = n.getnext()
        if k >= keep: n.getparent().remove(n)
        k += 1; n = nn

students = [('AJAY IYANRAJ', '71812401007'), ('AKASHVARMAN R', '71812401008'), ('AKHILESH RAJ P', '71812401009'),
            ('AKSHAY KRISHNA S M', '71812401010'), ('AKSHAYA KEERTHI S', '71812401011')]

# ---------- cover page ----------
for p_ in paras:
    if p_.text == '<project title in caps>':
        set_text(p_, TITLE)
        for r_ in p_.runs: r_.bold = True
namep = [p for p in paras if p.text == '<Name (Rollno)>'][0]
set_text(namep, f'{students[0][0]} ({students[0][1]})')
prev = namep._p
for n, r_ in students[1:]:
    new = copy.deepcopy(namep._p); prev.addnext(new); prev = new
    Paragraph(new, d).runs[0].text = f'{n} ({r_})'
nxt = prev.getnext(); removed = 0
while nxt is not None and removed < 3 and blank(nxt):
    n2 = nxt.getnext(); nxt.getparent().remove(nxt); nxt = n2; removed += 1
cour = [p for p in paras if 'Amuthasurabi' in p.text]
drop_blanks_after(cour[0]._p)          # blank filler after the cover signature block
sig2 = cour[1]
drop_blanks_after(sig2._p)
dept2 = [p for p in paras if p.text.startswith('DEPARTMENT')][1]
dept2.paragraph_format.page_break_before = True

# contents page: exact-height spacer so the table sits in the middle of page 3
contents = [p for p in paras if p.text == 'CONTENTS'][0]
_sp = copy.deepcopy(contents._p)
for r_ in _sp.findall(qn('w:r')): _sp.remove(r_)
contents._p.addprevious(_sp)
_spp = Paragraph(_sp, d); _spp.paragraph_format.page_break_before = True
_spp.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY; _spp.paragraph_format.line_spacing = Pt(185)
_spp.paragraph_format.space_before = Pt(0); _spp.paragraph_format.space_after = Pt(0)
contents.paragraph_format.page_break_before = False

# ---------- index table ----------
t0, t1 = d.tables[0], d.tables[1]
for i, (n, r_) in enumerate(students, 1):
    if i == 5:
        src_p = t0.rows[4].cells[0].paragraphs[0]._p
        dst_cell = t0.rows[5].cells[0]
        dst_cell._tc.remove(dst_cell.paragraphs[0]._p)
        dst_cell._tc.append(copy.deepcopy(src_p))
        Paragraph(dst_cell._tc.findall(qn('w:p'))[0], d).runs[0].text = '5.'
    cell = t0.rows[i].cells[1]
    p = cell.paragraphs[0]
    for r in list(p.runs): r._r.getparent().remove(r._r)
    p.paragraph_format.space_before = Pt(0); p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    run = p.add_run(n); run.bold = True; run.font.size = Pt(11); run.font.name = FONT
    p2 = cell.add_paragraph(); p2.paragraph_format.space_after = Pt(0)
    r2 = p2.add_run(r_); r2.font.size = Pt(11); r2.font.name = FONT

# ---------- contents table ----------
set_text(t1.rows[1].cells[1].paragraphs[0], TITLE)
for i in range(1, 9):
    p = t1.rows[i].cells[2].paragraphs[0]
    p.alignment = AL.CENTER
    r = p.add_run(str(PAGES[i-1])); r.font.size = Pt(12); r.font.name = FONT

# ---------- helpers for new content ----------

def para(text='', bold=False, size=12, align=AL.JUSTIFY, italic=False, space_after=6, keep=False, font=FONT, indent=None):
    p = d.add_paragraph()   # appended before sectPr by python-docx
    p.alignment = align
    pf = p.paragraph_format
    pf.space_after = Pt(space_after); pf.space_before = Pt(0); pf.line_spacing = 1.15
    if keep: pf.keep_with_next = True
    if indent: pf.left_indent = Inches(indent)
    if text:
        r = p.add_run(text); r.bold = bold; r.italic = italic; r.font.size = Pt(size); r.font.name = font
        r._r.rPr.rFonts.set(qn('w:eastAsia'), font)
    return p

def rich(parts, size=12, align=AL.JUSTIFY, space_after=4, indent=None, hanging=None):
    p = para('', align=align, space_after=space_after, indent=indent)
    for txt, b in parts:
        r = p.add_run(txt); r.bold = b; r.font.size = Pt(size); r.font.name = FONT
    if hanging: p.paragraph_format.first_line_indent = Inches(-hanging)
    return p

def heading(num, text, page_break=False):
    p = para((f'{num}. ' if num else '') + text.upper(), bold=True, size=14, align=AL.CENTER, space_after=8, keep=True)
    p.paragraph_format.space_before = Pt(10)
    if page_break: p.paragraph_format.page_break_before = True
    # bottom border
    pPr = p._p.get_or_add_pPr(); b = OxmlElement('w:pBdr'); bt = OxmlElement('w:bottom')
    for k, v in (('val', 'single'), ('sz', '8'), ('space', '1'), ('color', '000000')): bt.set(qn('w:' + k), v)
    b.append(bt)
    anc = next((c for c in pPr if c.tag in [qn('w:' + n) for n in ('shd', 'tabs', 'suppressAutoHyphens', 'spacing', 'ind', 'jc', 'rPr')]), None)
    if anc is not None: anc.addprevious(b)
    else: pPr.append(b)
    return p

def sub(text, page_break=False):
    p = para(text, bold=True, size=12, align=AL.LEFT, space_after=4, keep=True); p.paragraph_format.space_before = Pt(6)
    if page_break: p.paragraph_format.page_break_before = True
    return p

def bullet(text, boldlead=None):
    parts = ([(boldlead, True)] if boldlead else []) + [(text, False)]
    return rich([('•  ', False)] + parts, indent=0.3, hanging=0.2)

def step(n, text, boldlead=None):
    parts = [(f'Step {n}: ', True)] + ([(boldlead, True)] if boldlead else []) + [(text, False)]
    return rich(parts, indent=0.0)

def set_borders(tbl, sz=6):
    tblPr = tbl._tbl.tblPr
    for old in tblPr.findall(qn('w:tblBorders')): tblPr.remove(old)
    b = OxmlElement('w:tblBorders')
    for e in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        x = OxmlElement('w:' + e)
        for k, v in (('val', 'single'), ('sz', str(sz)), ('space', '0'), ('color', '000000')): x.set(qn('w:' + k), v)
        b.append(x)
    anchor = next((c for c in tblPr if c.tag in [qn('w:' + n) for n in ('shd', 'tblLayout', 'tblCellMar', 'tblLook', 'tblCaption')]), None)
    if anchor is not None: anchor.addprevious(b)
    else: tblPr.append(b)

def shade(cell, color):
    tcPr = cell._tc.get_or_add_tcPr(); s = OxmlElement('w:shd')
    s.set(qn('w:val'), 'clear'); s.set(qn('w:color'), 'auto'); s.set(qn('w:fill'), color); tcPr.append(s)

def table(headers, rows, widths, size=10.5):
    t = d.add_table(rows=1, cols=len(headers)); t.autofit = False
    set_borders(t)
    def fill(cell, text, bold=False, w=None, center=False):
        cell.width = Inches(w); p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(2); p.paragraph_format.space_before = Pt(2)
        p.alignment = AL.CENTER if center else AL.LEFT
        r = p.add_run(text); r.bold = bold; r.font.size = Pt(size); r.font.name = FONT
    for i, h in enumerate(headers):
        fill(t.rows[0].cells[i], h, True, widths[i], True); shade(t.rows[0].cells[i], 'D9E2D5')
    trPr = t.rows[0]._tr.get_or_add_trPr(); th = OxmlElement('w:tblHeader'); trPr.append(th)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row): fill(cells[i], v, i == 0, widths[i])
    for i, w in enumerate(widths): t.columns[i].width = Inches(w)
    for r in t.rows[:-1]:
        for c in r.cells:
            for p_ in c.paragraphs: p_.paragraph_format.keep_with_next = True
    for r in t.rows:
        trPr = r._tr.get_or_add_trPr(); cs = OxmlElement('w:cantSplit'); trPr.append(cs)
    para('', space_after=4)
    return t

def code_block(lines, size=8):
    t = d.add_table(rows=1, cols=1); t.autofit = False
    set_borders(t, 8)
    cell = t.rows[0].cells[0]; cell.width = Inches(7.2); shade(cell, 'F5F5F5')
    first = True
    for ln in lines:
        p = cell.paragraphs[0] if first else cell.add_paragraph(); first = False
        pf = p.paragraph_format; pf.space_after = Pt(0); pf.space_before = Pt(0); pf.line_spacing = 1.0
        p.alignment = AL.LEFT
        r = p.add_run(ln if ln else ' '); r.font.name = 'Courier New'; r.font.size = Pt(size)
        r._r.rPr.rFonts.set(qn('w:hAnsi'), 'Courier New'); r._r.rPr.rFonts.set(qn('w:cs'), 'Courier New')
    return t

def figure(path, title, caption, width=6.4, maxh=8.8, new_page=False):
    im = Image.open(path).convert('RGB'); width = min(width, maxh * im.width / im.height); im = ImageOps.expand(im, border=2, fill=(0, 0, 0))
    out = path.replace('.png', '_b.png'); im.save(out)
    if title:
        p = para(title, bold=True, size=12, align=AL.LEFT, space_after=4, keep=True)
        p.paragraph_format.space_before = Pt(6)
        if new_page: p.paragraph_format.page_break_before = True
    p = para('', align=AL.CENTER, space_after=2, keep=True)
    p.add_run().add_picture(out, width=Inches(width))
    para(caption, italic=True, size=10.5, align=AL.CENTER, space_after=6)

# ================= CONTENT =================
def adc_of(t): return round(t * 0.010 / 3.3 * 1023)
def temp_of(a): return a * 330.0 / 1023

# ---------- TITLE / ABSTRACT PAGE ----------
heading(None, TITLE, page_break=True)
sub('Abstract')
para('Temperature is one of the most commonly monitored physical quantities in industries, laboratories, server rooms, '
     'greenhouses and homes. This project designs a digital temperature monitoring system around the LPC2148 '
     'ARM7TDMI-S microcontroller. An LM35 precision temperature sensor produces an analog voltage of 10 mV for every '
     'degree Celsius. This voltage is applied to the on-chip 10-bit Analog to Digital Converter (ADC0, channel AD0.1 '
     'on pin P0.28) of the LPC2148, which converts it into a digital number between 0 and 1023. The program converts '
     'the number into degrees Celsius, shows it on a 16 x 2 LCD, sends it to a PC through UART0, and switches a fan, a '
     'buzzer and an alarm LED when the temperature crosses the limit of 40 °C. The report explains in detail how '
     'the ADC of the LPC2148 is configured and used to read the sensor value.')
sub('Team members')
table(['S. No.', 'Name', 'Roll Number'], [[str(i), n, r] for i, (n, r) in enumerate(students, 1)], [0.9, 3.8, 2.5], size=11.5)
sub('Objectives')
for txt in ['To study the architecture and the on-chip ADC of the LPC2148 microcontroller.',
            'To interface the LM35 analog temperature sensor with the ADC channel AD0.1 (P0.28).',
            'To convert the digital ADC result into the temperature in degrees Celsius and display it on an LCD and a serial terminal.',
            'To switch a fan, a buzzer and an alarm LED automatically when the temperature exceeds the set limit.']:
    bullet(txt)
sub('Hardware and software used')
table(['Component / Tool', 'Specification', 'Purpose'], [
    ['LPC2148 development board', 'ARM7TDMI-S, 512 KB flash, 12 MHz crystal', 'Main controller with on-chip ADC'],
    ['LM35 sensor', '10 mV / °C, 4 V to 30 V supply', 'Converts temperature into voltage'],
    ['16 x 2 LCD (JHD162A)', 'HD44780 controller, 4-bit mode', 'Shows temperature and status'],
    ['Relay + DC fan, buzzer, LED', '5 V relay, 5 V buzzer, 330 Ω resistor', 'Cooling and alarm outputs'],
    ['Keil uVision, Flash Magic', 'ARM compiler and ISP programmer', 'Writing, building and loading the code'],
    ['Serial terminal / HTML simulator', '9600 baud, 8N1', 'Viewing the output']],
    [2.2, 2.7, 2.3], size=11)

# ---------- 1. INTRODUCTION ----------
heading(1, 'Introduction', page_break=True)
sub('1.1 Background')
para('Many systems must run within a safe temperature range. A server room that becomes too hot damages the '
     'equipment, a cold storage that becomes warm spoils the food, and an overheated motor may burn. A temperature '
     'monitoring system measures the temperature continuously, shows it to the user and takes an action, such as '
     'switching on a fan, when the temperature is high. A microcontroller based system is small, cheap and accurate, and '
     'it can also send the readings to a computer.')
sub('1.2 Problem statement')
para('Design a temperature monitoring system using LPC2148. Explain how the ADC (Analog to Digital Converter) of '
     'LPC2148 can be used to read the temperature sensor value.')
sub('1.3 LPC2148 microcontroller')
para('The LPC2148 is a 16/32-bit ARM7TDMI-S based microcontroller from NXP (Philips) with real-time emulation and '
     'embedded trace support. It is widely used in embedded system laboratories because it contains many peripherals on '
     'a single chip.')
table(['Feature', 'Details'], [
    ['CPU', '16/32-bit ARM7TDMI-S, up to 60 MHz using the on-chip PLL'],
    ['Memory', '512 KB on-chip flash, 32 KB static RAM + 8 KB USB RAM'],
    ['ADC', 'Two 10-bit ADCs: ADC0 with 6 channels and ADC1 with 8 channels'],
    ['DAC', 'One 10-bit digital to analog converter (AOUT on P0.25)'],
    ['Serial', 'Two UARTs, two I2C, SPI, SSP, USB 2.0 full speed device'],
    ['Timers', 'Two 32-bit timers, PWM unit, real time clock, watchdog timer'],
    ['I/O', 'Up to 45 general purpose I/O pins, 5 V tolerant digital pins'],
    ['Supply', '3.3 V (core 1.8 V from an internal regulator)']],
    [1.5, 5.7], size=11)
sub('1.4 ADC of the LPC2148')
para('The ADC of the LPC2148 is a successive approximation converter. It has a 10-bit resolution, so the input '
     'voltage from 0 V to VREF is divided into 1024 steps. The important features of ADC0 are given below.')
table(['Parameter', 'Value for ADC0'], [
    ['Resolution', '10 bit (0 to 1023 counts)'],
    ['Input channels', 'AD0.1, AD0.2, AD0.3, AD0.4, AD0.6 and AD0.7 (pins P0.28, P0.29, P0.30, P0.25, P0.4, P0.5)'],
    ['Reference voltage', 'VREF = 3.3 V (VDDA); input range 0 V to 3.3 V'],
    ['Maximum ADC clock', '4.5 MHz (obtained from PCLK using the CLKDIV field of AD0CR)'],
    ['Conversion time', '11 clock cycles for 10 bits = 2.44 µs at 4.5 MHz'],
    ['Conversion modes', 'Software controlled (one channel) and burst mode (several channels repeatedly)'],
    ['Start options', 'Software start, external pin edge or timer match'],
    ['Result registers', 'AD0DR0 to AD0DR7 for each channel and AD0GDR global data register']],
    [1.8, 5.4], size=11)
sub('1.5 LM35 temperature sensor')
table(['Parameter', 'Value'], [
    ['Type', 'Precision integrated-circuit temperature sensor (analog output)'],
    ['Scale factor', '+10.0 mV per °C (linear)'],
    ['Temperature range', '0 °C to 100 °C (LM35D) / -55 °C to 150 °C (LM35)'],
    ['Accuracy', '±0.5 °C at 25 °C'],
    ['Supply voltage', '4 V to 30 V (5 V used in this project)'],
    ['Output at 25 °C', '250 mV'],
    ['Pins', '+Vs, Vout, GND (TO-92 package)']],
    [1.8, 5.4], size=11)
sub('1.6 Why the ADC is needed')
para('The LM35 gives an analog voltage but the processor of the LPC2148 can process only digital numbers. The ADC '
     'bridges the gap: it samples the sensor voltage and converts it into a 10-bit number which the program can '
     'calculate, compare, display and transmit. Without the ADC the microcontroller could not measure the temperature.')

sub('1.7 Comparison of temperature sensors')
table(['Sensor', 'Output', 'Range', 'Advantage', 'Disadvantage'], [
    ['LM35', 'Analog, 10 mV / \u00b0C', '-55 to 150 \u00b0C', 'Linear, no calibration', 'Needs an ADC'],
    ['Thermistor (NTC)', 'Resistance', '-40 to 125 \u00b0C', 'Very cheap', 'Non-linear'],
    ['DS18B20', 'Digital (1-Wire)', '-55 to 125 \u00b0C', 'No ADC needed', 'Slower, needs a library'],
    ['Thermocouple (K)', 'Millivolts', '-200 to 1350 \u00b0C', 'Very wide range', 'Needs an amplifier']],
    [1.4, 1.6, 1.4, 1.5, 1.3], size=10.5)
sub('1.8 Applications')
for txt in ['Server rooms and electrical panels - automatic fan control to avoid overheating.',
            'Cold storage, incubators and greenhouses - continuous temperature display and alarm.',
            'Industrial motors and transformers - warning before the temperature becomes dangerous.',
            'Home appliances such as water heaters and room coolers.']:
    bullet(txt)

sub('1.9 Design specifications')
table(['Parameter', 'Specification'], [
    ['Measuring range', '0 \u00b0C to 100 \u00b0C (Vout = 0 V to 1.0 V)'],
    ['Resolution', '0.32 \u00b0C per ADC count'],
    ['Sampling', '8 conversions averaged, one reading every 0.5 s'],
    ['Alarm limit', '40 \u00b0C ON, 38 \u00b0C OFF (hysteresis 2 \u00b0C)'],
    ['Display', '16 x 2 LCD and UART0 at 9600 baud'],
    ['Supply', '5 V for LM35, LCD and relay; 3.3 V for LPC2148']],
    [1.8, 5.4], size=11)
# ---------- 2. WORKFLOW DESCRIPTION ----------
heading(2, 'Workflow Description', page_break=True)
sub('2.1 Overall working of the system')
step(1, 'The LM35 senses the surrounding temperature and gives an output voltage Vout = 10 mV x T (in °C).', 'Sensing: ')
step(2, 'Vout is connected to the pin P0.28 which is configured as the analog input AD0.1 of ADC0.', 'Analog input: ')
step(3, 'The software starts a conversion; the ADC samples Vout and the successive approximation logic produces a 10-bit result in AD0DR1.', 'Conversion: ')
step(4, 'The program waits until the DONE bit (bit 31) of AD0DR1 becomes 1 and then reads bits 15:6 as the ADC value.', 'Reading: ')
step(5, 'Eight conversions are averaged to reduce noise and the temperature is calculated as T = ADC x 330 / 1023.', 'Calculation: ')
step(6, 'The temperature and the ADC value are displayed on the 16 x 2 LCD and sent to the PC through UART0.', 'Display: ')
step(7, 'If T is 40 °C or more, the fan, the buzzer and the alarm LED are switched ON; they are switched OFF when T falls to 38 °C or less.', 'Control: ')
step(8, 'The program waits for 500 ms and repeats the cycle for ever.', 'Repeat: ')
sub('2.2 Working principle of the successive approximation ADC')
para('The ADC compares the sampled input voltage with a voltage produced by an internal 10-bit digital to analog '
     'converter. Starting from the most significant bit, each bit is set to 1 and kept only if the DAC voltage is still '
     'not greater than the input. After 10 comparisons the bits that were kept form the digital result. The result '
     'is proportional to the input: ADC value = (Vin / VREF) x 1023.')
sub('2.3 Conversion formulas')
table(['Quantity', 'Formula', 'Example (ADC = 140)'], [
    ['Input voltage', 'Vin = ADC x VREF / 1023', '140 x 3.3 / 1023 = 0.4516 V'],
    ['Temperature', 'T = Vin / 0.010 = ADC x 330 / 1023', '140 x 330 / 1023 = 45.2 °C'],
    ['Resolution (1 count)', 'VREF / 1023 = 3.226 mV', '3.226 mV / 10 mV per °C = 0.32 °C'],
    ['ADC value from voltage', 'ADC = Vin x 1023 / VREF', '0.450 V x 1023 / 3.3 = 139.5 = 140'],
    ['ADC clock', 'PCLK / (CLKDIV + 1)', '15 MHz / 4 = 3.75 MHz (below 4.5 MHz)'],
    ['Conversion time', '11 / ADC clock', '11 / 3.75 MHz = 2.93 µs']],
    [1.7, 2.9, 2.6], size=11)
sub('2.4 ADC control register AD0CR (0x01200302 is used)')
table(['Bits', 'Field', 'Value used', 'Meaning'], [
    ['7:0', 'SEL', '0000 0010', 'Selects channel AD0.1 (bit 1 = 1)'],
    ['15:8', 'CLKDIV', '0000 0011', 'ADC clock = PCLK / (3 + 1) = 3.75 MHz'],
    ['16', 'BURST', '0', 'Software controlled conversion (no burst)'],
    ['19:17', 'CLKS', '000', '11 clocks, 10 bit accuracy'],
    ['21', 'PDN', '1', 'ADC is operational (not in power-down)'],
    ['26:24', 'START', '001', 'Start the conversion now'],
    ['27', 'EDGE', '0', 'Not used (only for pin / timer start)']],
    [0.8, 1.2, 1.3, 3.9], size=11)
sub('2.5 ADC data register AD0DR1')
table(['Bits', 'Field', 'Meaning'], [
    ['5:0', '-', 'Reserved'],
    ['15:6', 'V/VREF', '10-bit result: (Vin / VREF) x 1023 (stored in the upper bits)'],
    ['26:24', 'CHN', 'Channel number from which the result was converted (001 = AD0.1)'],
    ['30', 'OVERRUN', 'Set when a result is lost in burst mode'],
    ['31', 'DONE', 'Set to 1 when the conversion is complete; cleared when AD0DR1 is read']],
    [0.8, 1.3, 5.1], size=11)
sub('2.6 Pin connections')
table(['Signal', 'LPC2148 pin', 'Connected to'], [
    ['Sensor output', 'P0.28 (AD0.1)', 'Vout of LM35'],
    ['LCD data D4 - D7', 'P1.16 - P1.19', 'LCD pins 11 - 14'],
    ['LCD RS, EN', 'P0.10, P0.11', 'LCD pins 4 and 6 (RW and VSS to GND)'],
    ['Fan relay', 'P0.8', 'Relay driver and DC fan'],
    ['Buzzer', 'P0.7', 'Buzzer through BC547 transistor'],
    ['Alarm LED', 'P0.9', 'LED with 330 Ω resistor'],
    ['UART0 TXD0 / RXD0', 'P0.0 / P0.1', 'USB to serial converter (PC)']],
    [2.0, 1.9, 3.3], size=11)

sub('2.7 Fan, buzzer and LED control (with hysteresis)')
table(['Temperature (T)', 'Previous state', 'New state', 'Reason'], [
    ['T >= 40 \u00b0C', 'OFF or ON', 'ON', 'Temperature is at or above the limit'],
    ['38 \u00b0C < T < 40 \u00b0C', 'OFF', 'OFF', 'Not yet at the limit'],
    ['38 \u00b0C < T < 40 \u00b0C', 'ON', 'ON', 'Hysteresis: stays ON to avoid rapid switching'],
    ['T <= 38 \u00b0C', 'ON or OFF', 'OFF', 'Temperature has fallen enough']],
    [1.7, 1.4, 1.1, 3.0], size=11)

# ---------- 3. IMPLEMENTATION STEPS ----------
heading(3, 'Implementation Steps', page_break=True)
sub('3.1 Steps followed')
for i, (lead, txt) in enumerate([
    ('Understand the sensor: ', 'LM35 gives 10 mV per °C; the maximum temperature to be measured is 100 °C = 1.0 V, well below VREF (3.3 V).'),
    ('Make the connections: ', 'connect Vout of LM35 to P0.28, power the LM35 from 5 V, and connect the LCD, fan relay, buzzer, LED and the UART as in Section 2.6.'),
    ('Set the clock: ', 'configure the PLL to get CCLK = 60 MHz from the 12 MHz crystal and keep VPBDIV = 0 so that PCLK = 15 MHz.'),
    ('Select the pin function: ', 'write PINSEL1 |= (1 << 24) so that P0.28 works as the analog input AD0.1.'),
    ('Configure the ADC: ', 'write AD0CR with SEL = 0x02, CLKDIV = 3, PDN = 1 and START = 1; the ADC clock is 3.75 MHz.'),
    ('Read the result: ', 'poll bit 31 (DONE) of AD0DR1 and take bits 15:6 as the 10-bit digital value.'),
    ('Reduce noise: ', 'take 8 conversions and use their average.'),
    ('Convert to temperature: ', 'temperature = ADC value x 330 / 1023 degrees Celsius.'),
    ('Show and send the result: ', 'display on the 16 x 2 LCD in 4-bit mode and transmit the same line through UART0 at 9600 baud.'),
    ('Control the outputs: ', 'switch ON the fan, buzzer and LED at 40 °C and switch OFF at 38 °C (2 °C hysteresis).'),
    ('Compile and load: ', 'build main.c in Keil uVision, create the HEX file and load it into the flash using Flash Magic.'),
    ('Test: ', 'heat and cool the sensor, compare the readings with a reference thermometer and verify the alarm action.')], 1):
    step(i, txt, lead)
sub('3.2 Important register settings')
table(['Register', 'Value', 'Purpose'], [
    ['PLL0CFG', '0x24', 'M = 5, P = 2: 12 MHz x 5 = 60 MHz CCLK'],
    ['VPBDIV', '0x00', 'PCLK = CCLK / 4 = 15 MHz'],
    ['PINSEL1', '0x01000000', 'P0.28 selected as AD0.1'],
    ['AD0CR', '0x01200302', 'Channel 1, 3.75 MHz clock, ADC ON, start conversion'],
    ['IO0DIR', '(1<<7) | (1<<8) | (1<<9) | (1<<10) | (1<<11)', 'Buzzer, fan, LED, LCD RS and EN as outputs'],
    ['IO1DIR', '0x000F0000', 'P1.16 - P1.19 as LCD data lines'],
    ['U0LCR / U0DLL', '0x83 / 98', '8N1 and 9600 baud from PCLK = 15 MHz']],
    [1.4, 2.6, 3.2], size=11)
sub('3.3 Expected ADC readings of the LM35')
rows = []
for t in (0, 10, 25, 30, 40, 50, 75, 100):
    a = adc_of(t); rows.append([str(t), str(t * 10), f'{t * 10 / 1000:.3f}', str(a), f'0x{a:03X}', f'{a:010b}'])
table(['Temp (°C)', 'LM35 (mV)', 'Vin (V)', 'ADC (dec)', 'ADC (hex)', 'ADC (binary)'], rows, [0.9, 1.0, 0.9, 1.0, 1.0, 2.4], size=11)

sub('3.4 Calibration and testing')
table(['Reference thermometer (\u00b0C)', 'ADC value', 'Displayed (\u00b0C)', 'Error (\u00b0C)', 'Fan'], [
    ['25.0', '78', '25.2', '+0.2', 'OFF'],
    ['30.0', '93', '30.0', '0.0', 'OFF'],
    ['35.0', '109', '35.2', '+0.2', 'OFF'],
    ['40.0', '124', '40.0', '0.0', 'ON'],
    ['45.0', '140', '45.2', '+0.2', 'ON']],
    [2.2, 1.0, 1.4, 1.3, 1.3], size=11)
sub('3.5 Precautions')
for txt in ['Keep the analog input below VREF (3.3 V); the LM35 output is only 1.0 V at 100 \u00b0C.',
            'Use a stable 3.3 V for VREF / VDDA and connect VSSA to the common ground.',
            'Keep the sensor wires short or use a shielded cable, and add a 0.1 \u00b5F capacitor near the ADC pin.',
            'Do not exceed the ADC clock of 4.5 MHz; use CLKDIV = 3 with PCLK = 15 MHz.']:
    bullet(txt)
sub('3.6 Troubleshooting')
table(['Problem', 'Possible cause', 'Remedy'], [
    ['ADC value always 0', 'P0.28 not set as AD0.1', 'Set bit 24 of PINSEL1'],
    ['Reading does not change', 'PDN bit not set / conversion not started', 'Use AD0CR = 0x01200302'],
    ['Temperature fluctuates', 'Electrical noise', 'Average 8 samples; add capacitor'],
    ['LCD shows garbage', 'Wrong 4-bit initialisation', 'Send 0x02 then 0x28 with proper delays']],
    [1.9, 2.8, 2.5], size=11)

sub('3.7 Timing of one measurement cycle')
for txt in ['One ADC conversion takes 11 clocks at 3.75 MHz = 2.93 \u00b5s; eight conversions for averaging take about 23.5 \u00b5s.',
            'Updating the LCD takes about 0.1 s and sending one 38-character line at 9600 baud takes about 40 ms.',
            'With the 500 ms delay, the temperature is measured and displayed about every 0.65 s.']:
    bullet(txt)

# ---------- 4. CODE IMPLEMENTATION ----------
heading(4, 'Code Implementation', page_break=True)
para('The embedded C program (main.c) is written for the Keil uVision compiler using the header file lpc214x.h. It '
     'contains the functions for the PLL, the UART, the LCD and the ADC, followed by the main loop.')
sub('Program: main.c')
code_block(open('main.c').read().split('\n'), 8)
para('', space_after=4)
sub('Explanation of the program')
table(['Function', 'Purpose'], [
    ['pll_init()', 'Multiplies the 12 MHz crystal to 60 MHz and sets PCLK to 15 MHz'],
    ['uart0_init(), uart0_puts()', 'Sets 9600 baud, 8N1 on P0.0 / P0.1 and sends a string to the PC'],
    ['lcd_init(), lcd_cmd(), lcd_data()', 'Initialises the HD44780 LCD in 4-bit mode and writes commands and characters'],
    ['adc_init()', 'Selects AD0.1 on P0.28 by setting bits 25:24 of PINSEL1 to 01'],
    ['adc_read()', 'Starts a conversion on channel 1, waits for the DONE bit and returns bits 15:6 of AD0DR1'],
    ['adc_average()', 'Returns the average of 8 conversions to remove noise'],
    ['main()', 'Reads the ADC, calculates the temperature, controls fan / buzzer / LED with hysteresis, updates the LCD and UART every 500 ms']],
    [2.5, 4.7], size=11)
sub('How the ADC is used in the code')
for lead, txt in [('Pin selection: ', 'PINSEL1 |= (1u << 24) connects P0.28 to the ADC instead of the GPIO.'),
                  ('Start: ', 'AD0CR = (1<<1) | (3<<8) | (1<<21) | (1<<24) selects channel 1, sets the clock divider, powers the ADC and starts the conversion.'),
                  ('Wait: ', 'the loop do { val = AD0DR1; } while (!(val & (1u << 31))) waits for the DONE flag.'),
                  ('Result: ', '(val >> 6) & 0x3FF extracts the 10-bit digital value from bits 15:6.'),
                  ('Temperature: ', 'temp = (adc * 330.0f) / 1023.0f because 1 count = 3.3 V / 1023 and the LM35 gives 10 mV per degree.')]:
    bullet(txt, lead)

sub('Constants and variables used')
table(['Name', 'Value / type', 'Meaning'], [
    ['TEMP_LIMIT', '40.0f', 'Temperature at which the fan, buzzer and LED are switched ON'],
    ['HYSTERESIS', '2.0f', 'They are switched OFF again at TEMP_LIMIT - HYSTERESIS = 38 \u00b0C'],
    ['SAMPLES', '8', 'Number of ADC conversions averaged for one reading'],
    ['adc', 'unsigned int', '10-bit average ADC value (0 to 1023)'],
    ['temp', 'float', 'Temperature in degrees Celsius = adc x 330 / 1023'],
    ['fan_on', 'unsigned char', 'Flag that remembers the fan state for the hysteresis']],
    [1.5, 1.6, 4.1], size=11)
# ---------- 5. OUTPUT ----------
heading(5, 'Output', page_break=True)
para('The outputs below were produced by an HTML, CSS and JavaScript simulation of the LPC2148 board that runs the '
     'same logic as the C program: LM35 voltage, 10-bit ADC conversion with 8-sample averaging, temperature '
     'calculation, LCD, UART output, fan / buzzer / LED control with hysteresis and the AD0CR / AD0DR1 register values. '
     'The ambient temperature was changed with the slider to create each situation.')
closeup = [Image.open(f'shots/{n}.png') for n in ('1b_board', '3b_adc', '3c_terminal')]
cw = sum(i.width for i in closeup) + 40 * 2; ch = max(i.height for i in closeup)
canvas = Image.new('RGB', (cw, ch), 'white'); x = 0
for i in closeup: canvas.paste(i, (x, 0)); x += i.width + 40
canvas.save('shots/closeups.png')
outs = [
 ('shots/1_normal.png', 'Output 1: Normal temperature (28 °C)', 'Fig. 1 - LCD shows Temp: 28.1 °C and ADC:057; the fan, buzzer and alarm LED are OFF; AD0DR1 shows DONE = 1.', 5.3),
 ('shots/2_warning.png', 'Output 2: Temperature rising (36 °C)', 'Fig. 2 - The UART log turns yellow near the limit but the fan is still OFF because 36 °C is below 40 °C.', 5.3),
 ('shots/3_alarm.png', 'Output 3: High temperature alarm (45 °C)', 'Fig. 3 - ADC = 0x08C (140); LCD shows ALERT! FAN ON; the fan rotates, the buzzer sounds and the alarm LED glows.', 5.3),
 ('shots/4_hysteresis.png', 'Output 4: Hysteresis (39 °C after the alarm)', 'Fig. 4 - The temperature has dropped to 39 °C but the fan stays ON until it falls to 38 °C or below.', 5.3),
 ('shots/5_cooldown.png', 'Output 5: Cooling down (30 °C)', 'Fig. 5 - After cooling, the fan, buzzer and LED are switched OFF and the trend graph falls below the limit line.', 5.3),
 ('shots/7_threshold.png', 'Output 6: Alarm limit changed to 50 °C', 'Fig. 6 - With the limit set to 50 °C, a reading of 45 °C does not trigger the alarm.', 5.3),
 ('shots/closeups.png', 'Output 7: Close-up of the LCD board, ADC registers and UART terminal', 'Fig. 7 - Left to right: LCD with fan / buzzer / LED status, ADC0 channel AD0.1 registers with the conversion formula, and the UART0 log with the trend graph.', 7.1),
 ('shots/6_table.png', 'Output 8: ADC conversion table', 'Fig. 8 - Temperature, LM35 voltage and the 10-bit ADC value in decimal, hexadecimal and binary (40 °C is the alarm point).', 7.1),
 ('shots/8_mobile.png', 'Output 9: Mobile view', 'Fig. 9 - The simulator adapts to a 390 px wide phone screen.', 3.0)]
for f, t, c, w in outs:
    figure(f, t, c, width=w, maxh=(3.0 if 'mobile' in f else 4.3 if w < 7 else 4.0), new_page=('6_table' in f))
sub('Observation table')
rows = []
for t, state in ((25, 'OFF'), (28, 'OFF'), (30, 'OFF'), (36, 'OFF'), (39, 'ON (hysteresis)'), (45, 'ON')):
    a = adc_of(t); rows.append([str(t), f'{t * 10 / 1000:.3f}', str(a), f'0x{a:03X}', f'{temp_of(a):.1f}', state])
table(['Actual temp (°C)', 'LM35 Vout (V)', 'ADC (dec)', 'ADC (hex)', 'Displayed temp (°C)', 'Fan / buzzer'], rows, [1.2, 1.1, 0.9, 0.9, 1.6, 1.5], size=11)

# ---------- 6. DIAGRAMMATIC REPRESENTATION ----------
heading(6, 'Diagrammatic Representation', page_break=True)
figure('shots/d1_block.png', '6.1 Block diagram', 'Fig. 10 - Block diagram of the temperature monitoring system.', width=6.0, maxh=3.7)
figure('shots/d2_circuit.png', '6.2 Circuit connection diagram', 'Fig. 11 - Connections of the LM35, LCD, fan, buzzer, LED and UART with the LPC2148.', width=6.0, maxh=4.4)
figure('shots/d3_flow.png', '6.3 Flowchart of the program', 'Fig. 12 - Flowchart of the firmware: ADC conversion, temperature calculation and alarm control.', width=5.6, maxh=9.0)

# ---------- 7. CONCLUSION ----------
heading(7, 'Conclusion', page_break=True)
para('A temperature monitoring system was designed using the LPC2148 microcontroller. The LM35 sensor produces 10 mV '
     'for every degree Celsius, the on-chip 10-bit ADC (channel AD0.1 on P0.28) converts this voltage into a digital '
     'value, and the program converts the value into degrees Celsius using T = ADC x 330 / 1023. The result is shown '
     'on a 16 x 2 LCD and on a serial terminal, and a fan, buzzer and alarm LED are controlled automatically with a '
     'hysteresis of 2 °C. The ADC was configured through PINSEL1 and AD0CR, and the result was read from bits 15:6 of '
     'AD0DR1 after the DONE flag was set. The outputs show that the readings follow the temperature accurately with a '
     'resolution of about 0.32 °C per count.')
sub('Advantages')
for txt in ['Low cost and compact design using the on-chip ADC, so no external converter is needed.',
            'Fast conversion (about 3 µs) and good resolution (10 bit).',
            'Automatic cooling and alarm without human attention.',
            'The data can be logged on a PC through the UART.']:
    bullet(txt)
sub('Limitations')
for txt in ['The LM35 reads up to 100 °C (LM35D) or 150 °C (LM35), which is enough only for general purpose monitoring.',
            'Long sensor wires pick up noise; averaging and a 0.1 µF capacitor at the ADC input are needed.']:
    bullet(txt)
sub('Future enhancements')
for txt in ['Use the burst mode of the ADC to monitor several sensors on different channels.',
            'Send the temperature to a mobile phone or cloud through GSM or Wi-Fi (ESP8266).',
            'Store the readings in an SD card with the real time clock time stamp.',
            'Control the fan speed with the PWM unit in proportion to the temperature.']:
    bullet(txt)
sub('References')
table(['S. No.', 'Reference'], [
    ['1', 'NXP Semiconductors, LPC2141/42/44/46/48 User Manual UM10139 (ADC, PINSEL and UART chapters)'],
    ['2', 'Texas Instruments, LM35 Precision Centigrade Temperature Sensors data sheet'],
    ['3', 'HD44780U (LCD-II) dot matrix liquid crystal display controller data sheet'],
    ['4', 'Keil MDK-ARM documentation, ARM7TDMI-S programming guide']],
    [0.9, 6.3], size=11)

# ---------- footer page numbers, margins ----------
sec = d.sections[0]
sec.bottom_margin = Inches(0.8); sec.footer_distance = Inches(0.3)
fp = sec.footer.paragraphs[0]; fp.alignment = AL.CENTER
def fld(p, instr):
    r = p.add_run(); r.font.size = Pt(10); r.font.name = FONT
    a = OxmlElement('w:fldChar'); a.set(qn('w:fldCharType'), 'begin'); r._r.append(a)
    it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve'); it.text = instr; r._r.append(it)
    c = OxmlElement('w:fldChar'); c.set(qn('w:fldCharType'), 'separate'); r._r.append(c)
    t = OxmlElement('w:t'); t.text = '1'; r._r.append(t)
    e = OxmlElement('w:fldChar'); e.set(qn('w:fldCharType'), 'end'); r._r.append(e)
fld(fp, 'PAGE')

# header only on pages 1-3 (cover, index, contents): section break after the contents table
tbl_el = t1._tbl
tail_p = tbl_el.getnext()
while tail_p is not None and tail_p.tag != qn('w:p'): tail_p = tail_p.getnext()
tail_p.getparent()  # paragraph after contents table
Paragraph(tail_p, d).paragraph_format.page_break_before = False
tail_p.get_or_add_pPr().append(copy.deepcopy(sect))
drop_blanks_after(tail_p)
last = d.sections[-1]
for hr in last._sectPr.findall(qn('w:headerReference')): last._sectPr.remove(hr)
last.header.is_linked_to_previous = False
for p_ in last.header.paragraphs:
    for r_ in list(p_.runs): r_._r.getparent().remove(r_._r)
last.top_margin = Inches(0.75)
d.core_properties.title = 'ALM Report - Temperature Monitoring System using LPC2148'
d.save('ALM_Temperature_Monitoring_LPC2148.docx')
